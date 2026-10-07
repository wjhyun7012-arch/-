"""발표자료 v7 → v8: (1) 내용 수정 (2) 백업 71~80 삭제 (3) A4 가로 변환.
사용: python3 -I mk_deck8.py v7.pptx out.pptx   (fig8/s3_6.png, pptlib.py 필요)
"""
import sys, json, copy, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptlib import *
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

def remove_shape(prs, sn, sid, why=''):
    s = prs.slides[sn - 1]; sh = shape(s, sid); txt = sh.text_frame.text if sh.has_text_frame else ''
    sh._element.getparent().remove(sh._element); LOG.append((sn, f'[{sid}] 삭제', txt, '(삭제)', why))

def set_par_text(p, txt):
    rs = p.runs
    if rs:
        rs[0].text = txt
        for r in rs[1:]: r._r.getparent().remove(r._r)
    else: p.text = txt

def notes_append(prs, sn, txt, why=''):
    tf = prs.slides[sn - 1].notes_slide.notes_text_frame
    p0 = tf.paragraphs[-1]; newp = copy.deepcopy(p0._p); p0._p.addnext(newp)
    from pptx.text.text import _Paragraph
    set_par_text(_Paragraph(newp, p0._parent), txt); LOG.append((sn, '발표 노트 추가', '', txt, why))

# ================= (1) 내용 수정 =================
# ---- 1번 표지: 공개발표 → 예비심사, 날짜
rep(prs, 1, 4, '박사학위 논문 공개발표', '박사학위 논문 예비심사', '인수인계 4-9: 예비심사(10/22)')
rep(prs, 1, 7, '2026. 11. ○○.', '2026. 10. 22.', '예비심사 날짜')

# ---- 2번 목차: 큰 제목 여섯 줄만, 작은 글씨는 노트로, 뼈대 한 줄 삭제
s2 = prs.slides[1]; tb = shape(s2, 5); tf = tb.text_frame
old_lines = [p.text for p in tf.paragraphs]
big = [p for p in tf.paragraphs if p.runs and p.runs[0].font.bold]
small = [p for p in tf.paragraphs if not (p.runs and p.runs[0].font.bold)]
for p in small: p._p.getparent().remove(p._p)
for p in tf.paragraphs:
    for r in p.runs: r.font.size = Pt(28)
    p.space_after = Pt(18)
tb.top = Emu(1143000)
LOG.append((2, '[5] 목차', '\n'.join(old_lines), '\n'.join(p.text for p in tf.paragraphs) + '\n(28pt, 작은 글씨 6줄은 노트로)', '인수인계 4-1: 큰 제목 여섯 줄만 키워서'))
for sid in [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21]:
    remove_shape(prs, 2, sid, '뼈대 한 줄 삭제(인수인계 4-1)')
bot = shape(s2, 22); bot.top = Emu(5486400)
rep_notes(prs, 2, '뼈대 한 줄(만든 것 16 → 55 → 4 → ML0~ML5, 써 본 것 2.85 → ML2 → 3.15 → ML3)을 한 번 짚고, 맨 아래 한 줄을 그대로 말한다.',
          '뼈대 한 줄(만든 것 16종 → 55문항 → 핵심 4종 → ML0~ML5, 써 본 것 2.85 → ML2 → 3.15 → ML3)은 화면에 없으므로 말로만 한 번 짚고, 맨 아래 한 줄을 그대로 말한다.', '뼈대 한 줄이 화면에서 빠짐')
notes_append(prs, 2, '(화면에서 뺀 세부 — 나만 보는 것) Ⅰ 왜 지금인가 · 두 표준이 서로 비워 둔 곳 · 「LCA를 했다」의 두 상태 · 연구 문제 / Ⅱ ISO 14001의 전과정 관점 · LCA의 4단계 · 따른 표준 · 선행연구와 연구 공백 / Ⅲ LMA 3계층 · 조항 매핑 · 연계 지점 · 16종 · 55문항 · 판정 규칙 · 검증 / Ⅳ 적용 개요 · 객관성의 확보 · 영역별 결과 · 종합 판정 ML2 · 결과 요약 / Ⅴ 개선 과제 R1~R5 · 개선 로드맵과 재평가 · 장점과 단점, 한계 / 맺음 연구 전체 한 장 요약', '이사님 「나만 볼 수 있도록」')

# ---- 3번: 제목을 [그림 1-4] 캡션으로, 그림 안 두 상자 문구
rep(prs, 3, 2, '논문 전체 구조도 — 다섯 장의 큰 흐름', '논문의 구성과 장 간 연결', '인수인계 4-3 (가): [그림 1-4] 캡션')
set_img(prs, 3, 6, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fig8/s3_6.png'), '그림 머리말·3장 상자(네 지점 가운데 하나라도 빠지면 … / 모델의 타당성은 두 방향으로 검증)·4장 상자(낮은 점수의 원인 — 미비 사항 3건)를 [그림 1-4]대로, 컬러판 유지')

# ---- 4·8·12·19번: 「물음 — …」 줄 삭제, 12번 답은 노트로
remove_shape(prs, 4, 20, '인수인계 4-4 (답은 노트에 이미 있음)')
remove_shape(prs, 8, 10, '인수인계 4-4 (답은 노트에 이미 있음)')
remove_shape(prs, 12, 9, '인수인계 4-4 — 답은 노트로')
notes_append(prs, 12, '물음(모델을 만들었는데 검증은 했나)이 나오면 — 두 방향. 추적성 검증(55문항을 근거 조항까지)과 표준 적합성 검증(ISO/IEC 33004 요구사항 24건 자체점검 — 충족 22 · 해당 없음 1 · 보완 1). 전문가 내용타당도 조사(CVI)는 조사지 설계까지, 실시는 향후 과제 — 3.5.', '12번 화면의 물음 줄을 노트로')
remove_shape(prs, 19, 17, '인수인계 4-4 (답은 노트에 이미 있음)')

# ---- 26번 장점·단점·한계: 쉬운 말로(항목 이름은 [표 5-1] 그대로)
def set_box_lines(sn, sid, lines, size_pt, why):
    sh = shape(prs.slides[sn - 1], sid); tf = sh.text_frame; old = tf.text
    ps = tf.paragraphs
    while len(ps) < len(lines):
        newp = copy.deepcopy(ps[-1]._p); ps[-1]._p.addnext(newp); ps = tf.paragraphs
    for extra in ps[len(lines):]: extra._p.getparent().remove(extra._p)
    for p, txt in zip(tf.paragraphs, lines):
        set_par_text(p, txt)
        if size_pt and not p.runs[0].font.bold:
            for r in p.runs: r.font.size = Pt(size_pt)
    LOG.append((sn, f'[{sid}]', old, '\n'.join(lines), why))
set_box_lines(26, 8, ['장점',
    '· 문항의 추적성 — 문항마다 근거 조항을 따라갈 수 있다',
    '· 판정 기준의 명문화 — 점수 기준을 미리 적어, 같은 증거면 같은 점수',
    '· 종합점수에 묻히는 미비의 식별 — 평균이 높아도 꼭 있어야 할 네 곳이 빠지면 드러난다',
    '· 평가와 개선의 연결 — 고칠 목표가 다음 평가 기준과 같은 말',
    '· 표준 구조의 준용 — 판정 구조·평가 방법을 국제표준에서 가져옴',
    '· 수준이 뜻하는 바의 구분 — ML3까지 체계 갖추기, ML4부터 숫자로 확인'],
    14, '인수인계 부록 「26번」 문안을 항목 전부 두고 풀이만 짧게(이사님 10/7: 「다 넣어야 되는데 간단하게」)')
set_box_lines(26, 9, ['단점',
    '· 핵심 문항 4종에 대한 의존 — 네 문항이 수준을 정하는데, 그 선택은 전문가 검토 전',
    '· 판정 파라미터 — 과반 50% 등은 연구자가 정한 값, 상위 두 수준은 사례 없음',
    '· 평가의 부담 — 55문항 문서 확인, ISO 14001 심사와 LCA를 함께 알아야',
    '· 적용 대상의 제약 — LCA를 하지 않는 회사는 점수 차이가 안 난다',
    '· 문항 간 종속 · 문서 심사 방식 — 미비 하나가 여러 점수를 끌어내리고, 형식적 이행은 못 가려냄',
    '· 환경성과와의 관계 미검증 — 수준이 높으면 탄소·비용이 주는지는 미확인'],
    14, '같은 문안')
rep(prs, 26, 10, '한계 — 단일 사례 · 단일 시점 · 단일 평가자  /  전문가 내용타당도 조사(CVI) 미실시  /  ML4~ML5는 사례로 미검증',
    '한계 — 한 회사 · 한 시점 · 평가자 한 명  /  전문가 타당도 조사(CVI)는 아직 하지 않음  /  가장 높은 두 수준(ML4·ML5)은 사례로 확인 못 함', '같은 문안(빨간 상자)')
for sid in (8, 9): shape(prs.slides[25], sid).height = Emu(3310000)   # 상자 세로 늘림(16:9 기준 → A4에서 3.65M)
sh = shape(prs.slides[25], 10); sh.top = Emu(4800000)
sh = shape(prs.slides[25], 11); sh.top = Emu(5600000)

# ---- 30번 백업 구분: 「B-Q표」 빼기
rep(prs, 30, 2, ' · B-Q표 예상 질문 43문 전문', '', '인수인계 4-7: 43문 표는 예상 질문 v5 파일로만')

# ================= (2) 백업 71~80 삭제 =================
sldIdLst = prs.slides._sldIdLst; ids = list(sldIdLst)
for sldId in ids[70:]:
    rId = sldId.get(qn('r:id')); prs.part.drop_rel(rId); sldIdLst.remove(sldId)
LOG.append(('71~80', '슬라이드 삭제', '예상 질문 43문 표 10장(B-Q표 1~10)', '(삭제 — 예상 질문 v5 파일로 대체, 80 → 70장)', '인수인계 4-7'))

# ================= (3) A4 가로 변환 =================
W0, H0 = prs.slide_width, prs.slide_height
W1, H1 = 10692000, 7560000
sx, sy = W1 / W0, H1 / H0
def sc(v, f): return Emu(int(round(v * f)))
for s in prs.slides:
    for sh in s.shapes:
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        if sh.shape_type == 13:   # 그림: 비율 유지, 상자 안 가운데
            nw, nh = w * sx, h * sx
            sh.left, sh.top, sh.width, sh.height = sc(x, sx), sc(y * sy + (h * sy - nh) / 2, 1), sc(nw, 1), sc(nh, 1)
        else:
            sh.left, sh.top, sh.width, sh.height = sc(x, sx), sc(y, sy), sc(w, sx), sc(h, sy)
        if sh.has_table:
            t = sh.table
            for c in t.columns: c.width = sc(c.width, sx)
            for r in t.rows: r.height = sc(r.height, sy)
prs.slide_width, prs.slide_height = Emu(W1), Emu(H1)
LOG.append(('전체', '슬라이드 크기', '16:9 (12191695×6858000 EMU)', 'A4 가로 297×210mm (10692000×7560000 EMU) — 위치·크기 비례(가로 0.877·세로 1.102), 그림은 비율 유지해 가운데, 표는 열 너비·행 높이 비례, 글꼴 크기 유지', '이사님 결정 「슬라이드를 A4로 맞추고」'))

prs.save(OUT)
json.dump(LOG, open(os.path.join(os.path.dirname(OUT) or '.', 'log8.json'), 'w'), ensure_ascii=False, indent=1, default=str)
print(len(LOG), 'changes;', len(prs.slides), 'slides')
