# v29 -> v29a: 검토자 지적 반영(논문 일치·용어), 새 번호 기준
import sys, json
sys.path.insert(0, '.')
from pptx import Presentation
from pptlib import rep, rep_notes, LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text
    assert t.count(old) == 1, (sn, old, t.count(old)); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))

# 5
note_rep(5, '중대한 환경측면을 정하고', '중대한 환경측면을 결정하고', '논문 용어 「결정」(1.2.1)')
# 6 — 논문 고정 문구(2.3.2)
set_one(6, 5, '▶ ISO 14001은 정량 방법을, ISO 14044는 결과를 되풀이하여 관리하는 주기를 정하지 않았다 — 그 사이가 연계다.', '논문 고정 문구')
note_rep(6, '그 결과를 경영시스템에 넣는 일은 정하지 않았습니다.', '그 결과를 되풀이하여 관리하는 주기는 정하지 않았습니다.', '논문 고정 문구(2.3.2)')
note_rep(6, '문장은 2.3.2·2.3.3(B432·B435) 그대로.', '문장은 2.3.2(B432·B435) 그대로.', '절 번호')
# 7 — 33020·BARS, 논문 표현
note_rep(7, '척도는 행동기술척도에서 가져왔고, 그 적용 선례가 지도교수 연구진의 공공데이터 품질 성숙도 연구입니다.',
         '측정의 틀은 ISO/IEC 33020, 점수별 상태 기술은 행동기술척도에서 가져왔고, 그 적용 선례가 국내 공공데이터 품질 연구(PPDQM3·APDQM3)입니다.', '검토자 지적 — 3.3.2(B730), 논문 표현(B508)')
# 9 — 25개는 조항
set_one(9, 9, '전과정평가(LCA, ISO 14044)', '보류')  # 두 줄이라 아래에서 다시
sh9 = shape(prs.slides[8], 9); p0 = sh9.text_frame.paragraphs[0]; p0.runs[0].text = '전과정평가(LCA, ISO 14044)'
import copy
np_ = sh9.text_frame.add_paragraph(); r = np_.add_run(); r.text = '조항 25개(전조항 28개 가운데)'
src_r = shape(prs.slides[8], 8).text_frame.paragraphs[1].runs[0]
if src_r._r.rPr is not None: r._r.insert(0, copy.deepcopy(src_r._r.rPr))
if shape(prs.slides[8], 8).text_frame.paragraphs[1]._p.pPr is not None: np_._p.insert(0, copy.deepcopy(shape(prs.slides[8], 8).text_frame.paragraphs[1]._p.pPr))
LOG[-1] = (9, '[9]', '전과정평가(LCA, ISO 14044)\n요구사항 25개', '전과정평가(LCA, ISO 14044)\n조항 25개(전조항 28개 가운데)', '검토자 지적 — 25는 근거 조항 수(3.3.1)')
# 10 — 추가 항목
rep(prs, 10, 15, '추가 항목(should)을 적어 둠', '추가 항목(should 또는 체계화·지속 상태)을 적어 둠', '검토자 지적 — 3.3.2(B725)')
# 11
rep(prs, 11, 22, '핵심 문항 4종 모두 4점,', '핵심 문항 4종 모두 4점 이상,', '[표 3-7] ≥ 4')
note_rep(11, '연계가 아예 끊긴 것과 연계가 얕은 것은 다르기 때문에', '연계가 성립하지 않은 것과 성숙도가 낮은 것은 다르기 때문에', '논문 표현(3.4.1)')
# 12
rep(prs, 12, 27, '문서 심사(면담·현장 관찰 없음)', '문서 심사 — 면담·실행 상태의 현장 관찰은 포함하지 않음', '검토자 지적 — [표 5-1] 표현')
note_rep(12, '평가는 문서 심사만으로, 평가자 한 명이,', '평가는 문서 심사로 면담과 실행 상태의 현장 관찰은 포함하지 않고, 평가자 한 명이,', '검토자 지적')
note_rep(12, '낮은 점수의 원인 세 가지에서 개선 과제 다섯 건을 도출하였고', '구조적 미비 세 가지와 공통 원인 두 가지에서 개선 과제 다섯 건을 도출하였고', '4.5·4.6(B932·B949)')
# 14
note_rep(14, 'LCA를 다시 할 때가 아직 오지 않았습니다.', 'LCA 연 1회 재수행 규정은 있으나 아직 이행 전입니다.', '4.5.1 ③')
# 15
set_one(15, 5, '▶ 구조적 미비 셋과 공통 원인 둘에서 개선 과제 다섯 건 — 다 마치면 종합 3.15점, ML3이 전망된다.', '4.5·4.6')
note_rep(15, '낮은 점수의 원인 세 가지에서 과제 다섯 건을 도출했고', '구조적 미비 세 가지와 공통 원인 두 가지에서 과제 다섯 건을 도출했고', '4.5·4.6')
# 16
note_rep(16, '장단점 표의 단점은 전문가 조사와 다사례 적용으로 보완할 것입니다.', '단점 일곱 가운데 넷은 전문가 조사와 다사례 적용으로 보완하고, 셋은 설계에서 비롯되는 제약입니다.', '5.2(B1018)')
note_rep(16, '[표 5-1] 장점 6·단점 6 전문은', '[표 5-1] 장점 6건·단점 7건(핵심 문항 의존·판정 파라미터·평가 부담·적용 대상 제약·문항 간 종속·문서 심사·환경성과 미검증) 전문은', '단점 7건')
# 17
note_rep(17, '한계는 여덟 항목입니다. 단일 사례·단일 시점·단일 평가자와 상위 수준 미검증은 자료 조건의 경계이고, 전문가 조사와 다사례 적용으로 풉니다.\n문항 간 종속과 타당화 확장은 심화 과제, 표준 판의 범위와 진단·개선 고도화는 유지·고도화 과제이며, 표준화 경로의 검토는 학회·협회를 거친 뒤의 일입니다.',
         '한계는 여덟 항목입니다. 단일 사례·단일 시점·단일 평가자와 상위 수준 미검증은 자료 조건의 경계이고, 다사례·다평가자 적용과 개선 이행 뒤 재평가로 풉니다.\n문항 간 종속과 타당화 확장은 심화 과제로 전문가 조사가 여기에 들어가고, 표준 판의 범위와 진단·개선 고도화는 유지·고도화 과제이며, 외주처리 구간처럼 매핑에서 확인된 미측정 구간은 확장 과제, 표준화 경로의 검토는 학회·협회를 거친 뒤의 일입니다.', '5.3 여덟 항목·대응 경로(B1025~B1035)')
note_rep(17, '〔40초.', '〔45초.', '시간')
prs.save(dst); json.dump(LOG, open('log29a.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1); print('ok', len(LOG))
