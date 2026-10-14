# v29a -> v29b: 목차 상자 삭제, 4번 개요 복원, 6·7·8·17 쉬운 말, 16·18을 v28 쉬운 장(S-27·S-29)으로 되돌림, [그림 5-2]·[그림 5-1] 본 장 제거(백업 B-0c에 이미 있음)
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptlib import rep, LOG, shape
src, ref28, dst = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(src); p28 = Presentation(ref28)
IN = 914400
def set_note(sn, text, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; old = tf.text; tf.text = text; LOG.append((sn, '발표 노트(전체)', old, text, why))
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text; assert t.count(old) == 1, (sn, old); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def set_lines(sn, sid, lines, why, size=None):
    sh = shape(prs.slides[sn-1], sid); tf = sh.text_frame; old = tf.text
    p0 = tf.paragraphs[0]; pPr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    r0 = p0.runs[0]; rPr = copy.deepcopy(r0._r.rPr) if r0._r.rPr is not None else None
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    r0.text = lines[0]
    if size: r0.font.size = Pt(size)
    for ln in lines[1:]:
        p = tf.add_paragraph()
        if pPr is not None: p._p.insert(0, copy.deepcopy(pPr))
        r = p.add_run(); r.text = ln
        if rPr is not None: r._r.insert(0, copy.deepcopy(rPr))
        if size: r.font.size = Pt(size)
    LOG.append((sn, f'[{sid}]', old, '\n'.join(lines), why))
def remove_shape(sn, sid, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text if sh.has_text_frame else '[도형]'; sh._element.getparent().remove(sh._element); LOG.append((sn, f'[{sid}] 삭제', old, '(삭제)', why))

# 2 목차 — 진행 안내 상자 삭제
remove_shape(2, 22, '이사님 10/14 — 화면의 진행 안내 삭제')
# 4 개요 복원
s4 = prs.slides[3]
t = shape(s4, 2); t.text_frame.paragraphs[0].runs[0].text = '이 연구를 한 장으로 — 배경과 요지'; LOG.append((4, '제목', '연구 배경과 필요성', '이 연구를 한 장으로 — 배경과 요지', '개요 복원'))
set_one(4, 5, '▶ 두 표준은 실무에서 별도로 운영되고 연계 수준을 확인할 기준이 없었다 — 그 연계를 조항 단위로 확인하고 성숙도 수준으로 판정하는 모델을 만들었다.', '개요 요지')
set_one(4, 7, '근거  1.1, 1.2.1, [그림 1-1]', '개요')
el = copy.deepcopy(shape(p28.slides[3], 26)._element); s4.shapes._spTree.append(el)
sub = [x for x in s4.shapes if x.has_text_frame and x.text_frame.text.startswith('연계(linkage)란')][0]
set_lines(4, sub.shape_id, ['연계(linkage)란 — LCA 결과가 환경측면 결정·환경목표·설계·개발·모니터링에 쓰이는 것 (핵심 연계 지점 4곳)',
                            '판정 결과는 — ML0(도입 전)부터 ML5(혁신화)까지 여섯 수준 · 만든 것은 연계 프로세스 16종 · 평가 문항 55 · 핵심 문항 4종'], 'v28 요지 두 줄 복원')
shape(s4, 25).height = Emu(int(0.87*IN))
note_rep(4, '그래서 연계된 정도를 조항 단위로 확인하고 성숙도 수준으로 판정하는 모델을 만들었습니다.\n',
         '그래서 연계된 정도를 조항 단위로 확인하고 성숙도 수준으로 판정하는 모델을 만들었습니다.\n연계란 LCA 결과가 환경측면 결정, 환경목표, 설계·개발, 모니터링에 쓰이는 것이고, 판정 결과는 ML0부터 ML5까지 여섯 수준으로 나옵니다. 만든 것은 연계 프로세스 열여섯 가지, 평가 문항 쉰다섯 개, 핵심 문항 네 가지입니다.\n', '개요 두 문장')
note_rep(4, '〔40초.', '〔55초.', '시간')
# 6 쉬운 노트(숫자·용어는 논문대로, 지시문에 논문 문장)
set_note(6, '2장은 두 표준과 그 사이입니다. ISO 14001은 전과정을 생각하라고만 하고 어떻게 하라는 방법은 정하지 않았고, ISO 14044는 방법은 자세히 정했지만 그 결과를 되풀이하여 관리하는 주기는 정하지 않았습니다. 그 사이가 이 연구가 말하는 연계입니다.\n그림의 아래 네 칸이 두 표준이 만나는 자리입니다. 계획에서는 환경측면과 환경목표가 LCA의 범위 정의와 영향평가 결과를 받고, 실행에서는 운용기획이 목록분석과, 점검에서는 모니터링이 재수행 결과와 만나고, 개선에서는 해석의 권고가 다음 환경목표로 돌아갑니다.\n위 계단은 그 연계가 깊어질수록 올라가는 성숙도 수준이고, 뒤로 갈수록 요구되는 수준이 높아집니다.\n〔50초. 논문 문장(2.3.2 B432·B435) — 「Plan 단계에서 EMS의 적용범위(4.3)와 환경측면(6.1.2)·환경목표(6.2.1)는 LCA의 목적·범위 정의와 영향평가 결과를 입력으로 삼고, Do 단계에서 운용기획 및 관리(8.1)는 목록분석의 데이터 수집과 설계·구매 통제와 연결되며, Check 단계에서 모니터링(9.1)과 경영검토(9.3)는 영향평가·해석의 결과와 재수행에 의한 검증을 입력으로 삼고, Act 단계에서 개선(10)은 해석의 권고를 차기 환경목표로 피드백한다」. 점선 = 영역별 요구 수준(P·D는 ML3, C는 ML4, A는 ML5). [표 2-3]은 백업 S-5, ISO 14001 구조 S-8, LCA 4단계 S-9.〕', '쉬운 말로(이사님 10/14)')
# 7 쉬운 노트
set_note(7, '성숙도를 평가하는 방법은 새로 만들지 않았습니다. 소프트웨어와 데이터 분야에서 오래 쓰던 성숙도 평가의 국제표준에서 세 가지를 가져왔습니다. 모델을 세 계층으로 짜는 틀, 꼭 필요한 것이 빠지면 위로 못 올라가는 판정 구조, 그리고 점수를 매기는 척도입니다. 그 적용 선례가 국내 공공데이터 품질 연구입니다.\n오른쪽 환경 쪽 선행연구들은 환경경영 전반이나 제품 설계, LCA 활용 역량 가운데 한 측면씩만 다루었습니다.\n두 표준의 연계 자체를 조항 단위로 판정한 연구는 찾아보기 어려웠고, 그 빈 자리가 이 연구의 위치입니다.\n〔50초. 가져온 것 — 세 계층 구조 ISO/IEC 33004 · 최소충족·비보상 판정 KS X ISO 8000-62 · 측정 틀 ISO/IEC 33020, 점수별 상태 기술 BARS(행동기술척도). 선례 — 국내 공공데이터 연구(PPDQM3 2015·APDQM3 2017). 분야 축 — EMM·EcoM2·LCM CMM. 대비표 [표 2-4]·[표 2-5]는 백업 S-11, 따른 표준 S-10·B-2. 「8000-62를 왜 가져왔나」 → B-Q7, 「분야가 무엇인가」 → 전문 49번.〕', '쉬운 말로(이사님 10/14)')
# 8 아래 상자 — 셈은 지시문으로
set_lines(8, 27, ['ISO 14001 — 요구사항 조항 30개 전부가 문항이 됨   ·   ISO 14044 — 조항 25개가 문항의 근거(나머지 3개도 세부 조항의 문항·체크포인트로 반영) → 빠진 조항 없음. 셈은 부록 A.4',
                  '매우 높음 4개 조항 = 핵심 문항 4종(P-9·P-14·D-1·C-1) — 관리 주기의 계획 둘·실행 하나·점검 하나'], '셈은 지시문·노트북으로', 12.5)
shape(prs.slides[7], 27).height = Emu(int(0.95*IN))
# 17 쉬운 노트
set_note(17, '한계는 여덟 가지인데 크게 셋입니다. 한 회사를 한 번, 한 사람이 평가했다는 것, 전문가 검토를 아직 거치지 않았다는 것, 가장 높은 두 수준은 사례가 없어 확인하지 못했다는 것입니다.\n나머지는 문항 사이의 영향, 표준이 개정되면 매핑을 갱신해야 하는 것, 진단과 개선을 더 정교하게 하는 것, 매핑에서 새로 보인 측정 안 된 구간, 그리고 표준화 경로이고, 항목마다 대응 경로를 적어 두었습니다.\n〔45초. 여덟 항목(5.3) — ① 단일 사례·시점·평가자(→ 다사례·다평가자·개선 뒤 재평가) ② 문항 간 종속(→ 종속 구조 모형화) ③ ML4·ML5 사례 미검증(→ 단계적 축적) ④ 표준 판의 범위 2015(→ 2026년판 매핑 갱신) ⑤ 타당화 확장(→ CVI 조사 실시) ⑥ 진단·개선 고도화 ⑦ 미측정 구간(외주처리) ⑧ 표준화 경로. ML4·ML5는 A사에 재수행 기록이 없어 미확인, 규칙은 3.4.3. CVI → 전문 20번, 개정 → 40번.〕', '쉬운 말로(이사님 10/14)')
# 16 ← S-27(35번): 단점 일곱으로 나누고 장 표시·요지 수정, 본으로
s35 = prs.slides[34]
set_one(35, 3, '결론 — 기여 · 한계 · 답한 것', '본으로 복귀')
set_one(35, 5, '▶ 장점 여섯, 단점 일곱, 한계 셋 — 한계는 그대로 향후 연구 과제다.', '단점 7건(5.2)')
set_one(35, 7, '근거  5.2, 5.3, [표 5-1], [그림 5-3]', '근거')
set_lines(35, 9, ['단점', '· 핵심 문항 4종에 대한 의존 — 네 문항의 선택은 전문가 검토 전', '· 판정 파라미터 — 과반 50% 등은 연구자가 정한 값', '· 평가의 부담 — 55문항 문서 확인, 두 분야를 다 알아야', '· 적용 대상의 제약 — LCA를 안 하는 회사는 차이가 안 난다', '· 문항 간 종속 — 미비 하나가 여러 점수에 함께 영향', '· 문서 심사 방식 — 형식적 이행은 못 가림', '· 환경성과와의 관계 미검증 — 탄소·비용이 주는지는 확인 안 함'], '단점 7건으로 분리', 13.5)
set_note(35, '이 모델의 장점과 단점, 그리고 한계입니다. 장점 여섯, 단점 일곱, 한계 셋입니다.\n장점은 문항의 출처를 조항까지 따라갈 수 있다는 것, 같은 증거면 같은 점수가 나온다는 것, 종합점수에 묻히는 미비를 드러낸다는 것, 평가가 개선과 재평가로 이어진다는 것, 판정 구조와 평가 방법을 국제표준에서 가져왔다는 것, 수준의 뜻을 체계 단계와 숫자 확인 단계로 나눈 것입니다.\n단점은 판정이 핵심 문항 4종에 좌우된다는 것, 기준값이 연구자 확정값이라는 것, 평가에 손이 많이 간다는 것, LCA를 안 하는 회사는 차이가 안 난다는 것, 미비 하나가 여러 점수에 함께 영향을 준다는 것, 문서 심사라 형식적 이행을 가리기 어렵다는 것, 환경성과와의 관계는 검증하지 않았다는 것입니다. 이 가운데 넷은 전문가 조사와 다사례 적용으로 보완하고, 셋은 설계에서 비롯되는 제약입니다.\n한계는 한 회사·한 시점·평가자 한 명, 전문가 타당도 조사 미실시, 가장 높은 두 수준의 사례 미확인이고, 이 셋이 그대로 향후 연구 과제입니다.\n〔1분. 화면 글은 쉬운 말판, 항목 이름은 [표 5-1] 그대로. 단점 넷(핵심 문항 지정·판정 파라미터·문항 간 종속·환경성과)은 보완 가능, 셋(평가 부담·적용 대상·문서 심사)은 설계 제약(5.2). 장점·단점 전문은 백업 B-15a·B-15b, 기여 [그림 5-2]는 B-0c②, 한계 여덟 항목은 다음 장. 「환경성과와의 관계」 → 전문 46번.〕', '단점 7건·넷/셋(5.2)')
# 18 ← S-29(37번): 네 답을 쉬운 네 줄로, 본으로
s37 = prs.slides[36]
shape(s37, 2).text_frame.paragraphs[0].runs[0].text = '결론 — 네 가지 연구 문제의 답'; LOG.append((37, '제목', '결론', '결론 — 네 가지 연구 문제의 답', '본으로 복귀'))
set_one(37, 3, '결론 — 기여 · 한계 · 답한 것', '본으로 복귀')
set_lines(37, 6, ['① 두 표준의 조항은 서로 연계할 수 있는가 — 있다. 관리 주기(ISO 14001)와 평가 절차(ISO 14044)가 서로 맞물리는 자리가 있고, 한쪽은 정량 방법을, 다른 쪽은 되풀이하는 주기를 비워 두었다.',
                  '② 연계 프로세스와 평가 문항은 어떻게 구성되는가 — 맞물리는 자리를 업무로 묶으면 16종, 되어 있는지 묻는 문항이 55개, 문항마다 채점 기준과 증거 목록.',
                  '③ 성숙도는 어떤 규칙으로 판정하는가 — ML0~ML5 여섯 수준, 핵심 문항 4종이 모두 갖춰졌는지가 먼저이고 평균은 쓰지 않는다.',
                  '④ 그 판정이 실제 회사에 무엇을 뜻하는가 — A사는 2.85점이지만 핵심 문항 하나가 미달이라 ML2, 과제 다섯을 마치면 3.15점 ML3. 판정이 개선 계획으로 이어지고, 그 결과가 다음 주기의 입력이 된다.',
                  '두 표준의 조항을 연계하여 LMA를 만들고, 이 모델로 조직을 평가하여, 그 결과를 환경경영시스템(ISO 14001)의 지속적 개선으로 연결함으로써 전과정의 환경영향을 줄이는 데 기여하고자 한다.'], '네 답 쉬운 말(5.1) + 논문 마지막 문장', 15)
shape(s37, 6).height = Emu(int(5.6*IN))
set_note(37, '네 가지 연구 문제의 답입니다.\n첫째, 두 표준의 조항은 연계할 수 있습니다. 관리 주기와 평가 절차가 서로 맞물리는 자리가 있고, 한쪽은 정량 방법을, 다른 쪽은 되풀이하는 주기를 비워 두었기 때문입니다.\n둘째, 그 자리를 업무로 묶으면 연계 프로세스 열여섯 가지이고, 되어 있는지 묻는 문항이 쉰다섯 개입니다.\n셋째, 성숙도는 여섯 수준이고, 핵심 문항 네 가지가 모두 갖춰졌는지로 판정하며 평균은 쓰지 않습니다.\n넷째, A사에 적용해 판정과 개선 계획이 나오는 것을 확인했고, 그 결과가 다음 주기의 입력이 됩니다.\n두 표준의 조항을 연계하여 LMA를 만들고, 이 모델로 조직을 평가하여, 그 결과를 환경경영시스템(ISO 14001)의 지속적 개선으로 연결함으로써 전과정의 환경영향을 줄이는 데 기여하고자 합니다. 이상입니다.\n〔45초. 네 답 원문은 5.1 [그림 5-1](백업 B-0c①), 연구 전체 한 장 요약은 S-28. 마지막 문장은 논문 5.3 끝 문장 그대로.〕', '네 답 쉬운 말')
# 구분·감사 문구
rep(prs, 20, 2, 'S-5~S-29 본심사판에서 옮긴 17장(두 표준의 빈 곳 · 연구 목적과 문제 · ISO 14001 구조 · LCA 4단계 · 따른 표준 · 선행연구 · 연계 지점 맵 · 16종 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 개선 과제 표 · 로드맵·재평가 · 장점과 단점 · 연구 전체 요약 · 결론)',
    'S-5~S-28 본심사판에서 옮긴 15장(두 표준의 빈 곳 · 연구 목적과 문제 · ISO 14001 구조 · LCA 4단계 · 따른 표준 · 선행연구 · 연계 지점 맵 · 16종 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 개선 과제 표 · 로드맵·재평가 · 연구 전체 요약)', '백업 목록')
note_rep(19, 'S-5~S-29(본심사판에서 옮긴 17장)', 'S-5~S-28(본심사판에서 옮긴 15장)', '백업 수')
# ---- 순서: 본 = 1..15, 35, 17, 37, 19 ; 16·18 제거 ; 구분 20 ; S 블록에서 35·37 제외
N = len(prs.slides)
MAIN = list(range(1, 16)) + [35, 17, 37, 19]
DROP = [16, 18]
order = MAIN + [20] + [i for i in range(21, N+1) if i not in (35, 37)]
assert sorted(order + DROP) == list(range(1, N+1))
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for d in DROP: prs.part.drop_rel(ids[d-1].rId)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, n in enumerate(nums): out += str(old2new.get(int(n), int(n))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = s.notes_slide.notes_text_frame; t = tf.text; t2 = tr_nums(t)
    if t2 != t: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new}, open('log29b.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), 'slides', len(prs.slides))
