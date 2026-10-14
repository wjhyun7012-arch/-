# v28 15분판 -> v29: 논문 리뷰 구성(본 19장) — 1-2·2-5·2-7·4-10·5-1·5-2·5-3 그림 장, 매핑↔LMA 순서, 진행 안내 삭제
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from PIL import Image
from pptlib import rep, rep_notes, LOG, shape

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
IN = 914400

def set_note(sn, text, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; old = tf.text; tf.text = text
    LOG.append((sn, '발표 노트(전체)', old, text, why))

def set_lines(sh, lines, why='', sn=None, size=None):
    """도형 글을 줄 단위로 교체 — 첫 문단·첫 run 서식 유지"""
    tf = sh.text_frame; old = tf.text
    p0 = tf.paragraphs[0]; pPr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    r0 = p0.runs[0]; rPr = copy.deepcopy(r0._r.rPr) if r0._r.rPr is not None else None
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    r0.text = lines[0]
    if size: r0.font.size = Pt(size)
    for ln in lines[1:]:
        p = tf.add_paragraph()
        if pPr is not None: p._p.insert(0, copy.deepcopy(pPr))
        r = p.add_run(); r.text = ln
        if rPr is not None: r._r.insert(0, copy.deepcopy(rPr))
        if size: r.font.size = Pt(size)
    LOG.append((sn, f'[{sh.shape_id}]', old, '\n'.join(lines), why))

def set_title(sh, text, ref=None, size=27, why='', sn=None):
    tf = sh.text_frame; old = tf.text; p = tf.paragraphs[0]; runs = p.runs
    runs[0].text = text; runs[0].font.size = Pt(size)
    for r in runs[1:]: r._r.getparent().remove(r._r)
    if ref:
        r = p.add_run(); r.text = '  ' + ref; r.font.size = Pt(13); r.font.bold = False
        r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F); r.font.name = '맑은 고딕'
    LOG.append((sn, '제목', old, text + ('  ' + ref if ref else ''), why))

def set_one(sh, text):
    p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)

def remove_shape(sn, sid, why=''):
    s = prs.slides[sn-1]; sh = shape(s, sid); old = sh.text_frame.text if sh.has_text_frame else '[그림]'
    sh._element.getparent().remove(sh._element); LOG.append((sn, f'[{sid}] 삭제', old, '(삭제)', why))

def fit_box(path, box):
    l, t, r, b = box; W, H = r - l, b - t
    im = Image.open(path); a = im.width / im.height
    if W / H > a: h = H; w = h * a
    else: w = W; h = w / a
    return l + (W - w) / 2, t + (H - h) / 2, w, h

def add_pic(s, path, box):
    x, y, w, h = fit_box(path, box)
    return s.shapes.add_picture(path, Emu(int(x*IN)), Emu(int(y*IN)), Emu(int(w*IN)), Emu(int(h*IN)))

def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1

TPL = 5  # 복사 원본: 제목(2)·장 표시(3)·선(4)·요지(5)·쪽번호(6)·근거(7)
def new_fig_slide(title, ref, tag, tag_w, gist, base, fig, box, band=None, title_size=27):
    tpl = prs.slides[TPL-1]; s = prs.slides.add_slide(tpl.slide_layout)
    for sh in tpl.shapes:
        if sh.shape_id in (2, 3, 4, 5, 6, 7): s.shapes._spTree.append(copy.deepcopy(sh._element))
    renumber_ids(s)
    shs = {sh.shape_id: sh for sh in s.shapes}
    set_title(shs[2], title, ref, title_size)
    t3 = shs[3]; set_one(t3, tag)
    t3.width = Emu(int(tag_w*IN)); t3.left = Emu(int((11.15 - tag_w)*IN))
    set_one(shs[5], '▶ ' + gist)
    set_one(shs[7], '근거  ' + base)
    add_pic(s, fig, box)
    if band:
        bsrc = shape(prs.slides[TPL-1], 9); el = copy.deepcopy(bsrc._element); s.shapes._spTree.append(el); renumber_ids(s)
        bsh = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.startswith('상태 A')][0]
        bsh.top = Emu(int(band[1]*IN)); bsh.height = Emu(int(band[2]*IN)); set_lines(bsh, band[0], size=band[3] if len(band) > 3 else None)
    LOG.append((len(prs.slides), '새 장', '(없음)', f'{title} — {fig.split("/")[-1]}', '논문 그림 그대로'))
    return len(prs.slides)

from PIL import ImageOps
for _f in ('1-2','2-5','2-7','4-10','5-1','5-2','5-3'):
    _im = Image.open(f'figs29/fig{_f}.png').convert('RGB'); _bb = _im.convert('L').point(lambda v: 0 if v > 242 else 255).getbbox()
    if _bb: _im.crop((max(_bb[0]-20,0), max(_bb[1]-20,0), min(_bb[2]+20,_im.width), min(_bb[3]+20,_im.height))).save(f'figs29/fig{_f}.png')
N0 = len(prs.slides)  # 75
# =============== 기존 장 수정 (옛 번호 기준)
# 2 목차
rep(prs, 2, 22, '오늘은 모델의 뼈대와 적용 결과를 중심으로 말씀드리고, 조항별 판정과 문항 전문은 부록에 두었습니다.',
    '논문의 장 순서대로 — 조항별 판정과 문항 전문은 백업 슬라이드에 두었습니다.', '리뷰 구성')
set_note(2, '논문의 장 순서대로, 1장 왜 이 연구인가, 2장 무엇을 바탕으로 하였는가, 3장 무엇을 만들었는가, 4장 적용하면 어떻게 되고 무엇을 고치는가, 5장 결론입니다.\n〔15초. 진행 안내는 말하지 않는다. 세부는 구분 슬라이드 뒤 백업.〕', '진행 안내 삭제')

# 3 연구 절차와 논문 구성 [그림 1-2]
s3 = prs.slides[2]
set_title(shape(s3, 2), '연구 절차와 논문 구성', '[그림 1-2]', 27, '논문 그림으로', 3)
set_one(shape(s3, 5), '근거  1.3.2, 1.5, [그림 1-2]')
remove_shape(3, 6, '[그림 1-4] → 백업')
add_pic(s3, 'figs29/fig1-2.png', (0.54, 1.05, 11.14, 6.6))
bsrc = shape(prs.slides[TPL-1], 9); el = copy.deepcopy(bsrc._element); s3.shapes._spTree.append(el); renumber_ids(s3)
bsh = [x for x in s3.shapes if x.has_text_frame and x.text_frame.text.startswith('상태 A')][0]
bsh.top = Emu(int(6.72*IN)); bsh.height = Emu(int(0.9*IN))
set_lines(bsh, ['1·2단계 → 2장, 3.2.2 (RQ1 두 표준의 조항 매핑)   │   3·4단계 → 3장 (RQ2 연계 프로세스와 평가지표 · RQ3 성숙도 판정 규칙)   │   5·6단계 → 4장·5장 (RQ4 판정 결과의 의미와 개선)',
               '본 논문은 다섯 장 — 1장 서론 · 2장 이론적 배경 · 3장 모델 개발 · 4장 사례 적용 · 5장 결론'], '장 대응 띠', 3, 14)
set_note(3, '본 논문은 다섯 장으로 구성되어 있고, 연구는 여섯 단계로 진행했습니다.\n처음 두 단계는 두 표준의 원문을 분석하고 전 조항을 대조하여 어디가 대응하고 어디가 비어 있는지 확인한 것으로 2장과 3장 앞부분이며, 첫째 연구 문제에 해당됩니다.\n3·4단계는 모델을 만들고 검증한 것으로 3장이며, 둘째·셋째 연구 문제에 해당됩니다.\n5·6단계는 그 모델을 A사에 적용하고 개선 과제로 옮긴 것으로 4장과 5장이며, 넷째 연구 문제에 해당됩니다.\n〔1분 20초. 연구 문제 원문(1.2.2)은 백업 S-7, [그림 1-4] 논문 구성과 장별 요약은 백업 57·58. 여섯 단계는 Becker 8단계를 산출물 단위로 묶은 것(1.3.2·3.1.3).〕', '이사님 10/14 문안')

# 4 연구 배경과 필요성
s4 = prs.slides[3]
set_title(shape(s4, 2), '연구 배경과 필요성', None, 27, '리뷰 구성', 4)
set_one(shape(s4, 5), '▶ 두 표준은 실무에서 별도로 운영되고, 그 연계 수준을 조항 단위로 확인하는 기준이 없었다.')
set_one(shape(s4, 7), '근거  1.1, 1.2.1')
remove_shape(4, 26, '4곳·ML0~ML5는 9·11번에서')
sh25 = shape(s4, 25); sh25.height = Emu(int(1.3*IN))
set_note(4, '환경경영시스템(ISO 14001)과 전과정평가(LCA)는 각각 성숙한 표준이지만, 실무에서는 별도로 운영됩니다. 왼쪽은 회사가 환경을 관리하는 틀이고, 오른쪽은 제품이 각 단계에서 환경에 주는 영향을 숫자로 계산하는 방법입니다.\n두 표준은 별개의 표준이라 한쪽이 다른 쪽을 요구하지 않습니다. 전과정평가(LCA)를 했어도 그 결과가 경영시스템의 결정에 쓰이지 않는 경우가 많고, 둘이 어느 정도 연계되어 있는지 확인할 기준이 없었습니다.\n그래서 연계된 정도를 조항 단위로 확인하고 성숙도 수준으로 판정하는 모델을 만들었습니다.\n〔40초. 배경 — 제품 전과정 환경정보를 요구하는 규제(CBAM·ESG 공시) 1.1. 「ISO 14001은 많이 쓰이나」 → 국내 유효 인증 28,137건, 세계 676,232건(ISO Survey 2024) 2.1.1, 백업 S-5.〕', '배경·필요성으로 축소')

# 7(옛) 조항 매핑 → 새 8
s7 = prs.slides[6]
set_title(shape(s7, 2), '두 표준의 조항 매핑 — 30개 조항과 25개', None, 27, '30·25 설명', 7)
set_one(shape(s7, 5), '▶ 두 표준의 shall·should 문장을 전부 대조하고, 조항마다 같은 네 가지 질문으로 네 가지 등급을 판정하였다.')
set_one(shape(s7, 7), '근거  3.2.2, 3.3.1, [표 3-2], 부록 A.1·A.2·A.4')
rep(prs, 7, 13, '네 가지 판정 질문 — 30개 조항 전부에 같은 순서로', '네 가지 판정 질문 — ISO 14001 요구사항 조항 30개 전부에 같은 순서로', '30 명시') if '네 가지 판정 질문 — 30개 조항 전부에 같은 순서로' in shape(s7, 13).text_frame.text else None
sh27 = shape(s7, 27); sh27.top = Emu(int(6.17*IN)); sh27.height = Emu(int(1.08*IN))
set_lines(sh27, ['ISO 14001 — 요구사항 조항 30개(원표 34행에서 7.4.1~7.4.3은 하나로, 하위 조항만 가진 6.1·6.2 제외) 전부 문항화 → 44문항   ·   ISO 14044 — 전조항 28개 가운데 25개가 문항의 근거, 나머지 3개(4.2.1·4.2.3.1 일반 조항, 4.4.3)는 세부 조항의 문항·체크포인트로 반영 → 빠진 조항 없음(부록 A.4)',
                 '매우 높음 4개 조항 = 연계 지점 ①~④ = 핵심 문항 4종(P-9·P-14·D-1·C-1) — 관리 주기의 계획 둘·실행 하나·점검 하나. 높음·중간·낮음 조항은 빠지는 것이 아니라 지원 문항으로 편입(3.2.2)'], '30·25·핵심 4', 7, 12)
set_note(7, '3장의 출발은 매핑입니다. 두 표준의 의무와 권고 문장을 전부 대조했습니다.\nISO 14001은 요구사항 조항 서른 개가 모두 문항이 되었고, ISO 14044는 스물여덟 개 가운데 스물다섯 개가 문항의 근거가 되었으며 나머지 세 개도 세부 조항의 문항과 체크포인트로 들어가 빠진 조항은 없습니다.\n서른 개 조항 하나하나에 같은 네 가지 질문을 같은 순서로 물어 네 가지 등급으로 판정했습니다. 매우 높음이 네 조항, 환경측면, 환경목표, 운용기획 및 관리, 모니터링입니다.\n이 네 곳이 연계 지점이고 뒤의 핵심 문항 4종이 되며, 관리 주기의 계획에 둘, 실행과 점검에 하나씩 놓입니다.\n〔1분 30초. 30 = 원표 34행에서 7.4.1~7.4.3을 하나로, 상위 조항 6.1·6.2 제외(3.2.2). 25·28 = 부록 A.4. 대조 대상은 shall·should — may·can은 제외(3.1.1, KS A 0001 부속서 I). 55문항 근거 성격 shall 48·비교주장 시 조건부 shall 2·수행 시 조건부 shall 2·should 1·shall+should 2(3.3.1)는 노트북. 연계 지점 맵 [그림 3-4]는 백업 S-14, 등급 요건은 B-1a.〕', '이사님 10/14 문안')

# 6(옛) LMA 구성 → 새 9
s6 = prs.slides[5]
set_one(shape(s6, 5), '▶ 매핑에서 연계 프로세스 16종을 도출하고, 그 성과를 55문항으로 평가하며, 점수를 성숙도 규칙으로 판정한다.')
set_lines(shape(s6, 8), ['환경경영시스템(ISO 14001)', '요구사항 조항 30개'], 'LMA', 6)
set_lines(shape(s6, 9), ['전과정평가(LCA, ISO 14044)', '요구사항 25개'], 'LMA', 6)
set_lines(shape(s6, 11), ['앞 장의 조항 매핑'], 'LMA', 6)
set_lines(shape(s6, 12), ['매핑의 결과 = 연계 지점 · 그 가운데 핵심 네 곳 → 세 모델의 재료'], 'LMA', 6)
set_lines(shape(s6, 15), ['연계 프로세스 16종 — 계획 8 · 실행 2 · 점검 3 · 개선 3'], 'LMA', 6)
set_lines(shape(s6, 20), ['프로세스의 성과 하나에 문항 하나 — 확인 항목 개수로 1~5점'], 'LMA', 6)
set_note(6, '조항 매핑의 결과로 만든 것이 셋입니다.\n첫째 참조 모델은 연계 지점을 업무 단위로 묶은 연계 프로세스 열여섯 가지로, 계획 여덟, 실행 둘, 점검 셋, 개선 셋입니다. 둘째 평가 모델은 프로세스의 성과 하나에 문항 하나씩 쉰다섯 문항이고, 문항마다 채점 기준과 확인할 증거를 정했습니다. 셋째 성숙도 모델은 점수를 여섯 수준으로 바꾸는 규칙이고, 핵심 문항 넷 가운데 하나라도 빠지면 ML2에서 멈춥니다.\n이 셋을 합쳐 연계 성숙도 평가 모델, LMA라고 하며, 만든 뒤 추적성과 표준 적합성 두 방향으로 검증했습니다.\n〔1분 10초. 16종 이름은 백업 S-15([그림 3-6]), 논문 [그림 3-2] 전체는 B-7, 방법론 구성 [그림 3-1]은 B-2. 검증 — 3.5.1 추적성, 3.5.2 ISO/IEC 33004 24건(충족 22·해당 없음 1·보완 1).〕', '이사님 10/14 문안')

# 10 평가 모델 — D-1 예시 상자는 지시문으로
remove_shape(10, 18, 'D-1 예시는 노트 지시문·백업 B-10a')
set_note(10, '평가 모델은 채점표입니다. 채점은 세 단계입니다. 첫째, 문항마다 확인할 항목을 미리 정해 둡니다. 반드시 있어야 할 필수 항목과, 있으면 더 좋은 추가 항목입니다. 이것을 체크포인트라 합니다. 둘째, 회사 문서를 보고 그 항목이 충족되었는지 확인합니다. 말로 설명한 것은 인정하지 않고 문서와 기록만 봅니다. 셋째, 충족한 체크포인트의 개수가 점수입니다. 필수가 전부 있으면 3점, 이것이 기준선입니다. 추가까지 있으면 4점이나 5점, 필수가 빠지면 1점이나 2점입니다.\n그래서 점수는 평가자의 재량이 아니라 체크포인트 충족 개수로 정해집니다.\n〔55초. 예 — D-1 설계·개발 단계의 LCA 결과 반영: A사는 필수 3개 가운데 2개만 있어 2점, 4장에서 판정을 정한 문항. 점수별 기준 전문(BARS 앵커)은 백업 B-10a, 부록 C.〕', '예시·진행 안내 삭제')
# 11
_tf = prs.slides[10].notes_slide.notes_text_frame; _old = _tf.text
assert '이 모델을 한 회사에 적용한 결과를 보겠습니다.' in _old
_tf.text = _old.replace('이 모델을 한 회사에 적용한 결과를 보겠습니다.\n', ''); LOG.append((11, '발표 노트', '이 모델을 한 회사에 적용한 결과를 보겠습니다.', '(삭제)', '진행 안내 삭제'))
# 12
set_note(12, '4장은 이 모델을 한 회사에 적용한 결과입니다.\n첫째, 문서를 받았습니다. 제조기업 A사의 환경경영시스템(ISO 14001)과 2025년 제품 전과정평가(LCA)에 관한 문서 열여섯 종, 서른세 건입니다.\n둘째, 쉰다섯 문항을 채점하여 잠정 종합점수 2.89를 얻었습니다. 셋째, 잠정 점수를 증빙 문서로 하나씩 다시 확인하여 확정 종합점수는 2.85입니다. 올라간 것 일곱, 내려간 것 아홉입니다.\n넷째, 판정입니다. 핵심 문항 네 가지 가운데 D-1이 2점이어서 ML2, 관리화입니다. 다섯째, 낮은 점수의 원인 세 가지에서 개선 과제 다섯 건을 도출하였고, 모두 마치면 3.15, ML3이 전망됩니다.\n평가는 문서 심사만으로, 평가자 한 명이, 2026년 6월과 7월에 했습니다. 공식 등급을 주는 평가가 아니라 모델의 판정 논리가 실제 자료에서 작동하는지 확인하는 시범 적용입니다.\n〔1분. 객관성 두 가지(판정 규칙 사전 확정 · 잠정 점수의 증빙 재확인) 전문은 백업 S-21. ISO/IEC 33002 평가 클래스 3 상당(4.2.3). 「회사가 쓰겠다던가」 → 백업 B-Q9.〕', '진행 안내 삭제, 10초 단축')
# 13
set_note(13, '핵심 문항 4종 중 P-9는 4점, P-14와 C-1은 3점으로 통과하였으나, D-1이 2점으로 미달하였습니다.\n판정 경로는 아래 조건부터 차례로 묻습니다. ML1 진입 조건, 핵심 문항 4종이 모두 2점 이상이어서 통과입니다. ML2 진입 조건, 나머지 쉰한 문항 가운데 절반 이상이 3점이어서 통과입니다. 서른여섯 문항입니다. ML3 진입 조건, 핵심 문항 4종이 모두 3점 이상이어야 하는데 D-1이 2점이라 미충족입니다. 그래서 판정은 ML2, 관리화입니다.\n종합점수 2.85는 판정에 쓰지 않는 보조값입니다. 과반 기준이나 핵심 문항 기준을 올려도 이 판정은 변하지 않습니다. 판정은 점수의 크기가 아니라 핵심 문항의 충족 여부로 정해집니다.\n〔1분 10초. 영역별 평점 — 계획 3.23, 실행 3.00, 점검 2.10, 개선 2.33: LCA가 계획·실행에는 반영되나 점검·개선으로는 이어지지 않음. 전문은 백업 S-22, 그래프는 B-13a. 「너무 가혹하지 않은가」 → 백업 B-Q1.〕', '영역 평점을 지시문으로, 진행 안내 삭제')
# 13 제목이 장 표시와 겹침 → 20pt
_t13 = shape(prs.slides[12], 2)
for _r in _t13.text_frame.paragraphs[0].runs[:1]: _r.font.size = Pt(20)
LOG.append((13, '제목 크기', '23pt', '20pt', '장 표시와 겹침'))
# 18 감사(Q&A 지시문)
rep_notes(prs, 18, 'S-5~S-28(본심사판에서 옮긴 12장)', 'S-5~S-29(본심사판에서 옮긴 17장)', '백업 수')
# 19 구분
rep(prs, 19, 2, 'S-5~S-28 본심사판에서 옮긴 12장(두 표준의 빈 곳 · 연구 목적과 문제 · ISO 14001 구조 · LCA 4단계 · 따른 표준 · 선행연구 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 로드맵 · 연구 전체 요약)',
    'S-5~S-29 본심사판에서 옮긴 17장(두 표준의 빈 곳 · 연구 목적과 문제 · ISO 14001 구조 · LCA 4단계 · 따른 표준 · 선행연구 · 연계 지점 맵 · 16종 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 개선 과제 표 · 로드맵·재평가 · 장점과 단점 · 연구 전체 요약 · 결론)', '백업 목록')
# 본 → 백업으로 가는 다섯 장: 장 표시
for o, n in ((8, 14), (9, 15), (15, 25), (16, 27), (17, 29)):
    s = prs.slides[o-1]; sh = shape(s, 3); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]
    p.runs[0].text = f'백업  S-{n} (본심사판 {n}번)'
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    sh.width = Emu(int(3.3*IN)); sh.left = Emu(int((11.15-3.3)*IN))
    LOG.append((o, '장 표시', old, f'백업  S-{n} (본심사판 {n}번)', '본 → 백업'))

# =============== 새 장 (끝에 추가 → 뒤에서 재배열)
BOX = (0.54, 1.35, 11.14, 7.6)
n6 = new_fig_slide('두 표준의 연계 — 관리 주기와 평가 단계의 대응', '[그림 2-5]', 'Ⅱ. 학술적 배경 — 무엇을 바탕으로', 3.3,
    'ISO 14001은 방법을 정하지 않고, ISO 14044는 경영시스템에 넣는 일을 정하지 않는다 — 그 사이가 연계다.', '2.1~2.3, [표 2-3], [그림 2-5]', 'figs29/fig2-5.png', BOX, title_size=22)
set_note(n6, '2장은 두 표준과 그 사이입니다. ISO 14001은 전과정 관점을 고려하라고 요구하지만 그 방법은 정하지 않았고, ISO 14044는 평가의 방법과 절차를 요구사항으로 정했지만 그 결과를 경영시스템에 넣는 일은 정하지 않았습니다.\n그림의 아래 띠는 두 표준의 네 가지 대응이고, 위 계단은 성숙도 수준 ML0부터 ML5입니다. 계획 단계에서 적용범위와 환경측면·환경목표는 LCA의 목적·범위 정의와 영향평가 결과를 입력으로 삼고, 실행 단계의 운용기획 및 관리는 목록분석의 데이터 수집과 연결되며, 점검 단계의 모니터링과 경영검토는 영향평가·해석의 결과와 재수행 검증을 입력으로 삼고, 개선 단계는 해석의 권고를 차기 환경목표로 피드백합니다.\n점선은 관리 주기의 후반부일수록 요구되는 수준이 높아짐을 보입니다.\n〔50초. 단독·연계 운영 비교 [표 2-3]은 백업 S-5, ISO 14001 구조는 S-8, LCA 4단계는 S-9. 문장은 2.3.2·2.3.3(B432·B435) 그대로.〕', '이사님 10/14 문안')
n7 = new_fig_slide('성숙도 평가의 방법과 선행연구의 빈 자리', '[그림 2-7]', 'Ⅱ. 학술적 배경 — 무엇을 바탕으로', 3.3,
    '판정 체계는 선행 표준에서 가져왔고, 두 표준의 연계를 조항 단위로 판정한 연구는 찾아보기 어렵다.', '2.4, 2.5, [표 2-4], [표 2-5], [그림 2-7]', 'figs29/fig2-7.png', BOX)
set_note(n7, '성숙도 평가의 방법은 새로 만들지 않고 선행 표준에서 가져왔습니다. 왼쪽 방법론 축에서 세 계층 구조는 ISO/IEC 33004, 최소충족과 비보상 판정 구조는 KS X ISO 8000-62, 척도는 행동기술척도에서 가져왔고, 그 적용 선례가 지도교수 연구진의 공공데이터 품질 성숙도 연구입니다.\n오른쪽 분야 축의 환경·LCA 성숙도 연구들은 환경경영 전반이나 제품 설계, LCA 활용 역량 가운데 한 측면씩만 다룹니다.\n두 표준의 연계 자체를 조항 단위로 판정하는 모델은 찾아보기 어려웠고, 그 빈 자리가 이 연구의 위치입니다.\n〔50초. 선행 모델 대비표 [표 2-4]·[표 2-5]는 백업 S-11, 따른 표준은 S-10·B-2. 「8000-62를 왜 가져왔나」 → B-Q7, 「분야가 무엇인가」 → 전문 49번.〕', '이사님 10/14 문안')
n15 = new_fig_slide('개선 로드맵 — R1~R5와 장기 과제 R6', '[그림 4-10]', 'Ⅴ. 개선 과제 — 무엇을 고치나', 3.15,
    '낮은 점수의 원인 세 가지에서 개선 과제 다섯 건 — 다 마치면 종합 3.15점, ML3이 전망된다.', '4.6, [표 4-15], [표 4-18], [그림 4-10]', 'figs29/fig4-10.png', (0.54, 1.5, 11.14, 7.4))
set_note(n15, '판정은 진단에서 끝나지 않고 개선 과제로 이어집니다. 낮은 점수의 원인 세 가지에서 과제 다섯 건을 도출했고, 추진 시기 순서로 배치한 것이 이 로드맵입니다.\n단기는 문서화 과제 둘, 중기는 절차 신설 둘이고, 그 가운데 R3 설계·개발 절차 제정이 핵심 문항 D-1을 통과시켜 ML2에서 ML3으로 올리는 직접 수단입니다. 중기에서 장기에 걸친 R5는 LCA 연 1회 재수행의 정착입니다.\n다섯을 다 마치면 종합 3.15점, ML3이 전망되고, 외부 정밀검토 R6은 ML4를 향한 장기 과제입니다.\n〔1분. 과제별 대상 문항·점수 효과 [표 4-15]는 백업 S-25, [표 4-18] 로드맵·재평가 전문은 S-26, 과제별 카드는 B-14. 「3.15의 근거」 → 목표 상태가 채점 기준과 같은 문구, [표 4-16]·[표 4-17]. 시기 — 단기 1년 이내, 중기 3년 이내, 장기 3년 초과(4.6).〕', '이사님 10/14 문안')
n16 = new_fig_slide('학문적·실무적 기여와 장단점', '[그림 5-2]', '결론 — 기여 · 한계 · 답한 것', 3.0,
    '학문적 기여 셋, 실무적 쓰임 셋, 그리고 선행 모델과 견준 장단점.', '5.2, [표 5-1], [그림 5-2]', 'figs29/fig5-2.png', BOX)
set_note(n16, '학문적 기여는 셋입니다. 측정 대상이 새롭고, 방법은 표준을 그대로 따랐으며, 선행연구의 핵심 항목 선별 원리를 판정 규칙의 구성 요소로 확장했습니다.\n실무적으로는 운영 조직, 평가 수행 인력, 조직 내부 전개의 세 쓰임이 있고, 연계의 전면 도입을 처방하는 것이 아니라 자신의 수준을 보고 선택하게 하는 판단 정보입니다.\n장단점 표의 단점은 전문가 조사와 다사례 적용으로 보완할 것입니다.\n〔50초. [표 5-1] 장점 6·단점 6 전문은 백업 S-27·B-15a·B-15b. 「환경성과와의 관계」 → 전문 46번. 기대 효과(탄소 저감 등)는 모델 구조에서 예상되는 것이지 본 연구에서 확인한 것은 아님(5.2).〕', '이사님 10/14 문안')
n17 = new_fig_slide('한계와 향후 연구 과제', '[그림 5-3]', '결론 — 기여 · 한계 · 답한 것', 3.0,
    '한계 여덟 항목을 성격별로 묶고, 각각의 대응 경로를 정하였다.', '5.3, [그림 5-3]', 'figs29/fig5-3.png', BOX)
set_note(n17, '한계는 여덟 항목입니다. 단일 사례·단일 시점·단일 평가자와 상위 수준 미검증은 자료 조건의 경계이고, 전문가 조사와 다사례 적용으로 풉니다.\n문항 간 종속과 타당화 확장은 심화 과제, 표준 판의 범위와 진단·개선 고도화는 유지·고도화 과제이며, 표준화 경로의 검토는 학회·협회를 거친 뒤의 일입니다.\n〔40초. ML4·ML5는 A사에 재수행 기록이 없어 사례로 미확인, 규칙은 3.4.3. CVI는 조사지 설계까지(3.5.3, 부록 F) → 전문 20번. ISO 14001 개정 → 전문 40번.〕', '이사님 10/14 문안')
n18 = new_fig_slide('결론 — 네 가지 연구 문제의 답', '[그림 5-1]', '결론 — 기여 · 한계 · 답한 것', 3.0,
    '두 표준의 조항을 연계하는 LMA를 만들고 한 기업에 적용하여 작동함을 확인하였다.', '5.1, [그림 5-1]', 'figs29/fig5-1.png', BOX)
set_note(n18, '네 가지 연구 문제의 답입니다.\n첫째, 두 표준의 조항은 관리 주기와 평가 절차의 상호 보완 구조로 연계할 수 있습니다. 둘째, 연계 지점을 업무 단위로 묶은 열여섯 프로세스와 쉰다섯 문항으로 구체화할 수 있습니다. 셋째, 성숙도 수준은 핵심 문항 4종으로 조정한 비보상 판정 규칙으로 정합니다. 넷째, 한 기업의 판정과 개선 계획으로 작동함을 확인했고, 그 결과가 다음 주기의 입력이 됩니다.\n두 표준의 조항 단위로 연계를 확인하고 성숙도로 판정하는 모델을 만들어, 한 기업에 적용하여 작동함을 확인한 연구입니다. 이상입니다.\n〔40초. 결론 전문은 백업 S-29, 연구 전체 한 장 요약은 S-28.〕', '이사님 10/14 문안')

# =============== 순서 재배열
N = len(prs.slides); assert N == N0 + 6
MAIN = [1, 2, 3, 4, 5, n6, n7, 7, 6, 10, 11, 12, 13, 14, n15, n16, n17, n18, 18]
SEP = 19
SBLK = [20, 21, 22, 23, 24, 25, 8, 9, 26, 27, 28, 29, 15, 30, 16, 31, 17]   # S-5,7,8,9,10,11,14,15,18,19,21,22,25,26,27,28,29
order = MAIN + [SEP] + SBLK + list(range(32, N0+1))
assert sorted(order) == list(range(1, N+1)), (len(order), N)
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
# 쪽번호 + 노트의 「백업 N」 번호
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, n in enumerate(nums):
            out += str(old2new.get(int(n), int(n))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = s.notes_slide.notes_text_frame; t = tf.text; t2 = tr_nums(t)
    if t2 != t: tf.text = t2
prs.save(dst)
json.dump({'log': LOG, 'old2new': old2new, 'main': MAIN}, open('log29.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), 'slides', len(prs.slides), 'new main:', [old2new[o] for o in MAIN])
