"""v16b(13·23번 사본을 33·35번에 넣은 74장) → v17 (이사님 10/7).
10 따른 표준: 조항 번호 → 역할 한 줄 / 13 매핑: 네 단 막대 + Q1~Q4 짧게(표는 백업 B-1a=33번) / 19 제목 「반복 개발」 /
21 표 0.15in 위로 / 23 종합 판정: 판정 경로 도형 크게, 막대그래프는 백업 B-13a(=35번, [그림 4-3]) / 31번 백업 문구 / 쪽번호.
사용: python3 -I mk_deck17.py v16b.pptx v17.pptx
"""
import sys, os, json, re, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
from pptlib import _rep_tf
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); W, H = prs.slide_width, prs.slide_height
IN = 914400; L = 497190; RW = W - 2 * L; LOG = []
NAVY = RGBColor(0x1F, 0x38, 0x64); LIGHT = RGBColor(0xE8, 0xEE, 0xF7); RED = RGBColor(0xC0, 0x00, 0x00); PINK = RGBColor(0xFB, 0xEC, 0xEC)
INK = RGBColor(0x11, 0x11, 0x11); GREY = RGBColor(0x59, 0x59, 0x59); WHITE = RGBColor(0xFF, 0xFF, 0xFF); MID = RGBColor(0x5B, 0x7F, 0xB5); LGREY = RGBColor(0xF2, 0xF2, 0xF2)
def set_ea(run):
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn('a:ea')) is None:
        ea = etree.SubElement(rpr, qn('a:ea')); ea.set('typeface', '맑은 고딕')
def par(tf, text, size, bold=False, color=INK, align=None, first=False, after=None, before=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None: p.alignment = align
    if after is not None: p.space_after = Pt(after)
    if before is not None: p.space_before = Pt(before)
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = '맑은 고딕'; r.font.color.rgb = color; set_ea(r); return p
def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name='', anchor=MSO_ANCHOR.MIDDLE, margin=90000):
    sh = s.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h))); sh.name = name
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1.0)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(margin); tf.margin_top = tf.margin_bottom = Emu(45000); return sh
def tbox(s, x, y, w, h, name=''):
    tb = s.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h))); tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = Emu(45720); tf.margin_top = tf.margin_bottom = Emu(20000); return tb
def arrow_right(s, x, cy, w=260000, name='arr'):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(x)), Emu(int(cy - 100000)), Emu(int(w)), Emu(200000)); a.name = name
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0x8A, 0x8A, 0x8A); a.line.fill.background(); a.shadow.inherit = False; return a
def shp(s, sid):
    for sh in s.shapes:
        if sh.shape_id == sid: return sh
    raise KeyError(sid)
def remove(s, *sids):
    for sid in sids:
        sh = shp(s, sid); sh._element.getparent().remove(sh._element)
def set_text_keep(sh, text):
    tf = sh.text_frame; p = tf.paragraphs[0]; rs = p.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
def set_lines(sh, lines, size, head_bold=True):
    tf = sh.text_frame; old = tf.text
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    p0.runs[0].text = lines[0]; p0.runs[0].font.size = Pt(size); p0.runs[0].font.bold = head_bold
    for ln in lines[1:]: par(tf, ln, size, False, INK, PP_ALIGN.LEFT, before=2)
    return old
def rep(sn, sh, old, new, why):
    n = _rep_tf(sh.text_frame, old, new); assert n == 1, (sn, old, n); LOG.append((sn, sh.name, old, new, why))
def set_notes(sn, txt):
    tf = prs.slides[sn - 1].notes_slide.notes_text_frame
    for p in list(tf.paragraphs): p._p.getparent().remove(p._p)
    for line in txt.split('\n'):
        p = tf.add_paragraph(); r = p.add_run(); r.text = line; r.font.size = Pt(12); r.font.name = '맑은 고딕'; set_ea(r)
def label_backup(s, text):
    lb = shp(s, 3); set_text_keep(lb, text); lb.fill.background()
    for p in lb.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(12); r.font.bold = False; r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

# ===== 10번 따른 표준 =====
s = prs.slides[9]
T = {10: ['업무를 어떤 형식으로 적는가', 'ISO/IEC 33004가 요건을, KS X ISO 8000-61이 「제목·목적·성과·활동」 형식을 줌 — 원형은 국내 공공데이터 연구(김선호·이창수, 2013)'],
     13: ['점수를 어떻게 매기는가', 'ISO/IEC 33020이 등급의 틀을, ISO/IEC 33003이 척도가 갖출 요건을 줌 — 본 모델은 그 틀을 5점 상태 기술(BARS)로 구체화'],
     16: ['점수를 수준으로 어떻게 바꾸는가', 'KS X ISO 8000-62가 여섯 수준과 최소충족 원리를, ISO/IEC 33004가 수준의 요건을 줌 — 핵심 문항 4종에만 비보상 판정'],
     17: ['평가를 어떻게 수행하는가 — ISO/IEC 33002(평가 클래스)  ·  심사·평가자 자격·개선 계획의 운영 준거도 표준에서 — 조항 번호와 전체 목록은 백업 B-2']}
for sid, lines in T.items():
    sh = shp(s, sid); old = set_lines(sh, lines, 13 if sid != 17 else 12, head_bold=(sid != 17))
    if sid != 17:
        sh.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    LOG.append((10, f'[{sid}]', old, '\n'.join(lines), '이사님 10/7: 조항 번호 → 역할 한 줄(번호는 백업 B-2)'))
set_notes(10, """모델의 뼈대는 모두 표준에서 가져왔습니다.
측정 대상은 환경경영시스템(ISO 14001)과 전과정평가(LCA)의 연계입니다.
세 가지를 각각 어느 표준에서 가져왔는지 말씀드리면, 업무를 적는 형식은 ISO/IEC 33004와 KS X ISO 8000-61에서, 점수를 매기는 틀은 ISO/IEC 33020과 33003에서, 점수를 수준으로 바꾸는 원리는 KS X ISO 8000-62와 ISO/IEC 33004에서 가져왔습니다. 평가의 수행은 ISO/IEC 33002를 따랐습니다.
판정 구조는 KS X ISO 8000-62를, 모델 형식과 평가 방법은 ISO/IEC 33000 계열을 따랐고, 그 적용 선례는 국내 공공데이터 연구입니다.
그러면 이 대상을 다룬 연구는 이미 있었는가. 다음 장입니다.
〔조항 번호는 읽지 않는다 — 백업 B-2 [그림 1-3]·[표 1-1]. 선행 모델은 논문 제목으로 부른다. 1분〕""")

# ===== 13번 매핑 =====
s = prs.slides[12]; remove(s, 9, 10, 11)
q = shp(s, 8); set_lines(q, ['네 가지 판정 질문 — 30개 조항 전부에 같은 순서로', 'Q1 전과정 관점을 요구하거나 그 결과를 직접 쓰는가', 'Q2 LCA 없이는 의무(shall)를 다 지킬 수 없는가', 'Q3 LCA 결과가 주요 입력인가', 'Q4 보조 정보라도 주는가'], 13.5)
q.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY; q.width = Emu(int(4.6 * IN)); q.height = Emu(int(2.0 * IN))
t = tbox(s, L, q.top + q.height + 150000, int(4.6 * IN), int(1.4 * IN), 'q note'); par(t.text_frame, '「예」가 나오는 물음이 앞쪽일수록 연계 가능성이 높다 — 네 등급으로 매김. 조항별 판정 전문은 백업 B-1(흐름)·B-1a(등급표)·부록 A.', 12, False, GREY, PP_ALIGN.LEFT, first=True)
rx = L + int(4.9 * IN); rw = L + RW - rx; y = int(1.29 * IN)
bars = [('매우 높음', 4, '6.1.2 · 6.2.1 · 8.1 · 9.1.1 — 뒤의 핵심 문항 4종의 자리', True),
        ('높음', 12, 'LCA 결과가 주요 입력 — 없이도 되지만 품질이 떨어짐', False),
        ('중간', 10, 'LCA는 보조 정보', False),
        ('낮음', 4, '조항 본질이 LCA 영역 밖', False)]
bh = int(1.05 * IN); gap = 110000; maxw = rw - int(1.55 * IN)
for i, (nm, n, desc, hot) in enumerate(bars):
    lab = box(s, rx, y, int(1.45 * IN), bh, RED if hot else NAVY, name=f'bar lab{i}'); par(lab.text_frame, nm, 14, True, WHITE, PP_ALIGN.CENTER, first=True); par(lab.text_frame, f'{n}개 조항', 12, False, WHITE, PP_ALIGN.CENTER)
    bw = int(maxw * (0.45 + 0.55 * n / 12)); b = box(s, rx + int(1.55 * IN), y, bw, bh, PINK if hot else LIGHT, name=f'bar{i}', anchor=MSO_ANCHOR.MIDDLE)
    par(b.text_frame, desc, 13 if hot else 12.5, hot, RED if hot else INK, PP_ALIGN.LEFT, first=True)
    y += bh + gap
ft = shp(s, 11) if False else None
t2 = tbox(s, L, int(6.35 * IN), RW, int(0.75 * IN), 'map foot'); par(t2.text_frame, 'ISO 14001 30개 조항(7.4.1~7.4.3은 7.4로 통합)을 전부 판정 — 55문항 중 44는 ISO 14001 조항에, 11은 ISO 14044 고유 요건(목적·범위 정의, 목록분석, 보고, 정밀검토)에 대응', 11.5, False, GREY, PP_ALIGN.LEFT, first=True)
LOG.append((13, '화면', 'Q1~Q4 긴 문장 + 네 등급·조항 번호 32개 표', 'Q1~Q4 짧게 + 네 단 막대(매우 높음 4개 조항만 번호, 나머지는 개수) — 등급표는 백업 B-1a', '이사님 10/7: 13이 어려움'))
set_notes(13, """ISO 14001의 모든 조항, 서른 개에 같은 네 가지 질문을 같은 순서로 물었습니다. 전과정 관점을 요구하거나 그 결과를 직접 쓰는가, LCA 없이는 의무를 다 지킬 수 없는가, LCA 결과가 주요 입력인가, 보조 정보라도 주는가입니다.
「예」가 앞쪽 물음에서 나올수록 연계 가능성이 높고, 네 등급으로 매겼습니다.
매우 높음이 네 개 조항입니다. 환경측면 6.1.2, 환경목표 6.2.1, 운용기획 및 관리 8.1, 모니터링 9.1.1입니다. 이 네 곳이 뒤의 핵심 문항 4종의 자리가 됩니다. 높음이 열둘, 중간이 열, 낮음이 넷입니다.
네 지점이 어디에 놓이는지 그림으로 보겠습니다.
〔「전조항을 매핑했는데 핵심은 4개뿐인가」 → 매핑(30조항) → 55문항·16종 → 연계 성립 지점 4곳. 등급별 판정 요건 전문은 백업 B-1a, 흐름은 B-1. 1분 10초〕""")

# ===== 19번 제목 =====
s = prs.slides[18]; rep(19, shp(s, 2), '모델은 어떻게 만들어졌나 — 반복 개발의 경위', '모델은 어떻게 만들어졌나 — 반복 개발', '이사님 10/7')

# ===== 21번 표 위로 =====
s = prs.slides[20]; d = int(0.15 * IN)
shp(s, 10).top = Emu(shp(s, 10).top - d); shp(s, 13).top = Emu(shp(s, 13).top - d + 40000)
tb = shp(s, 10); shp(s, 11).top = Emu(tb.top + sum(r.height for r in tb.table.rows) + 30000)
LOG.append((21, '배치', '표 3.45in', '표 3.30in, 설명 줄 표 바로 아래', '이사님 10/7: 한 칸 위로'))

# ===== 23번 종합 판정 =====
s = prs.slides[22]; remove(s, 12, 13, 14)
for sid in (8, 9, 10, 11): shp(s, sid).top = Emu(int(1.35 * IN))
y = int(1.35 * IN) + int(0.99 * IN) + 300000
t = tbox(s, L, y, RW, 330000, 'path head'); par(t.text_frame, '판정 경로 — 아래 조건부터 차례로, 「아니오」가 나오는 곳이 수준', 13.5, True, NAVY, first=True); y += 360000
steps = [('ML1 진입', '핵심 문항 4종 모두 2점 이상', '4 · 3 · 2 · 3 — 통과', False),
         ('ML2 진입', '나머지 51문항 중 절반 이상 3점', '36/51 = 70.6% — 통과', False),
         ('ML3 진입', '핵심 문항 4종 모두 3점 이상', 'D-1 = 2점 — 미충족', True),
         ('판정', '', 'ML2 (관리화)\n종합점수 2.85', True)]
n = 4; ag = 260000; sw = (RW - (n - 1) * ag) // n; sh_h = int(1.55 * IN)
for i, (h1, h2, h3, hot) in enumerate(steps):
    x = L + i * (sw + ag); b = box(s, x, y, sw, sh_h, PINK if hot else LIGHT, RED if hot and i == 3 else None, name=f'path{i}', anchor=MSO_ANCHOR.MIDDLE, margin=70000); tf = b.text_frame
    par(tf, h1, 15, True, RED if hot else NAVY, PP_ALIGN.CENTER, first=True, after=3)
    if h2: par(tf, h2, 12, False, INK, PP_ALIGN.CENTER, after=3)
    for k, ln in enumerate(h3.split('\n')): par(tf, ln, 14 if k == 0 else 12, k == 0, RED if hot else INK, PP_ALIGN.CENTER)
    if i < n - 1: arrow_right(s, x + sw + 20000, y + sh_h // 2, ag - 40000, name=f'path arr{i}')
y += sh_h + 260000
b = box(s, L, y, RW, int(0.95 * IN), PINK, name='sens'); par(b.text_frame, '종합점수 2.85는 판정에 쓰지 않는 보조값 — 과반 기준을 50%에서 70%로, 핵심 문항 기준을 4점이나 5점으로 올려도 판정은 ML2 그대로 (민감도)', 13.5, True, RED, PP_ALIGN.CENTER, first=True)
y += int(0.95 * IN) + 120000
t = tbox(s, L, y, RW, 330000, 'sens note'); par(t.text_frame, '영역별 평점과 개선 전·후 그래프([그림 4-3])는 백업 B-13a, 판정 경로 원도([그림 4-2])는 논문 4.3.5', 11.5, False, GREY, first=True)
LOG.append((23, '화면', '핵심 문항 상자 4 + 판정 경로 그림(작음) + 막대그래프(작음) + 글 3줄', '핵심 문항 상자 4 + 판정 경로 도형 4칸(크게) + 민감도 한 줄 — 막대그래프는 백업 B-13a', '이사님 10/7: 23 아이디어'))
set_notes(23, """핵심 문항 4종 중 P-9는 4점, P-14와 C-1은 3점으로 shall 충족 기준 이상으로 통과하였으나, D-1이 2점으로 미달하였습니다.
판정 경로는 아래 조건부터 차례로 묻습니다. ML1 진입, 핵심 문항 4종이 모두 2점 이상인가. 통과입니다. ML2 진입, 나머지 쉰한 문항 가운데 절반 이상이 3점인가. 서른여섯, 70.6%로 통과입니다. ML3 진입, 핵심 문항 4종이 모두 3점 이상인가. D-1이 2점이라 미충족입니다. 그래서 판정은 ML2, 관리화입니다.
종합점수 2.85는 판정에 쓰지 않는 보조값입니다. 과반 기준을 50%에서 70%로 올려도, 핵심 문항의 기준을 4점이나 5점으로 올려도 이 판정은 변하지 않습니다.
판정을 가른 것은 점수의 크기가 아니라 필요조건의 충족 여부입니다. 그러면 낮은 점수는 어디에서 왔는가. 다음 장입니다.
〔「너무 가혹하지 않은가」 → 백업 B-Q1. 영역별 그래프는 백업 B-13a. 1분 20초〕""")

# ===== 새 백업 33(B-1a)·35(B-13a) =====
s = prs.slides[32]; remove(s, 5, 8, 10, 11, 12)
set_text_keep(shp(s, 2), '연계 가능성 판정 등급과 요건 — [표 3-2]'); label_backup(s, '백업  B-1a'); set_text_keep(shp(s, 7), '근거  3.2.2, [표 3-2], 부록 A·B.1')
tb = shp(s, 9); tb.left = Emu(L); tb.top = Emu(int(1.3 * IN)); tb.width = Emu(RW)
for c in tb.table.columns: c.width = Emu(int(c.width * RW / (5.61 * IN)))
for row in tb.table.rows:
    for cell in row.cells:
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size: r.font.size = Pt(r.font.size.pt + 1)
set_notes(33, '백업 — 질문 때만. 13번에서 뺀 [표 3-2] 네 등급·판정 요건·해당 조항 전문. 「높음·중간·낮음에 어떤 조항이 있나」가 나오면 이 장.')
LOG.append((33, '새 백업 B-1a', '(13번 사본)', '[표 3-2] 등급·요건·해당 조항 표 전문', '13번에서 옮김'))
s = prs.slides[34]; remove(s, 5, 8, 9, 10, 11, 12, 14)
set_text_keep(shp(s, 2), '영역별 평점과 개선 전·후 — [그림 4-3]'); label_backup(s, '백업  B-13a'); set_text_keep(shp(s, 7), '근거  4.3, 4.6.2, [그림 4-3]')
pic = shp(s, 13); nh = int(5.2 * IN); nw = int(pic.width * nh / pic.height)
if nw > RW: nw = RW; nh = int(pic.height * nw / pic.width)
pic.width = Emu(nw); pic.height = Emu(nh); pic.left = Emu(L + (RW - nw) // 2); pic.top = Emu(int(1.3 * IN))
set_notes(35, '백업 — 질문 때만. 23번에서 뺀 [그림 4-3] 영역별 평점(현행 · R1~R5 완수 시 재산정)과 핵심 문항 판정. 「영역별로 어디가 낮았나」·「개선 뒤 얼마나 오르나」가 나오면 이 장(26번 표와 같은 수치).')
LOG.append((35, '새 백업 B-13a', '(23번 사본)', '[그림 4-3] 영역별 평점 그래프 크게', '23번에서 옮김'))
# 31번 백업 문구
s = prs.slides[30]; rep(31, shp(s, 2), 'B-1~B-16 그림·표(B-10a 채점 기준 표 · B-11a 진입 조건 표 포함)', 'B-1~B-16 그림·표(B-1a 등급표 · B-10a 채점 기준 표 · B-11a 진입 조건 표 · B-13a 영역별 그래프 포함)', '백업 두 장 추가')
# 쪽번호
n_fix = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top > int(7.5 * IN) and sh.left > int(9.5 * IN) and re.fullmatch(r'\d{1,2}', sh.text_frame.text.strip()):
            if sh.text_frame.text.strip() != str(i): set_text_keep(sh, str(i)); n_fix += 1
LOG.append(('전체', '쪽번호', '', f'{n_fix}곳 재부여', '74장 = 본 30 + 구분 1 + 백업 43'))
prs.save(OUT); json.dump(LOG, open('log17.json', 'w'), ensure_ascii=False, indent=1)
print(len(LOG), 'changes;', len(prs.slides), 'slides; fixed', n_fix)
