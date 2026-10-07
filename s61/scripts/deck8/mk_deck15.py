"""v14a(54번 사본을 18번 뒤에 넣은 73장) → v15 (이사님 10/7 결정 8건).
1 16번 ③ (BARS 앵커) / 2 19→20번 정보 상자 네 줄 / 3 20→21번 상자 줄이고 표 올림 / 4 25→26번 그림 줄임 /
5 26→27번 장점·단점 짧게 14pt / 6 27→28번 표 내림 / 7 54번 타임라인 → 본문 19번(백업 원본 삭제, 29→30장) / 8 6번 상태 B 문구.
사용: python3 -I mk_deck15.py v14a.pptx v15.pptx
"""
import sys, os, json, copy, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
from pptx.text.text import _Paragraph
from pptlib import _rep_tf
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); W, H = prs.slide_width, prs.slide_height
IN = 914400; L = 497190; RW = W - 2 * L; LOG = []
NAVY = RGBColor(0x1F, 0x38, 0x64); INK = RGBColor(0x11, 0x11, 0x11); RED = RGBColor(0xC0, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59)
def set_ea(run):
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn('a:ea')) is None:
        ea = etree.SubElement(rpr, qn('a:ea')); ea.set('typeface', '맑은 고딕')
def par(tf, text, size, bold=False, color=INK, align=None, first=False, after=None, before=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None: p.alignment = align
    if after is not None: p.space_after = Pt(after)
    if before is not None: p.space_before = Pt(before)
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = '맑은 고딕'; r.font.color.rgb = color; set_ea(r); return p
def shp(s, sid):
    for sh in s.shapes:
        if sh.shape_id == sid: return sh
    raise KeyError(sid)
def byname(s, name): return [sh for sh in s.shapes if sh.name == name][0]
def set_text_keep(sh, text):
    tf = sh.text_frame; p = tf.paragraphs[0]; rs = p.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
def clear_tf(tf):
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]: r._r.getparent().remove(r._r)
    return p
def rep(sn, sh, old, new, why):
    n = _rep_tf(sh.text_frame, old, new); assert n == 1, (sn, old, n); LOG.append((sn, sh.name, old, new, why))
def rep_notes(sn, old, new, why):
    n = _rep_tf(prs.slides[sn - 1].notes_slide.notes_text_frame, old, new); assert n == 1, (sn, old, n); LOG.append((sn, '노트', old, new, why))
def set_notes(sn, txt):
    tf = prs.slides[sn - 1].notes_slide.notes_text_frame
    for p in list(tf.paragraphs): p._p.getparent().remove(p._p)
    for line in txt.split('\n'):
        p = tf.add_paragraph(); r = p.add_run(); r.text = line; r.font.size = Pt(12); r.font.name = '맑은 고딕'; set_ea(r)
def move(sh, dy): sh.top = Emu(int(sh.top + dy))

# ===== 8) 6번 상태 B =====
s = prs.slides[5]
rep(6, shp(s, 9), '상태 B — 그 결과로 중대한 환경측면을 다시 결정하고 환경목표를 정하며 다음 해에 다시 산정', '상태 B — 그 결과(LCA 영향평가)를 근거로 중대한 환경측면을 결정하고 환경목표를 정하며, 다음 해에 다시 측정해 확인', '이사님 10/7: 「다시 결정」이 오해 소지 — 그림 안 글자([그림 1-1])는 그대로')
rep_notes(6, '오른쪽 상태 B는 그 결과로 중대한 환경측면을 다시 정하고, 환경목표를 세우고, 다음 해에 다시 재서 확인하는 회사입니다.', '오른쪽 상태 B는 LCA 영향평가 결과를 근거로 중대한 환경측면을 결정하고, 환경목표를 세우고, 다음 해에 다시 측정해 확인하는 회사입니다.', '같은 이유 + 「재서」(규칙 W05) 제거')

# ===== 1) 16번 ③ (BARS 앵커) =====
s = prs.slides[15]; rep(16, byname(s, 'sc2'), '③ 센 개수가 점수', '③ 센 개수가 점수 (BARS 앵커)', '이사님 10/7: 16번에 앵커')

# ===== 7) 새 19번 타임라인(본문) =====
s19 = prs.slides[18]; s18 = prs.slides[17]
# 장 표시·▶ 줄을 18번에서 복사
old_label = shp(s19, 3); old_label._element.getparent().remove(old_label._element)
for src_id in (3, 5):
    el = copy.deepcopy(shp(s18, src_id)._element); s19.shapes._spTree.append(el)
lab = [sh for sh in s19.shapes if sh.has_text_frame and sh.text_frame.text.startswith('Ⅲ.')][0]
hd = [sh for sh in s19.shapes if sh.has_text_frame and sh.text_frame.text.startswith('▶')][0]
p = hd.text_frame.paragraphs[0]; p.runs[-1].text = '한 번에 설계하지 않고 전 조항 대조 → 55문항 → 채점 기준 → 표준 대조의 순서로 다듬었다.'
for r in p.runs[1:-1]: r._r.getparent().remove(r._r)
set_text_keep(shp(s19, 2), '모델은 어떻게 만들어졌나 — 반복 개발의 경위')
q = shp(s19, 6); q._element.getparent().remove(q._element)   # 「…를 묻는 경우 — 예상 질문 8·9·43번의 백업」 줄 삭제
# 근거 줄 추가(18번 것 복사)
el = copy.deepcopy(shp(s18, 7)._element); s19.shapes._spTree.append(el)
bs = [sh for sh in s19.shapes if sh.has_text_frame and sh.text_frame.text.startswith('근거')][0]; set_text_keep(bs, '근거  2.4.5, 3.1, 3.2.1, 3.2.3')
# 내용(타임라인) 전체를 ▶ 아래(1.35in)로 올림 — 지금 1.57in부터
dy = int(1.35 * IN) - shp(s19, 8).top
for sh in s19.shapes:
    if sh.shape_id in (2, 4, 5) or sh is lab or sh is hd or sh is bs: continue
    move(sh, dy)
set_notes(19, """모델은 한 번에 설계되지 않았습니다. 일곱 단계로 다듬었습니다.
2025년 9월에 두 표준 연계의 목적과 범위를 정리하였고, 10월에 ISO 14001 전 조항 검토표와 연계표, 성숙도표를 만들었습니다.
2026년 3월에 전조항 분석으로 평가 문항 쉰다섯 개와 판정 포인트를 정했고, 6월에 문항별 체크포인트와 BARS 앵커 체계, 핵심 문항 4종을 정하고 잠정 채점을 하였습니다. 7월에 체크포인트 단위로 전수 재검증하여 종합 2.85, ML2로 확정하였습니다.
8월에 ISO/IEC 33000 계열과 KS X ISO 8000-61, 8000-62에 대조하여 판정 규칙과 연계 프로세스 16종 편성을 보완하였고, 9월에 지도교수 검토와 용어·구조 정리, 33004 자체점검 24건을 마쳤습니다.
표준 대조가 모델을 실제로 보완하였습니다. 최소충족 판정은 KS X ISO 8000-62 4.5.1에서, 수준별 프로세스 집합의 도출 규칙은 ISO/IEC 33004 7.3.6에서, 프로세스 기술 형식은 KS X ISO 8000-61 5.3에서 가져왔습니다.
이제 이 도구로 한 기업을 평가한 결과입니다.
〔「처음 설계와 무엇이 달라졌나」 → 2025년 10월 성숙도표의 수준 이름이 현행 문항 점수 이름(1~5점)의 유래, 성숙도 수준 이름과는 별개(예상 질문 17번). 55문항을 16종으로 묶은 기준·명칭·경계는 연구자가 확정하고 지도교수 검토(3.2.1 K1~K4). 1분〕""")
LOG.append((19, '새 본 슬라이드', '백업 B-17 개발 경위 타임라인', '본문 19번 「모델은 어떻게 만들어졌나 — 반복 개발의 경위」(장 표시 Ⅲ, ▶ 줄, 근거 줄 추가, 예상 질문 안내 줄 삭제). 백업 원본은 삭제', '이사님 10/7: 7번 권장안'))

# ===== 2) 20번(옛 19) 정보 상자 네 줄 =====
s = prs.slides[19]; info = byname(s, 'ev info'); key = byname(s, 'ev key')
tf = info.text_frame; p0 = clear_tf(tf); p0.runs[0].text = '적용 개요'; p0.runs[0].font.bold = True; p0.runs[0].font.color.rgb = NAVY; p0.runs[0].font.size = Pt(12.5); p0.alignment = PP_ALIGN.LEFT; p0.space_after = Pt(3)
for ln in ['대상 — 제조기업 A사의 환경경영시스템(ISO 14001)과 2025년 제품 전과정평가(LCA)', '방법 — 문서 심사(면담·현장 관찰 없음), 평가자 1명', '기간 — 2026년 6~7월', '비고 — 연구 목적의 시범 적용(ISO/IEC 33002 평가 클래스 3 상당 · 독립성 범주 B 상당), 회사에 공식 성숙도 등급을 주는 평가가 아님']:
    par(tf, ln, 12.5, False, INK, PP_ALIGN.LEFT, after=3)
info.height = Emu(int(1.95 * IN)); key.top = Emu(info.top + info.height + 160000)
LOG.append((20, 'ev info', '대상 / 방법·기간 한 줄 / 지위', '「적용 개요」 머리 + 대상·방법·기간·비고 네 줄', '이사님 10/7'))

# ===== 3) 21번(옛 20) 상자 줄이고 표 올림 =====
s = prs.slides[20]
for sid in (8, 9):
    b = shp(s, sid); b.height = Emu(int(1.5 * IN))
    for i, p in enumerate(b.text_frame.paragraphs):
        for r in p.runs: r.font.size = Pt(12.5 if i == 0 else 11.5)
note = byname(s, 'obj note'); note.top = Emu(int(1.25 * IN) + int(1.5 * IN) + 50000)
tb = shp(s, 10); tb.top = Emu(int(3.45 * IN)); cap = shp(s, 11); cap.top = Emu(tb.top + sum(r.height for r in tb.table.rows) + 40000)
LOG.append((21, '배치', '상자 1.95in · 표 3.92in · 설명 7.91in(쪽번호와 겹침)', '상자 1.5in · 표 3.45in · 설명 표 바로 아래', '이사님 10/7'))

# ===== 4) 26번(옛 25) 그림 줄이고 표·설명 올림 =====
s = prs.slides[25]; pic = shp(s, 8); nh = int(4.0 * IN); nw = int(pic.width * nh / pic.height)
pic.left = Emu(int(pic.left + (pic.width - nw) / 2)); pic.width = Emu(nw); pic.height = Emu(nh); pic.top = Emu(int(1.3 * IN))
tb = shp(s, 9); tb.top = Emu(int(5.4 * IN)); tx = shp(s, 10); tx.top = Emu(tb.top + sum(r.height for r in tb.table.rows) + 40000); tx.height = Emu(int(0.6 * IN))
for p in tx.text_frame.paragraphs:
    for r in p.runs: r.font.size = Pt(11)
LOG.append((26, '배치', '그림 4.59in · 표 6.02in · 설명 7.55in(쪽번호와 겹침)', '그림 4.0in · 표 5.4in · 설명 표 바로 아래 11pt', '이사님 10/7'))

# ===== 5) 27번(옛 26) 장점·단점 짧게 =====
s = prs.slides[26]
texts = {8: ['장점', '· 문항의 추적성 — 문항마다 근거 조항을 따라갈 수 있다', '· 판정 기준의 명문화 — 같은 증거면 같은 점수', '· 종합점수에 묻히는 미비의 식별 — 필수 네 곳이 빠지면 드러난다', '· 평가와 개선의 연결 — 고칠 목표가 다음 평가 기준과 같은 말', '· 표준 구조의 준용 — 판정 구조·평가 방법을 국제표준에서', '· 수준이 뜻하는 바의 구분 — ML3까지 체계, ML4부터 숫자로 확인'],
         9: ['단점', '· 핵심 문항 4종에 대한 의존 — 네 문항의 선택은 전문가 검토 전', '· 판정 파라미터 — 과반 50% 등은 연구자가 정한 값', '· 평가의 부담 — 55문항 문서 확인, 두 분야를 다 알아야', '· 적용 대상의 제약 — LCA를 안 하는 회사는 차이가 안 난다', '· 문항 간 종속 · 문서 심사 방식 — 미비 하나가 여러 점수에, 형식적 이행은 못 가림', '· 환경성과와의 관계 미검증 — 탄소·비용이 주는지는 확인 안 함']}
for sid, lines in texts.items():
    b = shp(s, sid); old = b.text_frame.text; p0 = clear_tf(b.text_frame); p0.runs[0].text = lines[0]; p0.runs[0].font.size = Pt(15); p0.runs[0].font.bold = True; p0.space_after = Pt(4)
    for ln in lines[1:]: par(b.text_frame, ln, 14, False, INK, PP_ALIGN.LEFT, after=3)
    LOG.append((27, f'[{sid}]', old, '\n'.join(lines), '이사님 10/7: 풀이 짧게, 14pt (항목 이름은 [표 5-1] 그대로)'))

# ===== 6) 28번(옛 27) 표 내림 =====
s = prs.slides[27]
for sid in range(8, 19): move(shp(s, sid), int(0.2 * IN))
LOG.append((28, '배치', '표 1.29in', '표 1.49in, 「한 줄」 띠 6.29in', '이사님 10/7: 표가 ▶ 줄에 붙음'))

# ===== 7) 백업 원본(옛 54 → 지금 55) 삭제, 31번 문구 =====
lst = prs.slides._sldIdLst; ids = list(lst); sid = ids[54]; prs.part.drop_rel(sid.get(qn('r:id'))); lst.remove(sid)
s = prs.slides[30]; rep(31, shp(s, 2), ' · B-17 개발 경위(타임라인)', '', '개발 경위는 본문 19번으로')

# ===== 쪽번호 재부여 =====
n_fix = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top > int(7.5 * IN) and sh.left > int(9.5 * IN) and re.fullmatch(r'\d{1,2}', sh.text_frame.text.strip()):
            if sh.text_frame.text.strip() != str(i): set_text_keep(sh, str(i)); n_fix += 1
LOG.append(('전체', '쪽번호', '', f'{n_fix}곳 재부여', '본 슬라이드 30장(19번 추가), 백업 31~72'))
prs.save(OUT); json.dump(LOG, open('log15.json', 'w'), ensure_ascii=False, indent=1)
print(len(LOG), 'changes;', len(prs.slides), 'slides; fixed', n_fix)
