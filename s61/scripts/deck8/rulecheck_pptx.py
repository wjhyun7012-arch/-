"""규칙표 rules_vNN.json을 발표자료 pptx(화면 글·표 셀·노트)와 추가 문안에 적용. 사용: python3 -I rulecheck_pptx.py deck.pptx rules.json [extra.txt]"""
import sys, re, json, html
from pptx import Presentation
rules = json.load(open(sys.argv[2], encoding='utf8'))
prs = Presentation(sys.argv[1])
items = []  # (where, text)
for i, s in enumerate(prs.slides, 1):
    kind = '본' if i <= 29 else '백업'
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            items.append((f'{kind}{i} 화면[{sh.name}]', sh.text_frame.text))
        if sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    if cell.text_frame.text.strip(): items.append((f'{kind}{i} 표{r+1}행{c+1}열', cell.text_frame.text))
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        items.append((f'{kind}{i} 노트', s.notes_slide.notes_text_frame.text))
if len(sys.argv) > 3:
    for blk in open(sys.argv[3], encoding='utf8').read().split('\n=== '):
        if not blk.strip(): continue
        head, _, body = blk.partition('\n'); items.append(('추가:' + head.strip('= '), body))
def hits(rule, text):
    out = []
    for m in re.finditer(rule['pattern'], text):
        ctx = text[max(0, m.start() - 20):m.end() + 20].replace('\n', ' ')
        if any(a in text[max(0, m.start() - 30):m.end() + 30] for a in rule.get('allow', [])): continue
        out.append(ctx)
    return out
tot = 0
for kind, lst in (('위반', rules['forbidden']), ('경고', rules['warnings'])):
    for r in lst:
        found = []
        for where, text in items:
            for h in hits(r, text): found.append((where, h))
        if found:
            tot += len(found)
            print(f"\n[{r['id']}] {kind} {len(found)}건 — {r['rule']}")
            for where, h in found[:40]: print(f"   {where}: …{h}…")
print('\n총', tot, '건')
