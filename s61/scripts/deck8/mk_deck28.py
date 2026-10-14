# v27(15분판) / v26(본심사판) -> v28: 3번 노트 첫 문장, 제목 아홉 장 논문 캡션, 8번 제목, 표지 제목 「에 대한 연구」 + 영문 제목
import sys, json, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptlib import rep, rep_notes, LOG
src, dst, kind = sys.argv[1], sys.argv[2], sys.argv[3]   # kind: 15 | 30
prs = Presentation(src)
# 15분판 번호 → 본심사판 번호
M15 = {3:3, 5:6, 6:12, 8:14, 9:15, 10:16, 11:17, 13:23, 14:24, 15:25, 16:27}
S = (lambda n: n) if kind == '15' else (lambda n: M15[n])
# 1) 3번 노트 첫 문장
rep_notes(prs, S(3), '논문은 다섯 장이고, 이 그림이 장 사이의 연결입니다.', '이 논문은 다섯 장으로 구성되어 있고, 이 그림은 장 사이의 연결을 보여 줍니다.', '이사님 10/14 문안')
# 2)(3) 제목 — 논문 캡션 + [그림·표] 작은 회색
TITLES = {5: ('본 연구가 말하는 연계 — 두 상태', '[그림 1-1]'),
          6: ('연계 성숙도 평가 모델(LMA)의 구성', '[그림 3-2]'),
          8: ('연계 지점 맵 — 관리 조항과 평가 단계가 만나는 곳', '[그림 3-4]'),
          10: ('평가 모델의 구성 — 55문항과 채점 기준', '[그림 3-8]'),
          11: ('성숙도 모델의 구성 — 판정 규칙과 핵심 문항', '[그림 3-10]'),
          13: ('A사의 판정 경로 — 종합 판정 ML2 (관리화)', '[그림 4-2]'),
          14: ('A사 연계 상태 — 강점과 구조적 미비', '[그림 4-4]'),
          15: ('개선 과제 R1~R5와 장기 과제 R6', '[표 4-15]'),
          16: ('본 모델의 장점과 단점, 그리고 한계', '[표 5-1]')}
for n15, (ttl, ref) in TITLES.items():
    sn = S(n15); s = prs.slides[sn-1]; sh = [x for x in s.shapes if x.shape_id == 2][0]
    old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; r0 = p.runs[0]
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    r0.text = ttl
    if len(ttl) >= 22: r0.font.size = Pt(23)
    r2 = copy.deepcopy(r0._r); p._p.append(r2); rr = p.runs[-1]; rr.text = '  ' + ref; rr.font.size = Pt(13); rr.font.bold = False; rr.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    LOG.append((sn, '제목', old, f'{ttl}  {ref}', '이사님 10/14 — 논문 캡션 그대로 + 그림·표 번호'))
# 4) 표지 제목 「에 대한 연구」
s = prs.slides[0]; sh = [x for x in s.shapes if x.shape_id == 5][0]
old = sh.text_frame.text; ps = sh.text_frame.paragraphs
lines = ['환경경영시스템(ISO 14001)과', '전과정평가(LCA)의', '프로세스 연계 모델 개발 및', '성숙도 평가에 대한 연구']
while len(sh.text_frame.paragraphs) < 4:
    sh.text_frame._txBody.append(copy.deepcopy(ps[-1]._p))
ps = sh.text_frame.paragraphs
for p_, ln in zip(ps, lines):
    p_.runs[0].text = ln
    for r in p_.runs[1:]: r._r.getparent().remove(r._r)
    p_.runs[0].font.size = Pt(24)
LOG.append((1, '표지 제목', old, ' / '.join(lines), '논문 제목과 동일하게 — 「에 대한 연구」'))
# 5) 표지 영문 제목 — 국문 제목 아래 작은 글
src_box = [x for x in s.shapes if x.shape_id == 6][0]
el = copy.deepcopy(src_box._element); s.shapes._spTree.append(el); eb = s.shapes[-1]
el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr').set('id', '8'); eb.name = 'EnTitle'
tf = eb.text_frame
for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
p0 = tf.paragraphs[0]; p0.runs[0].text = 'A Study on the Development and Maturity Assessment of a Process Linkage Model between Environmental Management Systems (ISO 14001) and Life Cycle Assessment (LCA)'
for r in p0.runs[1:]: r._r.getparent().remove(r._r)
p0.runs[0].font.size = Pt(10.5); p0.runs[0].font.bold = False; p0.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
eb.left = sh.left; eb.width = sh.width; eb.top = Emu(int(4.55 * 914400)); eb.height = Emu(int(0.6 * 914400))
LOG.append((1, '표지 영문 제목', '(없음)', p0.runs[0].text, '심사지 영문 제목 칸 대비'))
prs.save(dst); json.dump(LOG, open(f'log28_{kind}.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1); print('ok', kind, len(LOG))
