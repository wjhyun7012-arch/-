# v28 15분판 -> v30: 원래 구성에 보태기 — 2장 한 장(S-5 복귀), 매핑↔LMA 순서·30·25, 검토자 지적(숫자·용어), 진행 안내 삭제, 목차 상자 삭제
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptlib import rep, LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def set_note(sn, text, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; old = tf.text; tf.text = text; LOG.append((sn, '발표 노트(전체)', old, text, why))
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text; assert t.count(old) == 1, (sn, old, t.count(old)); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def set_lines(sn, sid, lines, why, size=None):
    sh = shape(prs.slides[sn-1], sid); tf = sh.text_frame; old = tf.text
    p0 = tf.paragraphs[0]; pPr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    r0 = p0.runs[0]; rPr = copy.deepcopy(r0._r.rPr) if r0._r.rPr is not None else None
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    r0.text = lines[0]
    if size: r0.font.size = Pt(size)
    for ln in lines[1:]:
        p = tf.add_paragraph()
        if pPr is not None: p._p.insert(0, copy.deepcopy(pPr))
        r = p.add_run(); r.text = ln
        if rPr is not None: r._r.insert(0, copy.deepcopy(rPr))
        if size: r.font.size = Pt(size)
    LOG.append((sn, f'[{sid}]', old, '\n'.join(lines), why))
def remove_shape(sn, sid, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text if sh.has_text_frame else '[도형]'; sh._element.getparent().remove(sh._element); LOG.append((sn, f'[{sid}] 삭제', old, '(삭제)', why))
def set_title(sn, text, why):
    sh = shape(prs.slides[sn-1], 2); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    LOG.append((sn, '제목', old, text, why))

# ---- 옛 번호(v28) 기준 수정
# 2 목차
remove_shape(2, 22, '이사님 10/14 — 화면의 진행 안내 삭제')
set_note(2, '논문의 장 순서대로, 1장 왜 이 연구인가, 2장 무엇을 바탕으로 하였는가, 3장 무엇을 만들었는가, 4장 적용하면 어떻게 되고 무엇을 고치는가, 5장 결론입니다.\n〔15초. 진행 안내는 말하지 않는다. 세부는 구분 슬라이드 뒤 백업.〕', '진행 안내 삭제')
# 4 개요 — 첫 문장(진행 안내)만 삭제
note_rep(4, '먼저 이 연구를 한 장으로 말씀드리겠습니다.\n', '', '진행 안내 삭제')
# 5
note_rep(5, '중대한 환경측면을 정하고', '중대한 환경측면을 결정하고', '논문 용어 「결정」')
# 20 (S-5) → 2장 본 장
set_one(20, 3, 'Ⅱ. 학술적 배경 — 무엇을 바탕으로', '백업 → 본(2장)')
sh = shape(prs.slides[19], 3); sh.width = Emu(int(3.3*IN)); sh.left = Emu(int((11.15-3.3)*IN))
set_one(20, 7, '근거  2.1~2.5, [표 2-3], [그림 2-7]', '2장')
set_note(20, '2장은 두 표준과 그 사이입니다. ISO 14001은 환경측면의 결정에서 전과정 관점을 고려하도록 요구하되 정량적 평가 방법을 정하지 않았고, 전과정평가(LCA) 표준인 ISO 14044는 계산 방법을 정하되 그 결과를 되풀이하여 관리하는 주기는 정하지 않았습니다. 두 표준은 서로 비어 있는 곳을 상대가 채울 수 있습니다.\n아래 두 수치로 규모의 차이를 보실 수 있습니다. 환경경영시스템(ISO 14001) 유효 인증은 국내 28,137건으로 세계 2위이고, LCA 수행을 전제하는 환경성적표지 유효 인증은 707개 기업입니다. 두 체계를 함께 운영하는 조직은 아직 제한적입니다.\n성숙도를 평가하는 방법은 새로 만들지 않고 소프트웨어·데이터 분야에서 쓰던 국제표준에서 가져왔고, 두 표준의 연계 정도를 조항 단위로 판정한 선행연구는 찾아보기 어려웠습니다. 그 빈 자리가 이 연구의 위치입니다.\n〔1분 10초. 가져온 것 — 세 계층 구조 ISO/IEC 33004 · 최소충족·비보상 판정 KS X ISO 8000-62 · 측정 틀 ISO/IEC 33020, 점수별 상태 기술 BARS. 선례 — 국내 공공데이터 연구(PPDQM3·APDQM3). 따른 표준은 백업 S-10, 선행연구 대비표 [표 2-4]·[표 2-5]·[그림 2-7]은 S-11, ISO 14001 구조 S-8, LCA 4단계 S-9, 연계의 성숙도 수준 [그림 2-5]는 2.3. 「8000-62를 왜 가져왔나」 → B-Q7.〕', '2장 한 장 — 선행연구 한 문장 보탬, 진행 안내 삭제')
# 7 매핑 (→ 새 7)
set_title(7, '두 표준의 조항 매핑 — 30개 조항과 25개', '30·25')
set_one(7, 5, '▶ 두 표준의 shall·should 문장을 전부 대조하고, 조항마다 같은 네 가지 질문으로 네 가지 등급을 판정하였다.', '요지')
set_one(7, 7, '근거  3.2.2, 3.3.1, [표 3-2], 부록 A.1·A.2·A.4', '근거')
set_lines(7, 27, ['ISO 14001 — 요구사항 조항 30개 전부가 문항이 됨   ·   ISO 14044 — 조항 25개가 문항의 근거(나머지 3개도 세부 조항의 문항·체크포인트로 반영) → 빠진 조항 없음. 셈은 부록 A.4',
                  '매우 높음 4개 조항 = 핵심 문항 4종(P-9·P-14·D-1·C-1) — 관리 주기의 계획 둘·실행 하나·점검 하나'], '30·25·핵심 4(셈은 지시문·노트북)', 12.5)
sh = shape(prs.slides[6], 27); sh.top = Emu(int(6.2*IN)); sh.height = Emu(int(1.0*IN))
set_note(7, '3장의 출발은 매핑입니다. 두 표준의 의무와 권고 문장을 전부 대조했습니다.\nISO 14001은 요구사항 조항 서른 개가 모두 문항이 되었고, ISO 14044는 스물여덟 개 가운데 스물다섯 개가 문항의 근거가 되었으며 나머지 세 개도 세부 조항의 문항과 체크포인트로 들어가 빠진 조항은 없습니다.\n서른 개 조항 하나하나에 같은 네 가지 질문을 같은 순서로 물어 네 가지 등급으로 판정했습니다. 매우 높음이 네 조항, 환경측면, 환경목표, 운용기획 및 관리, 모니터링입니다.\n이 네 곳이 연계 지점이고 뒤의 핵심 문항 4종이 되며, 관리 주기의 계획에 둘, 실행과 점검에 하나씩 놓입니다.\n〔1분 30초. 30 = 원표 34행에서 7.4.1~7.4.3을 하나로, 상위 조항 6.1·6.2 제외(3.2.2). 25·28 = 부록 A.4. 대조 대상은 shall·should — may·can은 제외(3.1.1). 55문항 근거 성격 shall 48·비교주장 시 조건부 shall 2·수행 시 조건부 shall 2·should 1·shall+should 2(3.3.1)는 노트북. 등급 요건은 B-1a, 흐름 B-1.〕', '이사님 10/14 문안')
# 6 LMA (→ 새 8)
set_one(6, 5, '▶ 매핑에서 연계 프로세스 16종을 도출하고, 그 성과를 55문항으로 평가하며, 점수를 성숙도 규칙으로 판정한다.', '요지')
set_lines(6, 8, ['환경경영시스템(ISO 14001)', '요구사항 조항 30개'], 'LMA')
set_lines(6, 9, ['전과정평가(LCA, ISO 14044)', '조항 25개(전조항 28개 가운데)'], 'LMA')
set_lines(6, 11, ['앞 장의 조항 매핑'], 'LMA')
set_lines(6, 12, ['매핑의 결과 = 연계 지점 · 그 가운데 핵심 네 곳 → 세 모델의 재료'], 'LMA')
set_lines(6, 15, ['연계 프로세스 16종 — 계획 8 · 실행 2 · 점검 3 · 개선 3'], 'LMA')
set_lines(6, 20, ['프로세스의 성과 하나에 문항 하나 — 확인 항목 개수로 1~5점'], 'LMA')
set_note(6, '조항 매핑의 결과로 만든 것이 셋입니다.\n첫째 참조 모델은 연계 지점을 업무 단위로 묶은 연계 프로세스 열여섯 가지로, 계획 여덟, 실행 둘, 점검 셋, 개선 셋입니다. 둘째 평가 모델은 프로세스의 성과 하나에 문항 하나씩 쉰다섯 문항이고, 문항마다 채점 기준과 확인할 증거를 정했습니다. 셋째 성숙도 모델은 점수를 여섯 수준으로 바꾸는 규칙이고, 핵심 문항 넷 가운데 하나라도 빠지면 ML2에서 멈춥니다.\n이 셋을 합쳐 연계 성숙도 평가 모델, LMA라고 하며, 만든 뒤 추적성과 표준 적합성 두 방향으로 검증했습니다.\n〔1분 10초. 논문 [그림 3-2] 전체는 백업 B-7, 방법론 구성 [그림 3-1]은 B-2. 검증 — 3.5.1 추적성, 3.5.2 ISO/IEC 33004 24건(충족 22·해당 없음 1·보완 1).〕', '이사님 10/14 문안')
# 9 16종
note_rep(9, '각 업무의 성과가 되어 있는지 묻는 것이 다음 장의 평가 문항 쉰다섯 개입니다.', '각 업무의 성과가 되어 있는지 묻는 것이 평가 문항 쉰다섯 개입니다.', '진행 안내 삭제')
# 10
rep(prs, 10, 15, '추가 항목(should)을 적어 둠', '추가 항목(should 또는 체계화·지속 상태)을 적어 둠', '검토자 — 3.3.2')
note_rep(10, '점수를 성숙도 수준으로 바꾸는 규칙이 다음 장의 성숙도 모델입니다.\n', '', '진행 안내 삭제')
# 11
rep(prs, 11, 22, '핵심 문항 4종 모두 4점,', '핵심 문항 4종 모두 4점 이상,', '[표 3-7] ≥ 4')
note_rep(11, '연계가 아예 끊긴 것과 연계가 얕은 것은 다르기 때문에', '연계가 성립하지 않은 것과 성숙도가 낮은 것은 다르기 때문에', '논문 표현')
note_rep(11, '이 모델을 한 회사에 적용한 결과를 보겠습니다.\n', '', '진행 안내 삭제')
# 12
rep(prs, 12, 27, '문서 심사(면담·현장 관찰 없음)', '문서 심사 — 면담·실행 상태의 현장 관찰은 포함하지 않음', '검토자')
note_rep(12, '4장은 이 모델을 한 회사에 적용한 결과입니다. 평가가 어떻게 흘러갔는지 그림으로 먼저 보겠습니다.', '4장은 이 모델을 한 회사에 적용한 결과입니다.', '진행 안내 삭제')
note_rep(12, '낮은 점수의 원인 세 가지에서 개선 과제 다섯 건을 도출하였고', '구조적 미비 세 가지와 공통 원인 두 가지에서 개선 과제 다섯 건을 도출하였고', '4.5·4.6')
note_rep(12, '평가는 문서 심사만으로, 평가자 한 명이,', '평가는 문서 심사로 면담과 실행 상태의 현장 관찰은 포함하지 않고, 평가자 한 명이,', '검토자')
note_rep(12, '판정 결과는 다음 장입니다.\n', '', '진행 안내 삭제')
# 13
note_rep(13, '그러면 낮은 점수는 어디에서 왔는지, 다음 장에서 보겠습니다.\n', '', '진행 안내 삭제')
# 14
note_rep(14, 'LCA를 다시 할 때가 아직 오지 않았습니다.', 'LCA 연 1회 재수행 규정은 있으나 아직 이행 전입니다.', '4.5.1 ③')
# 15
note_rep(15, '미비 사항 3건에 대응하는 개선 과제 5건을 도출하였습니다.', '구조적 미비 3건과 공통 원인 2건에 대응하는 개선 과제 5건을 도출하였습니다.', '4.5·4.6')
note_rep(15, '끝으로 이 모델의 장점과 단점, 한계입니다.\n', '', '진행 안내 삭제')
# 16 장단점·한계 — 단점 7건
set_one(16, 5, '▶ 장점 여섯, 단점 일곱, 한계 셋 — 한계는 그대로 향후 연구 과제다.', '단점 7건(5.2)')
set_lines(16, 9, ['단점', '· 핵심 문항 4종에 대한 의존 — 네 문항의 선택은 전문가 검토 전', '· 판정 파라미터 — 과반 50% 등은 연구자가 정한 값', '· 평가의 부담 — 55문항 문서 확인, 두 분야를 다 알아야', '· 적용 대상의 제약 — LCA를 안 하는 회사는 차이가 안 난다', '· 문항 간 종속 — 미비 하나가 여러 점수에 함께 영향', '· 문서 심사 방식 — 형식적 이행은 못 가림', '· 환경성과와의 관계 미검증 — 탄소·비용이 주는지는 확인 안 함'], '단점 7건으로 분리', 13.5)
set_note(16, '이 모델의 장점과 단점, 그리고 한계입니다. 장점 여섯, 단점 일곱, 한계 셋입니다.\n장점은 문항의 출처를 조항까지 따라갈 수 있다는 것, 같은 증거면 같은 점수가 나온다는 것, 종합점수에 묻히는 미비를 드러낸다는 것, 평가가 개선과 재평가로 이어진다는 것, 판정 구조와 평가 방법을 국제표준에서 가져왔다는 것, 수준의 뜻을 체계 단계와 숫자 확인 단계로 나눈 것입니다.\n단점은 판정이 핵심 문항 4종에 좌우된다는 것, 기준값이 연구자 확정값이라는 것, 평가에 손이 많이 간다는 것, LCA를 안 하는 회사는 차이가 안 난다는 것, 미비 하나가 여러 점수에 함께 영향을 준다는 것, 문서 심사라 형식적 이행을 가리기 어렵다는 것, 환경성과와의 관계는 검증하지 않았다는 것입니다. 이 가운데 넷은 전문가 조사와 다사례 적용으로 보완하고, 셋은 설계에서 비롯되는 제약입니다.\n한계는 한 회사·한 시점·평가자 한 명, 전문가 타당도 조사 미실시, 가장 높은 두 수준의 사례 미확인이고, 이 셋이 그대로 향후 연구 과제입니다.\n〔1분. 화면 글은 쉬운 말판, 항목 이름은 [표 5-1] 그대로. 단점 넷(핵심 문항 지정·판정 파라미터·문항 간 종속·환경성과)은 보완 가능, 셋(평가 부담·적용 대상·문서 심사)은 설계 제약(5.2). 한계는 논문 5.3에 여덟 항목(단일 사례·시점·평가자 / 문항 간 종속 / ML4·ML5 미검증 / 표준 판의 범위 / 타당화 확장(CVI) / 진단·개선 고도화 / 미측정 구간 / 표준화 경로) — [그림 5-3] 백업 B-0c③. 장점·단점 전문은 B-15a·B-15b, 연구 전체 한 장 요약은 S-28. 「환경성과와의 관계」 → 전문 46번.〕', '단점 7건·넷/셋, 한계 여덟 항목 지시문')
# 18 감사 · 19 구분
note_rep(18, 'S-5~S-28(본심사판에서 옮긴 12장)', 'S-7~S-28(본심사판에서 옮긴 11장)', '백업 수')
rep(prs, 19, 2, 'S-5~S-28 본심사판에서 옮긴 12장(두 표준의 빈 곳 · 연구 목적과 문제', 'S-7~S-28 본심사판에서 옮긴 11장(연구 목적과 문제', '백업 목록')
# ---- 순서: 1 2 3 4 5 | 20 | 7 6 | 8 9 10 11 12 13 14 15 16 17 18 | 19 | 21..31 | 32..75
N = len(prs.slides)
MAIN = [1, 2, 3, 4, 5, 20, 7, 6, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]
order = MAIN + [19] + list(range(21, N+1))
assert sorted(order) == list(range(1, N+1))
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, n in enumerate(nums): out += str(old2new.get(int(n), int(n))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = s.notes_slide.notes_text_frame; t = tf.text; t2 = tr_nums(t)
    if t2 != t: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new, 'main': MAIN}, open('log30.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), 'slides', len(prs.slides), [old2new[o] for o in MAIN])
