import copy
from pptx import Presentation
LOG=[]
def _rep_par(p,old,new):
    runs=list(p.runs); full=''.join(r.text for r in runs)
    i=full.find(old)
    if i<0: return 0
    assert full.count(old)==1,(full,old)
    j=i+len(old); pos=0; spans=[]
    for k,r in enumerate(runs):
        a,b=pos,pos+len(r.text); pos=b
        if b>i and a<j or (a==b==i): spans.append((k,a,b))
    k0,a0,b0=spans[0]; k1,a1,b1=spans[-1]
    if k0==k1:
        t=runs[k0].text; runs[k0].text=t[:i-a0]+new+t[j-a0:]
    else:
        runs[k0].text=runs[k0].text[:i-a0]+new
        for k,a,b in spans[1:-1]: runs[k].text=''
        runs[k1].text=runs[k1].text[j-a1:]
    return 1
def _rep_tf(tf,old,new):
    return sum(_rep_par(p,old,new) for p in tf.paragraphs)
def shape(slide,sid):
    def walk(shs):
        for s in shs:
            if s.shape_id==sid: return s
            if s.shape_type==6:
                r=walk(s.shapes)
                if r: return r
    r=walk(slide.shapes); assert r is not None,sid; return r
def rep(prs,sn,sid,old,new,why=''):
    s=prs.slides[sn-1]; sh=shape(s,sid); n=_rep_tf(sh.text_frame,old,new)
    assert n==1,(sn,sid,old,n); LOG.append((sn,f'[{sid}]',old,new,why))
def rep_cell(prs,sn,sid,r,c,old,new,why=''):
    s=prs.slides[sn-1]; sh=shape(s,sid); cell=sh.table.cell(r,c); n=_rep_tf(cell.text_frame,old,new)
    assert n==1,(sn,sid,r,c,old,n); LOG.append((sn,f'표 {r+1}행 {c+1}열',old,new,why))
def rep_notes(prs,sn,old,new,why=''):
    s=prs.slides[sn-1]; n=_rep_tf(s.notes_slide.notes_text_frame,old,new)
    assert n==1,(sn,'notes',old,n); LOG.append((sn,'발표 노트',old,new,why))
def rep_notes_all(prs,old,new,why=''):
    c=0
    for i,s in enumerate(prs.slides,1):
        if s.has_notes_slide and old in s.notes_slide.notes_text_frame.text:
            c+=_rep_tf(s.notes_slide.notes_text_frame,old,new)
    LOG.append(('여러 장',f'발표 노트 {c}곳',old,new,why)); return c
def set_img(prs,sn,sid,path,why=''):
    s=prs.slides[sn-1]; sh=shape(s,sid); rid=sh._element.blipFill.blip.rEmbed
    part=s.part.related_part(rid); part._blob=open(path,'rb').read(); LOG.append((sn,f'[{sid}] 그림','(v6 그림)',path.split('/')[-1],why))
def add_row(prs,sn,sid,after,cells):
    s=prs.slides[sn-1]; sh=shape(s,sid); tbl=sh.table._tbl
    trs=tbl.tr_lst; new=copy.deepcopy(trs[after]); trs[after].addnext(new)
    t=sh.table
    for c,txt in enumerate(cells):
        tf=t.cell(after+1,c).text_frame; p=tf.paragraphs[0]
        for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
        rs=p.runs
        if rs:
            rs[0].text=txt
            for r in rs[1:]: r._r.getparent().remove(r._r)
        else: p.text=txt
    LOG.append((sn,f'표 {after+2}행 추가','',' | '.join(cells),''))
