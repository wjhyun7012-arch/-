"""v18 → v19: 용어 줄 재배치(4·5·8번) + 26번 아래 글 줄이기 (이사님 10/7)."""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Pt, Emu
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); LOG = []
def terms(sn, txt, why):
    s = prs.slides[sn - 1]; tb = [sh for sh in s.shapes if sh.name == f'Terms {sn}'][0]
    p = tb.text_frame.paragraphs[0]; rs = p.runs; old = ''.join(r.text for r in rs); rs[1].text = txt
    for r in rs[2:]: r._r.getparent().remove(r._r)
    LOG.append((sn, '용어 줄', old, '용어  ' + txt, why))
terms(4, '전과정 관점 = 위 여섯 단계를 함께 보는 것  ·  전과정평가(LCA: Life Cycle Assessment)  ·  환경측면 = 활동·제품 가운데 환경에 영향을 줄 수 있는 요소  ·  중대한 환경측면 = 그 가운데 영향이 큰 것으로 결정한 것 → 환경목표에 반영', '환경측면·중대한 환경측면은 4번 왼쪽 상자에 처음 나옴')
terms(5, '전과정 목록분석(LCI: Life Cycle Inventory) = 제품 한 단위(기능단위)에 들어가고 나오는 원료·에너지·배출을 모두 세는 단계  ·  전과정 영향평가(LCIA: Life Cycle Impact Assessment) = 그 목록이 기후변화·산성화 등에 미치는 영향을 계산하는 단계', '표 2행이 되며 환경측면·PDCA가 5번 화면에서 사라짐')
terms(8, '핵심 문항 = 연계가 성립하려면 반드시 갖추어야 할 네 곳을 묻는 문항 — 하나라도 빠지면 평균이 높아도 수준이 올라가지 않음  ·  PDCA(Plan–Do–Check–Act) = 계획–실행–점검–조치의 관리 순환', 'PDCA는 8번(조항 구조)에 처음 나옴')
# 26번 아래 글
s = prs.slides[25]; sh = [x for x in s.shapes if x.shape_id == 10][0]; tf = sh.text_frame; old = tf.text
for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
p = tf.paragraphs[0]; rs = p.runs; rs[0].text = '완수 시 값은 지금 점수로 다시 계산한 전망이며 이행 결과가 아니다 · R6(외부 정밀검토)은 ML4를 향한 장기 과제라 수치에 넣지 않음 · 단기 1년 이내(LCA 연 1회 재수행 주기) · 중기 3년 이내(ISO 14001 인증 갱신 주기) · 장기 3년 초과 — 4.6.3'
for r in rs[1:]: r._r.getparent().remove(r._r)
for r in p.runs: r.font.size = Pt(11)
LOG.append((26, '[10] 아래 글', old, rs[0].text, '이사님 10/7: 아래 글 줄이기 — 두 문단을 한 줄로'))
prs.save(OUT); json.dump(LOG, open('log19.json', 'w'), ensure_ascii=False, indent=1); print(len(LOG), 'changes')
