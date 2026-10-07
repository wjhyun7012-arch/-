"""v13 → v14: 「서로를 전제하지 않는다」 → (가) 「별개의 표준이라 한쪽이 다른 쪽을 요구하지 않는다」(이사님 10/7). 28번 화면은 논문 문장 그대로, 노트에서 풀이."""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptlib import _rep_tf
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); LOG = []
def rep_shape(sn, name, old, new, why):
    s = prs.slides[sn - 1]; sh = [x for x in s.shapes if x.name == name][0]
    n = _rep_tf(sh.text_frame, old, new); assert n == 1, (sn, name, n); LOG.append((sn, name, old, new, why))
def rep_notes(sn, old, new, why):
    n = _rep_tf(prs.slides[sn - 1].notes_slide.notes_text_frame, old, new); assert n == 1, (sn, old, n); LOG.append((sn, '노트', old, new, why))
rep_shape(4, 'key line', '두 표준은 서로를 전제하지 않아 실무에서는 별도로 운영되는 경우가 많다', '두 표준은 별개의 표준이라 한쪽이 다른 쪽을 요구하지 않고, 실무에서는 별도로 운영되는 경우가 많다', '이사님 10/7 (가)안 — 「별개 표준」은 논문 1.1')
rep_notes(4, '두 표준은 서로를 전제하지 않아, 실무에서는 별도로 운영되는 경우가 많습니다.', '두 표준은 별개의 표준이라 한쪽이 다른 쪽을 요구하지 않습니다. 그래서 실무에서는 별도로 운영되는 경우가 많습니다.', '같은 이유')
rep_notes(1, '두 국제표준이 서로를 전제하지 않는 탓에 현장에서 따로 운영되는 것을,', '두 국제표준이 별개의 표준이라 현장에서 따로 운영되는 것을,', '같은 이유')
rep_notes(28, '환경경영시스템(ISO 14001)과 전과정평가(LCA)는 각각 성숙한 표준이나 서로를 전제하지 않습니다.', '환경경영시스템(ISO 14001)과 전과정평가(LCA)는 각각 성숙한 표준이나 서로를 전제하지 않습니다. 곧 한쪽이 다른 쪽을 요구하지 않습니다.', '화면은 논문 마지막 문단 그대로, 노트에서 한 번 풀어 읽음')
prs.save(OUT); json.dump(LOG, open('log14.json', 'w'), ensure_ascii=False, indent=1); print(len(LOG), 'changes')
