# v20 -> v21: 백업 75번 B-Q13 「생성형 AI 활용 범위와 저자 책임」 추가
import sys, copy, json, unicodedata
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from lxml import etree
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
LOG = []
base = prs.slides[73]                       # 74번 B-Q12
layout = base.slide_layout
s = prs.slides.add_slide(layout)
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
for sh in base.shapes: s.shapes._spTree.append(copy.deepcopy(sh._element))
def by_id(sl, i): return [x for x in sl.shapes if x.shape_id == i][0]
def set_text(sh, txt):
    p = sh.text_frame.paragraphs[0]; p.runs[0].text = txt
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for extra in sh.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)
set_text(by_id(s, 2), '생성형 AI 활용 범위와 저자 책임')
set_text(by_id(s, 3), '백업  B-Q13')
set_text(by_id(s, 5), '75')
body = by_id(s, 6); body._element.getparent().remove(body._element)
# ▶ 길잡이 — 32번 Guide 복사
g_src = [x for x in prs.slides[31].shapes if x.name == 'Guide 32'][0]
el = copy.deepcopy(g_src._element); s.shapes._spTree.append(el); hd = s.shapes[-1]
el.find('.//p:cNvPr', NS).set('id', '6'); hd.name = 'Guide 75'
set_text(hd, '▶ 연구와 결정은 연구자가, 정리·대조·형식은 AI가 보조 — 조항·수치는 원문과 연구자가 확인')
hd.top = Emu(g_src.top); hd.left = Emu(g_src.left); hd.width = Emu(g_src.width); hd.height = Emu(int(0.463 * IN))
# 세 상자 — 23번 path0 복사
box_src = [x for x in prs.slides[22].shapes if x.name == 'path0'][0]
BOXES = [
 ('연구자가 한 것', '1F3864', 'E8EEF7', [
   '논문 전체 설계 — 연구 문제, 구성안, 로드맵',
   '두 표준 전 조항 대조(2025.10 검토표 → 2026.3 전조항 분석)',
   '55문항·판정 포인트, 채점 기준의 원칙, 핵심 문항 4종 지정',
   'A사 문서 16종·기록물 33건 수집, 55문항 판정·재검증',
   '개선 과제 R1~R5 도출, 결론 — 그리고 모든 결정']),
 ('AI를 도구로 쓴 것', '595959', 'F2F2F2', [
   '연구자가 세운 구성안·로드맵의 검토·수정',
   '55문항 → 연계 프로세스 16종 편성 초안(기준·명칭·경계는 연구자가 확정, 지도교수 검토)',
   '표준 원문 확보 뒤 대조 보조 — ISO/IEC 33004 자체점검 초안, 판정 규칙의 표준 정합',
   '문장 다듬기·용어 통일, 표·그림·상호 참조 정합 점검',
   '그림 작성, 검토표 작성, 발표자료 편집']),
 ('원칙과 확인', 'C00000', 'FBEAEA', [
   '표준 조항·인용은 원문을 확인한 것만 넣음',
   '수치·판정은 연구자가 산출·확인',
   '지도교수 검토를 거침',
   '대학원 지침(2026.2.2)·학위청구논문 심사 안내·학위수여규정·연구윤리준수확인서(서식 6)에 생성형 AI 조항 없음 확인(2026-09-29)',
   '논문에 별도 표기는 두지 않고, 물으면 범위를 밝힘'])]
L0 = 0.54 * IN; GAP = 0.15 * IN; W = (10.60 * IN - 2 * GAP) / 3; TOP = 1.40 * IN; H = 4.95 * IN
for k, (ttl, col, fill, lines) in enumerate(BOXES):
    el = copy.deepcopy(box_src._element); s.shapes._spTree.append(el); b = s.shapes[-1]
    el.find('.//p:cNvPr', NS).set('id', str(7 + k)); b.name = f'ai box{k}'
    b.left = Emu(int(L0 + k * (W + GAP))); b.width = Emu(int(W)); b.top = Emu(int(TOP)); b.height = Emu(int(H))
    el.find('.//a:solidFill/a:srgbClr', NS).set('val', fill)
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.TOP; tf.word_wrap = True
    ps = tf.paragraphs
    p0, p1 = ps[0], ps[1]
    for extra in ps[2:]: extra._p.getparent().remove(extra._p)
    p0.runs[0].text = ttl; p0.runs[0].font.size = Pt(16); p0.runs[0].font.color.rgb = __import__('pptx.dml.color', fromlist=['RGBColor']).RGBColor.from_string(col)
    p0._p.find('a:pPr', NS).find('a:spcAft', NS).find('a:spcPts', NS).set('val', '900')
    tmpl = copy.deepcopy(p1._p); p1._p.getparent().remove(p1._p)
    for ln in lines:
        np_ = copy.deepcopy(tmpl); b.text_frame._txBody.append(np_)
        np_.find('a:pPr', NS).set('algn', 'l'); np_.find('a:pPr', NS).find('a:spcAft', NS).find('a:spcPts', NS).set('val', '800')
        r = np_.find('a:r', NS); r.find('a:rPr', NS).set('sz', '1400'); r.find('a:t', NS).text = '· ' + ln
# 아래 근거 줄 — 23번 sens note 복사
n_src = [x for x in prs.slides[22].shapes if x.name == 'sens note'][0]
el = copy.deepcopy(n_src._element); s.shapes._spTree.append(el); nb = s.shapes[-1]
el.find('.//p:cNvPr', NS).set('id', '10'); nb.name = 'ai note'
set_text(nb, '근거  예상 질문 43번 답 · 면담 자료 5.2 나-3 · 생성형 AI 조항 유무는 대학원 지침·심사 안내·학위수여규정·서식 6에서 확인(2026-09-29)')
nb.top = Emu(int(6.55 * IN)); nb.left = Emu(int(L0)); nb.width = Emu(int(10.60 * IN))
# 노트
NOTE = ('생성형 AI는 도구로 썼고, 그 범위를 밝히겠습니다. 논문 전체 설계와 두 표준의 전 조항 대조, 55문항과 채점 기준, A사 판정과 개선 과제, 그리고 모든 결정은 제가 했습니다. '
        'AI는 제가 세운 구성안과 로드맵의 검토·수정, 55문항을 연계 프로세스 16종으로 편성하는 초안, 표준 원문을 확보한 뒤의 대조 보조, 문장 다듬기와 상호 참조 점검, 그림과 검토표 작성에 썼습니다. '
        '표준 조항과 인용은 원문을 확인한 것만 넣었고, 수치와 판정은 제가 산출하여 지도교수님 검토를 거쳤습니다. '
        '대학원 지침과 심사 안내, 학위수여규정, 서식 6에 생성형 AI 조항이 없음을 확인하여 논문에 별도 표기는 두지 않았습니다.')
ntf = s.notes_slide.notes_text_frame; ntf.text = NOTE
LOG.append((75, '새 장 B-Q13', '(없음)', '제목 「생성형 AI 활용 범위와 저자 책임」 · ▶ 길잡이 · 세 상자(연구자가 한 것 / AI를 도구로 쓴 것 / 원칙과 확인) · 근거 줄 · 노트 대본', '이사님 10/8 「백업으로 가자」 — 예상 질문 43번 답 기준, 「구성안·로드맵은 연구자가 설계, AI는 검토·수정」'))
prs.save(dst); json.dump(LOG, open('log21.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('slides', len(prs.slides))
