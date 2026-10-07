"""v15 → v16: 「연계(linkage)」 첫 등장 네 곳(이사님 10/7). 4번 제목·작은 글, 6번 빨간 줄, 14번 용어 줄."""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptlib import _rep_tf
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); LOG = []
def rep(sn, pick, old, new, why):
    s = prs.slides[sn - 1]; sh = [x for x in s.shapes if x.has_text_frame and pick(x)][0]
    n = _rep_tf(sh.text_frame, old, new); assert n == 1, (sn, old, n); LOG.append((sn, sh.name, old, new, why))
rep(4, lambda x: x.shape_id == 2, '두 표준과 연계', '두 표준과 연계(linkage)', '이사님 10/7: 연계 첫 등장에 영문')
rep(4, lambda x: x.name == 'sub lines', '연계란 — LCA 결과가', '연계(linkage)란 — LCA 결과가', '같은 이유')
rep(6, lambda x: x.shape_id == 9, '본 연구가 말하는 연계 = 상태 B', '본 연구가 말하는 연계(linkage) = 상태 B', '같은 이유')
rep(14, lambda x: x.name == 'Terms 14', '연계 지점 = 매핑에서 확인된', '연계 지점(linkage point) = 매핑에서 확인된', '논문 1.2.1 영문 표기')
prs.save(OUT); json.dump(LOG, open('log16.json', 'w'), ensure_ascii=False, indent=1); print(len(LOG), 'changes')
