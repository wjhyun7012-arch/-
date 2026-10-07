"""v11 → v12 (이사님 10/7 확정 문안 + 규칙표 v18 반영).
(1) 4번을 「이 연구를 한 장으로 — 두 표준과 연계」 개요 장으로 재구성
(2) 5·7·13번 용어 줄 수정, 약어 영문 전체를 용어 줄 첫 등장에 괄호로
(3) 규칙 위반·경고 수정: 재는→측정/판정, ML3가→ML3이, 향후 과제→향후 연구 과제, PDCA 풀이→계획–실행–점검–조치
사용: python3 -I mk_deck12.py v11.pptx v12.pptx
"""
import sys, os, json, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
from pptx.text.text import _Paragraph, _Run
from pptlib import _rep_tf
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); W, H = prs.slide_width, prs.slide_height
LOG = []
NAVY = RGBColor(0x1F, 0x38, 0x64); LIGHT = RGBColor(0xE8, 0xEE, 0xF7); RED = RGBColor(0xC0, 0x00, 0x00); PINK = RGBColor(0xFB, 0xEC, 0xEC); INK = RGBColor(0x11, 0x11, 0x11); GREY = RGBColor(0x59, 0x59, 0x59); GREEN = RGBColor(0x2E, 0x6B, 0x4A); LGREEN = RGBColor(0xEA, 0xF4, 0xEE)

def set_ea(run):
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn('a:ea')) is None:
        ea = etree.SubElement(rpr, qn('a:ea')); ea.set('typeface', '맑은 고딕')
def add_par(tf, text, size, bold=False, color=INK, align=None, first=False, space_after=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None: p.alignment = align
    if space_after is not None: p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = '맑은 고딕'; r.font.color.rgb = color; set_ea(r)
    return p
def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=''):
    sh = s.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h)); sh.name = name
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(110000); tf.margin_top = tf.margin_bottom = Emu(60000)
    return sh
def set_text_keep(sh, text):
    tf = sh.text_frame; p = tf.paragraphs[0]; rs = p.runs
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
def set_notes(sn, txt):
    tf = prs.slides[sn - 1].notes_slide.notes_text_frame; old = tf.text
    p0 = tf.paragraphs[0]; tmpl = copy.deepcopy(p0._p)
    for p in list(tf.paragraphs): p._p.getparent().remove(p._p)
    for line in txt.split('\n'):
        np_ = copy.deepcopy(tmpl); tf._txBody.append(np_); para = _Paragraph(np_, tf); rs = para.runs
        if rs:
            rs[0].text = line
            for r in rs[1:]: r._r.getparent().remove(r._r)
        else: para.text = line
    return old
def rep_notes(sn, old, new, why):
    tf = prs.slides[sn - 1].notes_slide.notes_text_frame; n = _rep_tf(tf, old, new); assert n == 1, (sn, old, n); LOG.append((sn, '노트', old, new, why))
def set_terms(sn, txt, why):
    s = prs.slides[sn - 1]; tb = [sh for sh in s.shapes if sh.name == f'Terms {sn}'][0]
    p = tb.text_frame.paragraphs[0]; rs = p.runs; old = ''.join(r.text for r in rs)
    rs[1].text = txt
    for r in rs[2:]: r._r.getparent().remove(r._r)
    LOG.append((sn, '용어 줄', old, '용어  ' + txt, why))

# ================= (1) 4번 개요 장 =================
s4 = prs.slides[3]
# 기존 내용 도형 삭제(제목·장 표시·선·▶·쪽번호·근거·용어 줄은 유지)
for sh in list(s4.shapes):
    if sh.shape_id in (2, 3, 4, 6, 7) or sh.name == 'Terms 4': continue
    if sh.has_text_frame and sh.text_frame.text.startswith('▶'): continue
    sh._element.getparent().remove(sh._element)
title = [sh for sh in s4.shapes if sh.shape_id == 2][0]; old_title = title.text_frame.text; set_text_keep(title, '이 연구를 한 장으로 — 두 표준과 연계')
hd = [sh for sh in s4.shapes if sh.has_text_frame and sh.text_frame.text.startswith('▶')][0]
hp = hd.text_frame.paragraphs[0]; old_hd = hd.text_frame.text; hp.runs[-1].text = '관리하는 틀과 측정하는 방법, 이 둘이 회사 안에서 이어져 도는지를 등급으로 판정한다.'
for r in hp.runs[1:-1]: r._r.getparent().remove(r._r)
basis = [sh for sh in s4.shapes if sh.shape_id == 7][0]; set_text_keep(basis, '근거  1.1, 1.2.1, [그림 1-1]')
L = 497190; RW = W - 2 * L
# 두 상자 + 가운데 「연계?」
bw = int(RW * 0.40); gap = RW - 2 * bw; y0 = 1330000; bh = 2700000
lb = box(s4, L, y0, bw, bh, LIGHT, NAVY, name='EMS box'); tf = lb.text_frame
add_par(tf, '환경경영시스템(ISO 14001)', 18, True, NAVY, PP_ALIGN.CENTER, first=True, space_after=6)
add_par(tf, '환경을 관리하는 틀', 15, True, INK, PP_ALIGN.CENTER, space_after=6)
add_par(tf, '환경측면을 결정하고 영향을 평가해 중대한 환경측면을 결정하고, 그것을 환경목표에 반영해 실행하고, 점검하여 지속적으로 개선한다', 14, False, INK, PP_ALIGN.CENTER)
rb = box(s4, L + bw + gap, y0, bw, bh, LGREEN, GREEN, name='LCA box'); tf = rb.text_frame
add_par(tf, '전과정평가(LCA)', 18, True, GREEN, PP_ALIGN.CENTER, first=True, space_after=6)
add_par(tf, '환경영향을 측정하는 방법', 15, True, INK, PP_ALIGN.CENTER, space_after=6)
add_par(tf, '제품이 원료 채취 → 설계 → 생산 → 운송·배송 → 사용 → 폐기의 각 단계에서 환경에 주는 영향을 숫자로 계산해 영향이 가장 큰 단계(Hotspot)를 찾는다 (ISO 14044)', 14, False, INK, PP_ALIGN.CENTER)
# 가운데 연결선 + 「연계?」
cx0 = L + bw; cy = y0 + bh // 2
ln = s4.shapes.add_connector(1, Emu(cx0 + 40000), Emu(cy), Emu(cx0 + gap - 40000), Emu(cy)); ln.line.color.rgb = NAVY; ln.line.width = Pt(2.25); ln.name = 'link line'
qb = box(s4, cx0 + gap // 2 - 520000, cy - 300000, 1040000, 600000, RGBColor(0xFF, 0xFF, 0xFF), NAVY, shape=MSO_SHAPE.OVAL, name='link q'); qb.text_frame.margin_left = qb.text_frame.margin_right = Emu(20000)
add_par(qb.text_frame, '연계?', 20, True, NAVY, PP_ALIGN.CENTER, first=True)
# 빨간 강조 한 줄
y1 = y0 + bh + 300000; h1 = 800000
kb = box(s4, L, y1, RW, h1, PINK, None, name='key line')
add_par(kb.text_frame, '둘은 서로를 전제하지 않아 따로 돈다 — 얼마나 이어졌는지를 측정하고 판정하는 모델, 이것이 이 연구가 만든 것이다', 16, True, RED, PP_ALIGN.CENTER, first=True)
# 아래 작은 글씨 두 줄
y2 = y1 + h1 + 260000
tb = s4.shapes.add_textbox(Emu(L), Emu(y2), Emu(RW), Emu(760000)); tb.name = 'sub lines'; tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = Emu(45720)
add_par(tf, '이어졌다는 것은 — LCA 결과가 환경측면의 결정·영향평가·환경목표·설계·점검에 실제로 쓰이는 것', 15, False, INK, first=True, space_after=4)
add_par(tf, '판정 결과는 — ML0(도입 전)부터 ML5(혁신화)까지 여섯 수준', 15, False, INK)
LOG.append((4, '화면 재구성', f'제목 「{old_title}」 / {old_hd} / 규제·함께 운영·별도 운영·측정 기준 부재 4줄 + 아래 문장', '제목 「이 연구를 한 장으로 — 두 표준과 연계」 / ▶ 관리하는 틀과 측정하는 방법… / 두 상자(환경경영시스템(ISO 14001)·전과정평가(LCA)) + 연계? + 빨간 한 줄 + 작은 글 2줄', '이사님 10/7: 개요를 1장(Ⅰ 첫 장)에서 개념만, CBAM·ESG는 노트로'))
set_terms(4, '전과정 관점 = 위 여섯 단계를 함께 보는 것  ·  전과정평가(LCA: Life Cycle Assessment)', '상자와 겹치는 단계 나열 줄임, 약어 영문 전체')
old4 = set_notes(4, """먼저 이 연구를 한 장으로 말씀드리겠습니다.
왼쪽 환경경영시스템(ISO 14001)은 회사가 환경을 관리하는 틀입니다. 환경측면을 결정하고, 그 영향을 평가해 중대한 환경측면을 결정하고, 그것을 환경목표에 반영해 실행하고, 점검하여 지속적으로 개선합니다.
오른쪽 전과정평가(LCA)는 제품이 원료 채취, 설계, 생산, 운송과 배송, 사용, 폐기의 각 단계에서 환경에 주는 영향을 숫자로 계산해 영향이 가장 큰 단계, 이것을 핫스팟이라 하는데, 그것을 찾는 방법입니다.
이 둘은 서로를 전제하지 않아서 현장에서는 따로 돕니다. LCA 보고서를 받아 놓고 경영시스템에서는 쓰지 않는 식입니다.
이 연구는 둘이 얼마나 이어졌는지를 측정하고 판정하는 모델을 만든 것입니다. 이어졌다는 것은 LCA 결과가 환경측면의 결정과 영향평가, 환경목표, 설계, 점검에 실제로 쓰인다는 뜻이고, 판정 결과는 여섯 수준으로 나옵니다.
배경으로는 제품 전과정의 환경정보를 요구하는 규제가 늘고 있다는 점이 있습니다.
〔1분. 「ISO 14001은 많이 쓰이나」가 나오면 — 국내 유효 인증 28,137건, 세계 676,232건(ISO Survey 2024), 2.1.1. CBAM·ESG 공시는 1.1〕""")
LOG.append((4, '노트', old4, '(개요 대본 — 위 화면과 같은 순서, 1분)', '같은 이유'))

# ================= (2) 용어 줄 =================
set_terms(5, '환경측면 = 조직의 활동·제품 가운데 환경에 영향을 줄 수 있는 요소(예: 전력 사용, 원료의 탄소 배출)  ·  중대한 환경측면 = 환경측면의 영향을 평가해 영향이 큰 것으로 결정한 것 — 이것을 환경목표에 반영한다(6.1.2 → 6.2.1)  ·  전과정 목록분석(LCI: Life Cycle Inventory) = 제품 한 단위(기능단위)에 들어가고 나오는 원료·에너지·배출을 모두 세는 단계  ·  전과정 영향평가(LCIA: Life Cycle Impact Assessment) = 그 목록이 기후변화·산성화 등에 미치는 영향을 계산하는 단계  ·  PDCA(Plan–Do–Check–Act) = 계획–실행–점검–조치의 관리 순환', '이사님 10/7: 중대한 환경측면·LCI·LCIA 각각 / 규칙 T10 PDCA 풀이 / 약어 영문')
set_terms(7, '매핑(mapping) = map은 「지도」지만 여기서는 「짝지어 대응시키다」는 뜻 — 두 표준의 조항을 하나하나 서로 대응시키는 일  ·  연계 프로세스 = 두 표준이 맞닿는 지점을 업무 단위로 묶은 것  ·  RQ = 연구 문제(Research Question)', '이사님 10/7: 매핑의 두 뜻')
set_terms(13, 'shall = 「반드시 해야 한다」는 의무, should = 권고 — KS 번역본에서 어느 문장이 의무이고 권고인지는 KS A 0001(표준의 서식과 작성방법) 부속서 I의 문장 말미 형태로 가름  ·  정밀검토 = LCA 결과를 제3자나 내부 전문가가 검토하는 절차(ISO 14044 6절)', '이사님 10/7: KS A 0001 근거')
set_terms(12 if False else 16, '체크포인트 = 문항마다 「이것이 있는가」를 확인하는 항목(필수·추가)  ·  BARS 앵커(BARS: Behaviorally Anchored Rating Scale, 행동기술척도) = 1~5점 각각이 어떤 상태인지 말로 적어 둔 기준 — 평가자가 달라도 같은 점수가 나오게 함', '약어 영문')
set_terms(18, '추적성 검증 = 문항마다 어느 표준 조항에서 나왔는지 따라갈 수 있는지 확인  ·  내용타당도 지수(CVI: Content Validity Index) = 전문가들이 「이 문항이 적절한가」를 평가한 비율  ·  객관적 증거 = 문서·기록처럼 남이 보아도 확인되는 증거', '약어 영문')
# LMA 영문: 12번에는 용어 줄이 없음 → 10번 용어 줄 끝에 덧붙임
s10 = prs.slides[9]; tb10 = [sh for sh in s10.shapes if sh.name == 'Terms 10'][0]; r10 = tb10.text_frame.paragraphs[0].runs[1]
old10 = r10.text; r10.text = old10 + '  ·  이 셋을 합쳐 연계 성숙도 평가 모델(LMA: Linkage Maturity Assessment)'; LOG.append((10, '용어 줄', old10, r10.text, '약어 영문 LMA(12번에 용어 줄 없음)'))

# ================= (3) 규칙 수정 =================
rep_notes(6, '그 정도를 재는 자가 성숙도 모델입니다.', '그 정도를 판정하는 것이 성숙도 모델입니다.', '규칙 W05 「재다」→측정/판정')
rep_notes(8, '계획–실행–점검–개선으로 돕니다', '계획–실행–점검–조치로 돕니다', '규칙 T10 PDCA 풀이')
rep_notes(12, '실시는 향후 과제 — 3.5.', '실시는 향후 연구 과제 — 3.5.', '규칙 T27')
rep_notes(25, 'ML3가 전망됩니다', 'ML3이 전망됩니다', '규칙 T61·T70 약호+3 조사')
rep_notes(27, 'ML3가 전망됩니다', 'ML3이 전망됩니다', '규칙 T61·T70')
prs.save(OUT)
json.dump(LOG, open('log12.json', 'w'), ensure_ascii=False, indent=1)
print(len(LOG), 'changes')
