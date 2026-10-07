"""v19 → v20: 백업 32~61번 30장에 ▶ 길잡이 한 줄(본 슬라이드 4번의 ▶ 줄을 복사해 글만 바꿈). 내용은 ▶ 아래로 내리고 넘치면 그림·표를 줄임."""
import sys, os, json, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu, Pt
from guides20 import G
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC); IN = 914400; LOG = []
hd_src = [sh for sh in prs.slides[3].shapes if sh.has_text_frame and sh.text_frame.text.startswith('▶')][0]
HD_TOP = int(hd_src.top); HD_H = int(hd_src.height); HD_LEFT = int(hd_src.left); HD_W = int(hd_src.width); FOOT = int(7.5 * IN)
for sn, txt in G.items():
    s = prs.slides[sn - 1]
    el = copy.deepcopy(hd_src._element); s.shapes._spTree.append(el)
    hd = s.shapes[-1]
    new_id = max(x.shape_id for x in s.shapes) + 1; el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr').set('id', str(new_id))
    p = hd.text_frame.paragraphs[0]; p.runs[-1].text = txt
    for r in p.runs[1:-1]: r._r.getparent().remove(r._r)
    hd.name = f'Guide {sn}'
    # ▶ 줄 실제 높이: 글자 수로 줄 수 추정(한 줄 약 52자) → 내용 시작 위치
    import unicodedata
    em = sum(1.0 if unicodedata.east_asian_width(c) in 'WF' else 0.55 for c in txt)
    lines = 1 if em <= 50 else 2
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    hd.left = Emu(HD_LEFT); hd.width = Emu(HD_W); hd.top = Emu(HD_TOP); hd.text_frame.vertical_anchor = MSO_ANCHOR.TOP; hd.text_frame.auto_size = MSO_AUTO_SIZE.NONE; hd.text_frame.word_wrap = True
    hd.height = Emu(int(HD_H * (1.0 if lines == 1 else 1.8)))
    CONTENT_TOP = HD_TOP + hd.height + 40000
    content = [sh for sh in s.shapes if sh.shape_id not in (2, 3, 4) and sh is not hd and sh.top < FOOT and not (sh.has_text_frame and sh.text_frame.text.strip().isdigit()) and not (sh.has_text_frame and sh.text_frame.text.startswith('근거'))]
    if not content: continue
    top = min(sh.top for sh in content); bottom = max(sh.top + sh.height for sh in content)
    need_top = CONTENT_TOP; shift = need_top - top
    avail = FOOT - 60000 - need_top; cur_h = bottom - top
    if cur_h + max(0, shift) * 0 <= avail and shift <= 0:
        # 이미 ▶ 아래에 있으면 그대로
        pass
    elif cur_h <= avail:
        for sh in content: sh.top = Emu(sh.top + shift)
    else:
        k = avail / cur_h
        for sh in content:
            rel = sh.top - top
            if sh.shape_type == 13:
                nw = int(sh.width * k); nh = int(sh.height * k); sh.left = Emu(sh.left + (sh.width - nw) // 2); sh.width = Emu(nw); sh.height = Emu(nh)
            elif sh.has_table:
                for r in sh.table.rows: r.height = Emu(int(r.height * k))
            sh.top = Emu(int(need_top + rel * k))
    hd.left = Emu(HD_LEFT); hd.width = Emu(HD_W); hd.top = Emu(HD_TOP)
    LOG.append((sn, '▶ 길잡이', '', txt, f'{"그대로" if shift <= 0 and cur_h <= avail else ("내림" if cur_h <= avail else f"축소 ×{avail/cur_h:.2f}")}'))
prs.save(OUT); json.dump(LOG, open('log20.json', 'w'), ensure_ascii=False, indent=1); print(len(LOG), 'guides')
