# v31 -> v31b: 이사님 10/14 검토 묶음 ①~⑯ (발표자료 부분)
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from PIL import Image
from pptlib import rep, LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def set_note(sn, text, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; old = tf.text; tf.text = text; LOG.append((sn, '발표 노트(전체)', old, text, why))
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text; assert t.count(old) == 1, (sn, old, t.count(old)); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def set_lines(sn, sid, lines, why, size=None, keep_first=False):
    """keep_first=True: 첫 문단은 그대로 두고 둘째 문단부터 교체"""
    sh = shape(prs.slides[sn-1], sid); tf = sh.text_frame; old = tf.text
    base = tf.paragraphs[1] if keep_first and len(tf.paragraphs) > 1 else tf.paragraphs[0]
    pPr = copy.deepcopy(base._p.pPr) if base._p.pPr is not None else None
    r0 = base.runs[0]; rPr = copy.deepcopy(r0._r.rPr) if r0._r.rPr is not None else None
    start = 1 if keep_first else 0
    for p in list(tf.paragraphs)[start:]: p._p.getparent().remove(p._p)
    for ln in lines:
        p = tf.add_paragraph()
        if pPr is not None: p._p.insert(0, copy.deepcopy(pPr))
        r = p.add_run(); r.text = ln
        if rPr is not None: r._r.insert(0, copy.deepcopy(rPr))
        if size: r.font.size = Pt(size)
    LOG.append((sn, f'[{sid}]', old, tf.text, why))
def set_title(sn, text, ref, size, why):
    sh = shape(prs.slides[sn-1], 2); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; runs = p.runs
    runs[0].text = text; runs[0].font.size = Pt(size)
    for r in runs[1:]: r._r.getparent().remove(r._r)
    if ref:
        r = p.add_run(); r.text = '  ' + ref; r.font.size = Pt(13); r.font.bold = False; r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F); r.font.name = '맑은 고딕'
    LOG.append((sn, '제목', old, text + ('  ' + ref if ref else ''), why))
def set_tag(sn, text, width, why):
    sh = shape(prs.slides[sn-1], 3); set_one(sn, 3, text, why); sh.width = Emu(int(width*IN)); sh.left = Emu(int((11.15-width)*IN))
def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1
def crop(src_png, dst_png):
    im = Image.open(src_png).convert('RGB'); bb = im.convert('L').point(lambda v: 0 if v > 242 else 255).getbbox()
    im.crop((max(bb[0]-20, 0), max(bb[1]-20, 0), min(bb[2]+20, im.width), min(bb[3]+20, im.height))).save(dst_png)
def add_pic(s, path, box):
    l, t, r, b = box; W, H = r-l, b-t; im = Image.open(path); a = im.width/im.height
    if W/H > a: h = H; w = h*a
    else: w = W; h = w/a
    return s.shapes.add_picture(path, Emu(int((l+(W-w)/2)*IN)), Emu(int((t+(H-h)/2)*IN)), Emu(int(w*IN)), Emu(int(h*IN)))
def new_fig_slide(title, ref, tag, tag_w, gist, base, fig, box, band, band_top, band_h, band_size, title_size=25):
    tpl = prs.slides[6]; s = prs.slides.add_slide(tpl.slide_layout)   # 원본: 7번 두 상태
    for sh in tpl.shapes:
        if sh.shape_id in (2, 3, 4, 5, 6, 7, 9): s.shapes._spTree.append(copy.deepcopy(sh._element))
    renumber_ids(s); n = len(prs.slides)
    set_title(n, title, ref, title_size, '새 장'); set_tag(n, tag, tag_w, '새 장')
    set_one(n, 5, '▶ ' + gist, '요지'); set_one(n, 7, '근거  ' + base, '근거')
    add_pic(s, fig, box)
    b = shape(s, 8); b.top = Emu(int(band_top*IN)); b.height = Emu(int(band_h*IN))
    set_one(n, 8, band, '설명 띠');
    for r in shape(s, 8).text_frame.paragraphs[0].runs: r.font.size = Pt(band_size)
    return n

# ① 6번 LCA — 4.1·5절·6절
set_lines(6, 9, ['한 번의 수행은 이 안에서 완결되지만, 되풀이하는 관리 주기는 ISO 14044의 범위가 아니다.',
                 '4.1 일반 요구사항은 네 단계 전체에 적용 · 5절 보고(A-7·A-8) · 6절 정밀검토(내부·외부·패널 — P-21·C-5~C-7)도 평가 문항의 근거'], '4.1·5절·6절(이사님)', 14, keep_first=True)
sh = shape(prs.slides[5], 9); sh.top = Emu(int(6.45*IN)); sh.height = Emu(int(1.2*IN))
note_rep(6, '한 번의 평가는 이 안에서 끝납니다.', '네 단계 앞에 일반 요구사항이 있고, 뒤에 보고와 정밀검토가 따로 있습니다. 한 번의 평가는 이 안에서 끝납니다.', '4.1·5·6절')
# ② 7번 두 상태 노트
note_rep(7, '그 정도를 판정하는 것이 성숙도 모델입니다.', '그 정도를 여섯 수준으로 판정하는 것이 성숙도 모델입니다.', '정도 → 수준')
# ③ 9번 성숙도 모델 상자 두 줄
set_lines(9, 15, ['ML0~ML5', '핵심 문항 4종'], '두 줄(이사님)', 10.5, keep_first=True)
# ④ 10번 Q1~Q4 물음표
for sid in (14, 15, 16, 17):
    sh = shape(prs.slides[9], sid); r = sh.text_frame.paragraphs[1].runs[-1]
    if not r.text.endswith('?'): r.text = r.text + '?'
LOG.append((10, '[14~17]', '…쓰는가 / …없는가 / …입력인가 / …주는가', '끝에 물음표', '이사님'))
# ⑤ 11번 세 상자 두 줄
set_lines(11, 15, ['연계 프로세스 16종', '계획 8 · 실행 2 · 점검 3 · 개선 3'], '두 줄(이사님)')
set_lines(11, 19, ['평가 문항 55 · 문항별 판정 기준', '증거 목록'], '두 줄(이사님)')
set_lines(11, 23, ['ML0~ML5 여섯 수준', '핵심 문항 4종'], '두 줄(이사님)')
# ⑥ 15번 A사 라벨 → ML2 줄, 요지·노트 쉽게
sh = shape(prs.slides[14], 25); sh.top = Emu(int(3.63*IN)); sh.left = Emu(int(0.5*IN)); sh.width = Emu(int(0.94*IN)); LOG.append((15, '[25] A사 라벨', 'ML3 줄 옆', 'ML2 줄 옆', '이사님'))
shape(prs.slides[14], 5).text_frame.paragraphs[0].runs[1].text = '꼭 있어야 할 네 문항 가운데 하나라도 3점이 안 되면, 평균이 높아도 ML2에서 멈춘다.'; LOG.append((15, '요지', '핵심 문항 하나라도 3점 미만이면 평균과 관계없이 상한은 ML2다.', '꼭 있어야 할 네 문항 가운데 하나라도 3점이 안 되면, 평균이 높아도 ML2에서 멈춘다.', '쉬운 말'))
set_note(15, '점수를 수준으로 바꾸는 규칙입니다. 수준은 ML0에서 ML5까지 여섯 계단이고, 아래 계단의 조건을 다 갖춰야 다음 계단으로 올라갑니다.\n계단을 오르는 열쇠가 핵심 문항 네 가지입니다. 환경측면 결정, 환경목표, 설계·개발, 모니터링에 LCA 결과가 쓰이는가입니다. 네 가지가 모두 2점이면 ML1, 거기에 나머지 쉰한 문항의 절반이 3점이면 ML2, 네 가지가 모두 3점이면 ML3입니다. ML4와 ML5는 점검과 개선 영역까지 갖춰야 합니다.\n평균은 쓰지 않습니다. 핵심 네 가지 가운데 하나라도 3점이 안 되면 다른 점수가 아무리 높아도 ML2에서 멈춥니다. 꼭 있어야 할 것이 빠졌는데 평균으로 덮지 않는다, 이것이 최소충족 판정이고 데이터 품질 성숙도 표준에서 가져온 원리입니다.\n〔1분 10초. 진입 조건·설계 근거 표 전문은 백업 B-11a, 수준별 기준은 B-11, 판정 트리 [그림 3-9]는 B-9. 비교주장 시 조건부 shall(C-7·A-9)은 3.4.3 — 백업 B-Q5. 연계가 성립하지 않은 것과 성숙도가 낮은 것은 다르다(3.4.1).〕', '쉬운 말(이사님)')
# ⑦⑧⑨ 16번
set_lines(16, 18, ['A사 문서 16종(33건)', '매뉴얼 · 절차서 · 기록 · 보고서'], '접수', keep_first=True)
shape(prs.slides[15], 18).text_frame.paragraphs[0].runs[0].text = '① 문서 접수'; LOG.append((16, '[18]', '① 문서 받기', '① 문서 접수', '이사님'))
shape(prs.slides[15], 22).text_frame.paragraphs[2].runs[0].text = '점수 오른 문항 7 · 내린 문항 9 → 확정 2.85'; LOG.append((16, '[22]', '상향 7 · 하향 9 → 확정 2.85', '점수 오른 문항 7 · 내린 문항 9 → 확정 2.85', '이사님'))
shape(prs.slides[15], 17).text_frame.paragraphs[0].runs[1].text = '평가 클래스 3 = ISO/IEC 33002가 정한 평가 등급 가운데 가장 가벼운 것 — 모델이 작동하는지 보는 연구용 평가  ·  상향·하향 = 잠정 점수를 증빙으로 다시 확인해 점수가 오른·내린 문항 수(55문항 중 16문항 변동, [표 4-4])  ·  D-1 = 핵심 문항 4종의 하나, 설계·개발 단계의 LCA 결과 반영(14·15번)'; LOG.append((16, '용어', '평가 클래스 3 …', '+ 상향·하향, D-1', '이사님'))
set_one(16, 28, '목적은 모델의 판정 논리가 실제 자료에서 작동하는지 확인하는 것', '진행 안내 삭제')
note_rep(16, '첫째, 문서를 받았습니다.', '첫째, 문서를 접수했습니다.', '접수')
note_rep(16, '〔1분 10초.', '〔1분 10초. 2.89→2.85 — 재검증에서 16문항이 ±1점 변동: 오른 7(P-2·P-8·P-12·P-13·P-16·P-18 3→4, C-2 1→2), 내린 9(P-3 3→2 예고 역방향, P-10 4→3, P-21 2→1, P-23·D-5·C-4·C-8·C-9·A-6 3→2), 합계 −2점 ÷ 55 = −0.04. 핵심 문항 4종은 변동 없음. 재검증 전 분류 [표 4-3](확정 15·확인 16·상승 가능 7·하락 가능 11·핵심 4·미결 2), 내역 [표 4-4]. 확정치가 잠정치보다 낮게 수렴 = 재검증이 후하게 주는 쪽을 억제(비관대성, 4.2.3).', '상향·하향 내역')
# ⑪ 17번 = [그림 4-1] 새 장
crop('figs29/fig4-1.png', 'figs29/fig4-1c.png')
n17 = new_fig_slide('영역별 평점과 핵심 문항의 판정', '[그림 4-1]', 'Ⅳ. 성숙도 평가 — 적용하면 어떻게 되나', 4.07,
    '판정은 점수의 크기가 아니라 핵심 문항의 충족 여부로 정해진다 — D-1 하나가 2점이라 ML2.', '4.3, 4.3.5, [그림 4-1], [그림 4-2]', 'figs29/fig4-1c.png', (0.54, 1.35, 11.14, 6.5),
    '판정 경로 — ML1 진입: 핵심 4종 모두 2점 이상(4·3·2·3) 통과 → ML2 진입: 나머지 51문항 중 36문항(70.6%)이 3점 이상 통과 → ML3 진입: 핵심 4종 모두 3점 이상이어야 하는데 D-1이 2점 → 미충족 → 판정 ML2(관리화). 종합점수 2.85는 판정에 쓰지 않는 보조값', 6.6, 0.65, 13)
set_note(n17, 'A사를 이 계단에 올려 보겠습니다. 핵심 문항 네 가지 가운데 셋은 통과했고, 설계·개발의 D-1 하나가 2점으로 모자랍니다.\n아래부터 묻습니다. 네 가지 모두 2점 이상인가, 예, ML1 통과입니다. 나머지 쉰한 문항의 절반 이상이 3점인가, 서른여섯 문항이라 예, ML2 통과입니다. 네 가지 모두 3점 이상인가, D-1이 2점이라 아니오. 그래서 ML2, 관리화입니다.\n네 칸이 계획·실행·점검·개선의 영역 평점입니다. LCA가 계획과 실행에는 들어와 있지만 점검과 개선으로는 이어지지 않는 모습이고, 낮은 점수의 원인은 설계 절차, 정밀검토 체계, 재수행 세 곳입니다. 강점은 LCA가 문서 체계에 들어와 있다는 것입니다.\n종합점수 2.85는 참고값일 뿐 판정에는 쓰지 않습니다. 기준을 더 엄하게 바꿔 봐도 판정은 같습니다. 점수가 얼마나 높은가가 아니라 꼭 있어야 할 네 가지가 있는가로 정해지기 때문입니다.\n〔1분 20초. 영역 평점 계획 3.23·실행 3.00·점검 2.10·개선 2.33, 전문은 백업 S-22. 판정 경로 상자 장은 S-23, 연계 상태 [그림 4-4]는 S-24, 개선 전·후 그래프 [그림 4-3]은 B-13a. 「너무 가혹하지 않은가」 → B-Q1. 민감도 — 과반 기준 50→70%, 핵심 기준 4·5점으로 올려도 판정 ML2 그대로(4.3.5).〕', '이사님 10/14 — [그림 4-1]')
# 옛 17·18 → 백업
set_tag(17, '백업  S-23 (본심사판 23번)', 3.3, '본 → 백업'); set_tag(18, '백업  S-24 (본심사판 24번)', 3.3, '본 → 백업')
# ⑮ 20번 노트
note_rep(20, '가장 높은 두 수준의 사례 미확인이고, 이 셋이 그대로 향후 연구 과제입니다.', '가장 높은 두 수준은 A사가 ML2라 재수행 같은 상위 수준의 기록이 없어 사례로 확인하지 못한 것이고, 이 셋이 그대로 향후 연구 과제입니다. 규칙이 없는 것이 아닙니다.', '이사님 — 기록이 없어서')
# ⑯ 21번 결론 = [그림 5-1] 새 장
crop('figs29/fig5-1.png', 'figs29/fig5-1c.png')
n21 = new_fig_slide('결론 — 네 가지 연구 문제의 답', '[그림 5-1]', '결론 — 답한 것 · 기여 · 남은 것', 3.32,
    '두 표준의 조항을 연계하는 LMA를 만들고 한 기업에 적용하여 작동함을 확인하였다.', '5.1, 5.3, [그림 5-1]', 'figs29/fig5-1c.png', (0.54, 1.35, 11.14, 6.55),
    '두 표준의 조항을 연계하여 LMA를 만들고, 이 모델로 조직을 평가하여, 그 결과를 환경경영시스템(ISO 14001)의 지속적 개선으로 연결함으로써 전과정의 환경영향을 줄이는 데 기여하고자 한다.', 6.65, 0.6, 14)
set_note(n21, '네 가지 연구 문제의 답입니다.\n첫째, 두 표준의 조항은 연계할 수 있습니다. 관리 주기와 평가 절차가 서로 맞물리는 자리가 있고, 한쪽은 정량 방법을, 다른 쪽은 되풀이하는 주기를 비워 두었기 때문입니다.\n둘째, 그 자리를 업무로 묶으면 연계 프로세스 열여섯 가지이고, 되어 있는지 묻는 문항이 쉰다섯 개입니다.\n셋째, 성숙도는 여섯 수준이고, 핵심 문항 네 가지가 모두 갖춰졌는지로 판정하며 평균은 쓰지 않습니다.\n넷째, A사에 적용해 판정과 개선 계획이 나오는 것을 확인했고, 그 결과가 다음 주기의 입력이 됩니다.\n두 표준의 조항을 연계하여 LMA를 만들고, 이 모델로 조직을 평가하여, 그 결과를 환경경영시스템(ISO 14001)의 지속적 개선으로 연결함으로써 전과정의 환경영향을 줄이는 데 기여하고자 합니다. 이상입니다.\n〔50초. 결론 글 상자 장은 백업 S-29, 연구 전체 한 장 요약은 S-28, 기여 [그림 5-2]는 B-0c②.〕', '이사님 10/14 — [그림 5-1]')
set_tag(21, '백업  S-29 (본심사판 29번)', 3.3, '본 → 백업')
# 구분·감사
rep(prs, 23, 2, 'S-7~S-28 본심사판에서 옮긴 9장(연구 목적과 문제 · ISO 14001 구조 · 선행연구 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 로드맵 · 연구 전체 요약)',
    'S-7~S-29 본심사판에서 옮긴 12장(연구 목적과 문제 · ISO 14001 구조 · 선행연구 · 모델 한 장 요약 · 개발 경위 · 객관성 · 영역별 결과 · 판정 경로 · 연계 상태 · 로드맵 · 연구 전체 요약 · 결론)', '백업 목록')
note_rep(22, 'S-7~S-28(본심사판에서 옮긴 9장)', 'S-7~S-29(본심사판에서 옮긴 12장)', '백업 수')
# ---- 순서
N = len(prs.slides)
MAIN = [1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16, n17, 19, 20, n21, 22]
SBLK = [24,25,26,27,28,29,30, 17, 18, 31, 32, 21]
order = MAIN + [23] + SBLK + list(range(33, N-1))   # 새 장 둘(N-1, N)은 이미 MAIN에
assert sorted(order) == list(range(1, N+1)), (len(order), N, sorted(set(range(1,N+1))-set(order)))
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, n in enumerate(nums): out += str(old2new.get(int(n), int(n))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = s.notes_slide.notes_text_frame; t = tf.text; t2 = tr_nums(t)
    if t2 != t: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new, 'main': MAIN}, open('log31b.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), 'slides', len(prs.slides), [old2new[o] for o in MAIN])

# (보정) 6번 추가 줄은 빨간 강조가 아니라 본문색
_p = Presentation(dst); _sh = shape(_p.slides[5], 9); _par = _sh.text_frame.paragraphs[-1]
for _r in _par.runs: _r.font.bold = False; _r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
_p.save(dst); print('6번 줄 색 보정')
