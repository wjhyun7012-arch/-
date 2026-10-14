# v31b -> v32: 감사 뒤 「질의응답」 장 추가
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptlib import LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1
tpl = prs.slides[6]; s = prs.slides.add_slide(tpl.slide_layout)
for sh in tpl.shapes:
    if sh.shape_id in (2, 4, 6, 9, 9): s.shapes._spTree.append(copy.deepcopy(sh._element))
# 본문 상자 하나 더 복사
s.shapes._spTree.append(copy.deepcopy(shape(tpl, 9)._element)); renumber_ids(s)
shs = {sh.shape_id: sh for sh in s.shapes}
t = shs[2]; p = t.text_frame.paragraphs[0]; p.runs[0].text = '질의응답'
for r in p.runs[1:]: r._r.getparent().remove(r._r)
def fill(sh, text, size, bold, color, top, h):
    tf = sh.text_frame; p0 = tf.paragraphs[0]
    for q in list(tf.paragraphs)[1:]: q._p.getparent().remove(q._p)
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    r = p0.runs[0]; r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    p0.alignment = PP_ALIGN.CENTER; sh.top = Emu(int(top*IN)); sh.height = Emu(int(h*IN)); sh.left = Emu(int(0.54*IN)); sh.width = Emu(int(10.6*IN))
boxes = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith('상태 A')]
fill(boxes[0], '환경경영시스템(ISO 14001)과 전과정평가(LCA)의 프로세스 연계 모델 개발 및 성숙도 평가에 대한 연구', 20, True, '1F3864', 3.0, 0.9)
fill(boxes[1], '연계 프로세스 16종 · 평가 문항 55 · 성숙도 수준 ML0~ML5 · 핵심 문항 4종   │   A사 종합 2.85 → ML2(관리화), 개선 과제 5건 완수 시 3.15 → ML3(체계화) 전망', 15, False, '111111', 4.1, 0.8)
n = len(prs.slides)
# 노트: 감사(21)의 질의응답 지시문을 새 장으로
tf21 = prs.slides[20].notes_slide.notes_text_frame; qa = tf21.text
s.notes_slide.notes_text_frame.text = qa
tf21.text = '〔인사. 바로 다음 질의응답 장을 띄운 채로 질문을 받는다.〕'
LOG.append((n, '새 장', '(없음)', '질의응답 — 연구 제목 + 핵심 숫자 한 줄, 질의응답 지시문 노트', '이사님 10/14'))
# 순서: 21 감사 뒤에
N = len(prs.slides); order = list(range(1, 22)) + [N] + list(range(22, N))
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, v in enumerate(nums): out += str(old2new.get(int(v), int(v))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = sl.notes_slide.notes_text_frame; tx = tf.text; t2 = tr_nums(tx)
    if t2 != tx: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new}, open('log32.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(prs.slides))
