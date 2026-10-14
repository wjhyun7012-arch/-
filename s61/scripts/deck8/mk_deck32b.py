# v32 -> v32b: ① 6번 첫 줄 15pt ② 9번 BARS ③ 10번 ↔ B-1([그림 3-3]) ④ 연계 지점 맵을 매핑 다음으로 ⑤ 16번 상향·하향 ⑥ 감사 장에 Q & A, 22번 삭제
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptlib import rep, LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text; assert t.count(old) == 1, (sn, old, t.count(old)); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1
# ① 6번 첫 줄 15pt
for r in shape(prs.slides[5], 9).text_frame.paragraphs[0].runs: r.font.size = Pt(15)
LOG.append((6, '[9] 첫 줄', '16pt', '15pt', '「다.」 줄 넘김(이사님)'))
# ② 9번 BARS
rep(prs, 9, 13, '본 모델은 그 틀을 5점 상태 기술(BARS)로 구체화', '본 모델은 그 틀을 행동기술척도(BARS)의 5점 상태 기술로 구체화', '논문 용어(이사님)')
sh18 = shape(prs.slides[8], 18); r = sh18.text_frame.paragraphs[0].runs[-1]; r.text = r.text + '  ·  행동기술척도(BARS: Behaviorally Anchored Rating Scales) = 점수마다 「어떤 상태면 몇 점」을 글로 적어 둔 척도 — 본 모델은 체크포인트 충족 개수로 1~5점'
LOG.append((9, '용어', '…LMA)', '+ 행동기술척도(BARS) 풀이', '이사님'))
# ③ B-1(36) → 본 10번으로 손질
b1 = prs.slides[35]
t = shape(b1, 2); p = t.text_frame.paragraphs[0]; p.runs[0].text = 'ISO 14001 조항의 연계 가능성 판정 질문 적용 흐름'; p.runs[0].font.size = Pt(21)
for r in p.runs[1:]: r._r.getparent().remove(r._r)
r = p.add_run(); r.text = '  [그림 3-3]'; r.font.size = Pt(13); r.font.bold = False; r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F); r.font.name = '맑은 고딕'
tag = shape(b1, 3); set_one(36, 3, 'Ⅲ. 모델링 — 무엇을 만들었나', 'B-1 → 본 10번'); tag.width = Emu(int(3.08*IN)); tag.left = Emu(int((11.15-3.08)*IN)); tag.top = Emu(int(0.12*IN)); tag.height = Emu(int(0.36*IN))
for r in tag.text_frame.paragraphs[0].runs: r.font.size = Pt(14)
g = shape(b1, 7); g.height = Emu(int(0.46*IN)); set_one(36, 7, '▶ 조항마다 네 물음을 같은 순서로 물어 네 등급 가운데 하나로 매긴다 — 「예」가 앞에서 나올수록 연계 가능성이 높다.', '요지')
for r in g.text_frame.paragraphs[0].runs: r.font.size = Pt(15)
pic = shape(b1, 6); pic.top = Emu(int(1.35*IN)); pic.height = Emu(int(4.9*IN)); pic.left = Emu(int(0.54*IN)); pic.width = Emu(int(10.6*IN))
# 근거 줄·아래 줄은 7번(두 상태)에서 복사
tpl = prs.slides[6]
for sid in (7, 9): b1.shapes._spTree.append(copy.deepcopy(shape(tpl, sid)._element))
renumber_ids(b1)
for sh in b1.shapes:
    if not sh.has_text_frame: continue
    tx = sh.text_frame.text
    if tx.startswith('근거'): set_one(36, sh.shape_id, '근거  3.2.2, 3.3.1, [표 3-2], [그림 3-3], 부록 A.1·A.2·A.4', '근거')
    if tx.startswith('상태 A'):
        set_one(36, sh.shape_id, 'ISO 14001 — 요구사항 조항 30개 전부 판정·문항화   ·   ISO 14044 — 조항 25개가 문항의 근거(나머지 3개도 세부 조항의 문항·체크포인트로 반영) → 빠진 조항 없음(부록 A.4)   ·   매우 높음 4개 조항 = 핵심 문항 4종(P-9·P-14·D-1·C-1) — 계획 둘·실행 하나·점검 하나', '30·25·핵심 4')
        sh.top = Emu(int(6.4*IN)); sh.height = Emu(int(0.8*IN))
        for r in sh.text_frame.paragraphs[0].runs: r.font.size = Pt(13)
old10 = prs.slides[9].notes_slide.notes_text_frame.text
b1.notes_slide.notes_text_frame.text = old10.replace('서른 개 조항 하나하나에 같은 네 가지 질문을 같은 순서로 물어 네 가지 등급으로 판정했습니다.', '서른 개 조항 하나하나에 같은 네 가지 질문을 같은 순서로 물어 네 가지 등급으로 판정했습니다. 네 질문은 그 조항을 지키는 데 LCA가 얼마나 필요한가를 센 것부터 차례로 묻는 것입니다. 전과정 관점을 직접 요구하는가, 없으면 의무를 못 지키는가, 주요 입력인가, 보조 정보라도 되는가이고, 예가 앞에서 나올수록 등급이 높습니다.').replace('등급 요건은 B-1a, 흐름 B-1.', '등급 요건은 B-1a, 네 질문 상자 장은 백업 B-1. 「왜 이 네 질문인가」 → [표 3-2]의 네 등급 정의를 판정 가능한 물음으로 바꾼 것(3.2.2 B551).')
LOG.append((36, '발표 노트(전체)', '(B-1 지시문)', '매핑 대본 + 네 질문의 뜻', '이사님 10/14'))
# 옛 10번 → 백업 B-1
o10 = prs.slides[9]; tg = shape(o10, 3); set_one(10, 3, '백업  B-1', '본 → 백업'); tg.width = Emu(int(1.6*IN)); tg.left = Emu(int((11.15-1.6)*IN))
o10.notes_slide.notes_text_frame.text = '네 가지 판정 질문을 상자로 보여 달라는 경우. 흐름은 본 10번 [그림 3-3], 등급 요건은 B-1a. 3.2.2, [표 3-2].'
# ④ LMA 구성(11) 위 글
set_one(11, 11, '조항 매핑의 결과', '연계 지점 맵이 앞에 옴')
# ⑤ 16번 상향·하향
shape(prs.slides[15], 22).text_frame.paragraphs[2].runs[0].text = '상향 7 · 하향 9 → 확정 2.85'; LOG.append((16, '[22]', '점수 오른 문항 7 · 내린 문항 9', '상향 7 · 하향 9', '논문 용어(이사님)'))
# ⑥ 감사 장에 Q & A, 22 삭제
s21 = prs.slides[20]; s21.shapes._spTree.append(copy.deepcopy(shape(s21, 2)._element)); renumber_ids(s21)
qa = [sh for sh in s21.shapes if sh.has_text_frame and sh.text_frame.text.strip() == '감사합니다'][-1]
qa.text_frame.paragraphs[0].runs[0].text = 'Q & A'; qa.text_frame.paragraphs[0].runs[0].font.size = Pt(28); qa.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
qa.top = Emu(int(4.6*IN)); qa.height = Emu(int(1.0*IN))
s21.notes_slide.notes_text_frame.text = prs.slides[21].notes_slide.notes_text_frame.text
LOG.append((21, '감사 장', '감사합니다', '감사합니다 + Q & A, 질의응답 지시문 노트', '이사님 10/14'))
# ---- 순서: 1..9, 36(B-1→10), 12(지점 맵), 11(LMA), 13..21, [22 삭제], 23.., 10(옛 매핑→B-1 자리)
N = len(prs.slides)
order = list(range(1, 10)) + [36, 12, 11] + list(range(13, 22)) + list(range(23, 36)) + [10] + list(range(37, N+1))
DROP = [22]
assert sorted(order + DROP) == list(range(1, N+1)), (len(order), N)
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for d in DROP: prs.part.drop_rel(ids[d-1].rId)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, v in enumerate(nums): out += str(old2new.get(int(v), int(v))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = sl.notes_slide.notes_text_frame; tx = tf.text; t2 = tr_nums(tx)
    if t2 != tx: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new}, open('log32b.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), len(prs.slides))

# (보정) 10번 장 표시를 다른 장과 같은 상자 모양으로 (7번 장 표시 복사)
_p = Presentation(dst); _s = _p.slides[9]
_old = shape(_s, 3); _old._element.getparent().remove(_old._element)
_s.shapes._spTree.append(copy.deepcopy(shape(_p.slides[6], 3)._element)); renumber_ids(_s)
_tag = [x for x in _s.shapes if x.has_text_frame and x.text_frame.text.startswith('Ⅰ')][0]
_tag.text_frame.paragraphs[0].runs[0].text = 'Ⅲ. 모델링 — 무엇을 만들었나'
for _r in _tag.text_frame.paragraphs[0].runs[1:]: _r._r.getparent().remove(_r._r)
_tag.width = Emu(int(3.08*IN)); _tag.left = Emu(int((11.15-3.08)*IN))
_p.save(dst); print('10번 장 표시 보정')
