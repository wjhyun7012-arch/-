"""v8 → v9: 본 슬라이드 1~29 발표 노트를 존댓말 대본(논문 문장 기준)으로 교체. 화면 글은 손대지 않음.
사용: python3 -I mk_deck9.py v8.pptx v9.pptx  → log9.json(전·후)
"""
import sys, json, copy, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Pt
from notes9 import N
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
LOG = []
for i, s in enumerate(prs.slides, 1):
    if i > 29: break
    ns = s.notes_slide; tf = ns.notes_text_frame
    old = tf.text
    new = N[i]
    # 첫 문단의 run 서식을 유지해 문단 단위로 다시 씀
    p0 = tf.paragraphs[0]
    tmpl = copy.deepcopy(p0._p)
    for p in list(tf.paragraphs): p._p.getparent().remove(p._p)
    for line in new.split('\n'):
        np_ = copy.deepcopy(tmpl); tf._txBody.append(np_)
        from pptx.text.text import _Paragraph
        para = _Paragraph(np_, tf)
        rs = para.runs
        if rs:
            rs[0].text = line
            for r in rs[1:]: r._r.getparent().remove(r._r)
        else:
            para.text = line
        for r in para.runs:
            r.font.name = '맑은 고딕'; r.font.size = Pt(12)
    LOG.append((i, old, new))
prs.save(OUT)
json.dump(LOG, open(os.path.join(os.path.dirname(OUT) or '.', 'log9.json'), 'w'), ensure_ascii=False, indent=1)
print(len(LOG), 'notes replaced')
