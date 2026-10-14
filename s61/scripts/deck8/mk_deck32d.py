# v32c -> v32 최종: ⑪ 표지 영문 제목 두 줄(between부터 아랫줄)
import sys, copy, json
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptlib import LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
sh = shape(prs.slides[0], 8); tf = sh.text_frame; p0 = tf.paragraphs[0]
full = ''.join(r.text for r in p0.runs); a, b = full.split(' between ')
r0 = p0.runs[0]; r0.text = a
for r in p0.runs[1:]: r._r.getparent().remove(r._r)
p1 = tf.add_paragraph()
if p0._p.pPr is not None: p1._p.insert(0, copy.deepcopy(p0._p.pPr))
r1 = p1.add_run(); r1.text = 'between ' + b
if r0._r.rPr is not None: r1._r.insert(0, copy.deepcopy(r0._r.rPr))
sh.height = Emu(int(0.7*IN)); sh.width = Emu(int(6.7*IN))
LOG.append((1, '[8] 영문 제목', full, a + '\n' + 'between ' + b, 'between부터 아랫줄(이사님)'))
prs.save(dst); json.dump(LOG, open('log32d.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1); print('ok')
