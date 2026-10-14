# v30 -> v31: 개요 뒤에 두 표준 설명 두 장([그림 2-2] 새 장, S-9 복귀) → 두 상태 → 빈 곳 → 따른 표준(S-10 복귀) → 매핑. 본 22장.
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from PIL import Image
from pptlib import rep, LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def set_note(sn, text, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; old = tf.text; tf.text = text; LOG.append((sn, '발표 노트(전체)', old, text, why))
def note_rep(sn, old, new, why):
    tf = prs.slides[sn-1].notes_slide.notes_text_frame; t = tf.text; assert t.count(old) == 1, (sn, old, t.count(old)); tf.text = t.replace(old, new); LOG.append((sn, '발표 노트', old, new, why))
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def set_title(sn, text, ref, size, why):
    sh = shape(prs.slides[sn-1], 2); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; runs = p.runs
    runs[0].text = text; runs[0].font.size = Pt(size)
    for r in runs[1:]: r._r.getparent().remove(r._r)
    if ref:
        r = p.add_run(); r.text = '  ' + ref; r.font.size = Pt(13); r.font.bold = False; r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F); r.font.name = '맑은 고딕'
    LOG.append((sn, '제목', old, text + ('  ' + ref if ref else ''), why))
def set_tag(sn, text, width, why):
    sh = shape(prs.slides[sn-1], 3); set_one(sn, 3, text, why); sh.width = Emu(int(width*IN)); sh.left = Emu(int((11.15-width)*IN))
def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1
def add_pic(s, path, box):
    l, t, r, b = box; W, H = r-l, b-t; im = Image.open(path); a = im.width/im.height
    if W/H > a: h = H; w = h*a
    else: w = W; h = w/a
    return s.shapes.add_picture(path, Emu(int((l+(W-w)/2)*IN)), Emu(int((t+(H-h)/2)*IN)), Emu(int(w*IN)), Emu(int(h*IN)))

# 그림 2-2 여백 자르기
im = Image.open('figs29/fig2-2.png').convert('RGB'); bb = im.convert('L').point(lambda v: 0 if v > 242 else 255).getbbox()
im.crop((max(bb[0]-20, 0), max(bb[1]-20, 0), min(bb[2]+20, im.width), min(bb[3]+20, im.height))).save('figs29/fig2-2c.png')

# ---- 새 장: ISO 14001은 어떤 표준인가 [그림 2-2] (원본: 5번 두 상태의 상자 복사)
TPL = 5
tpl = prs.slides[TPL-1]; s = prs.slides.add_slide(tpl.slide_layout)
for sh in tpl.shapes:
    if sh.shape_id in (2, 3, 4, 5, 6, 7, 9): s.shapes._spTree.append(copy.deepcopy(sh._element))
renumber_ids(s); NEW = len(prs.slides)
set_title(NEW, '환경경영시스템(ISO 14001)은 어떤 표준인가', '[그림 2-2]', 23, '새 장')
set_one(NEW, 5, '▶ ISO 14001은 공통 구조 위의 환경 모듈, 전과정평가(LCA)는 구조 밖의 별개 표준 — 그 사이가 연계다.', '요지')
set_one(NEW, 7, '근거  2.1.2, [그림 2-2]', '근거')
add_pic(s, 'figs29/fig2-2c.png', (0.54, 1.35, 11.14, 6.55))
band = shape(s, 8); band.top = Emu(int(6.62*IN)); band.height = Emu(int(0.6*IN))
set_one(NEW, 8, '왼쪽 공통 구조(4~10절)는 품질·안전보건·환경 등 모든 경영시스템 표준이 같이 쓴다 → ISO 14001은 그 위의 환경 모듈 — 전과정 관점의 고려(6.1.2)·일관성(8.1)을 요구하되 방법은 정하지 않음 → 방법인 전과정평가(ISO 14040, ISO 14044)는 구조 밖 별개 표준 → 그 사이의 연계 프로세스가 본 연구의 대상', '설명 띠')
set_note(NEW, '두 표준이 각각 무엇인지부터 보겠습니다. 왼쪽은 품질, 안전보건, 환경 등 모든 경영시스템 표준이 같이 쓰는 공통 구조로, 조직 상황부터 개선까지 4절에서 10절이고 PDCA 순환입니다.\n오른쪽이 그 위에 얹는 분야별 요구사항이고, 환경 모듈인 ISO 14001에 전과정 관점을 고려하라는 요구가 있습니다.\n그런데 어떻게 고려하는지의 방법인 전과정평가(LCA)는 ISO 14040과 ISO 14044라는 별개 표준에 있어 경영시스템 구조 밖에 놓입니다. 이 둘 사이의 연계 프로세스가 본 연구의 대상입니다.\n〔50초. 2.1.2(B355). 공통 구조는 ISO 14001:2015가 채택한 상위 문서 구조. ISO 14001 절별 구조·PDCA는 백업 S-8 [그림 2-3].〕', '이사님 10/14 — 표준 설명 장')
# ---- S-9 → 본: LCA는 어떤 표준인가 [그림 2-4]
set_title(23, '전과정평가(LCA, ISO 14044)는 어떤 표준인가', '[그림 2-4]', 23, 'S-9 복귀')
set_tag(23, 'Ⅰ. 개요 — 왜 이 연구인가', 2.72, 'S-9 복귀')
set_one(23, 7, '근거  2.2.2, [그림 2-4]', '근거')
set_note(23, 'LCA는 네 단계입니다. 왼쪽 위에서 아래로, 무엇을 얼마나 볼지 정하고, 들어가고 나오는 것을 세고, 그것이 환경에 미치는 영향을 계산합니다. 오른쪽 해석은 세 단계 모두와 오가며 결론을 냅니다.\n한 번의 평가는 이 안에서 끝납니다. 그런데 「언제 다시 하는가」는 이 표준에 없습니다. 이 반복 구조가 PDCA 순환과 같은 꼴이라는 점이 뒤의 연계 프로세스 설계의 전제입니다.\n〔50초. 2.2.2(B392). 단계별 조항 — 목적 및 범위 정의 4.2, 목록분석 4.3, 영향평가 4.4, 해석 4.5. 「왜 ISO 14044인가(14040 아님)」 → 전문 4번.〕', 'S-9 노트 — 성숙도 문장 빼고 설계 전제 한 문장')
# ---- S-10 → 본: 따른 표준
set_tag(24, 'Ⅱ. 학술적 배경 — 무엇을 바탕으로', 3.3, 'S-10 복귀')
for sid in (9, 12, 15):
    sh = shape(prs.slides[23], sid); p2 = sh.text_frame.paragraphs[1]
    for r in p2.runs: r.font.size = Pt(10.5)
LOG.append((24, '[9·12·15] 아래 줄', '13pt', '10.5pt', '한 줄에 들어가도록(이사님)'))
note_rep(24, '그러면 이 대상을 다룬 연구는 이미 있었는지, 다음 장에서 보겠습니다.\n', '', '진행 안내 삭제')
# ---- 4 개요: 두 표준 설명 두 문장 축약
note_rep(4, '왼쪽 환경경영시스템(ISO 14001)은 회사가 환경을 관리하는 틀입니다. 환경측면을 결정하고, 그 영향을 평가해 중대한 환경측면을 결정하고, 그것을 환경목표에 반영해 실행하고, 점검하여 지속적으로 개선합니다.\n오른쪽 전과정평가(LCA)는 제품이 원료 채취, 설계, 생산, 운송과 배송, 사용, 폐기의 각 단계에서 환경에 주는 영향을 숫자로 계산해 영향이 가장 큰 단계, 이것을 핫스팟이라 하는데, 그것을 찾는 방법입니다.\n',
         '왼쪽 환경경영시스템(ISO 14001)은 회사가 환경을 관리하는 틀이고, 오른쪽 전과정평가(LCA)는 제품이 원료 채취부터 폐기까지 각 단계에서 환경에 주는 영향을 숫자로 계산해 영향이 가장 큰 단계, 핫스팟을 찾는 방법입니다.\n', '5·6번에서 자세히')
note_rep(4, '〔1분.', '〔50초.', '시간')
# ---- 6 빈 곳: 앞 두 장과 겹치는 부분 정리
note_rep(6, '2장은 두 표준과 그 사이입니다. ISO 14001은 환경측면의 결정에서 전과정 관점을 고려하도록 요구하되 정량적 평가 방법을 정하지 않았고, 전과정평가(LCA) 표준인 ISO 14044는 계산 방법을 정하되 그 결과를 되풀이하여 관리하는 주기는 정하지 않았습니다.',
         '2장은 두 표준과 그 사이입니다. 앞서 보신 대로 ISO 14001은 전과정 관점을 요구하되 정량적 평가 방법을 정하지 않았고, ISO 14044는 계산 방법을 정하되 그 결과를 되풀이하여 관리하는 주기는 정하지 않았습니다.', '5·6번과 겹침 정리')
note_rep(6, '성숙도를 평가하는 방법은 새로 만들지 않고 소프트웨어·데이터 분야에서 쓰던 국제표준에서 가져왔고, 두 표준의 연계 정도를 조항 단위로 판정한 선행연구는 찾아보기 어려웠습니다. 그 빈 자리가 이 연구의 위치입니다.',
         '두 표준의 연계 정도를 조항 단위로 판정한 선행연구는 찾아보기 어려웠고, 그 빈 자리가 이 연구의 위치입니다.', '방법 출처는 9번(따른 표준)에서')
note_rep(6, '〔1분 10초. 가져온 것', '〔1분. 가져온 것', '시간')
# ---- 구분·감사
rep(prs, 20, 2, 'S-7~S-28 본심사판에서 옮긴 11장(연구 목적과 문제 · ISO 14001 구조 · LCA 4단계 · 따른 표준 · 선행연구', 'S-7~S-28 본심사판에서 옮긴 9장(연구 목적과 문제 · ISO 14001 구조 · 선행연구', '백업 목록')
note_rep(19, 'S-7~S-28(본심사판에서 옮긴 11장)', 'S-7~S-28(본심사판에서 옮긴 9장)', '백업 수')
# ---- 순서
N = len(prs.slides)
MAIN = [1, 2, 3, 4, NEW, 23, 5, 6, 24, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19]
order = MAIN + [20] + [i for i in range(21, N) if i not in (23, 24)]
assert sorted(order) == list(range(1, N+1)), (len(order), N)
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(m):
        nums = re.findall(r'\d+', m.group(0)); seps = re.findall(r'[·~,]\s?', m.group(0))
        out = m.group(0)[:m.group(0).index(nums[0])]
        for k, n in enumerate(nums): out += str(old2new.get(int(n), int(n))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = s.notes_slide.notes_text_frame; t = tf.text; t2 = tr_nums(t)
    if t2 != t: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new, 'main': MAIN}, open('log31.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), 'slides', len(prs.slides), [old2new[o] for o in MAIN])

# ---- (추가) 4번 핵심 줄: 「이것이 이 연구가 만든 것이다」를 아랫줄로 (이사님 10/14)
prs2 = Presentation(dst)
sh = shape(prs2.slides[3], 25); tf = sh.text_frame; p0 = tf.paragraphs[0]; r0 = p0.runs[0]
full = r0.text; head, tail = full.rsplit(', 이것이', 1)
r0.text = head
p1 = tf.add_paragraph(); import copy as _c
if p0._p.pPr is not None: p1._p.insert(0, _c.deepcopy(p0._p.pPr))
r1 = p1.add_run(); r1.text = '이것이' + tail
if r0._r.rPr is not None: r1._r.insert(0, _c.deepcopy(r0._r.rPr))
sh.height = Emu(int(1.0*IN))
prs2.save(dst); print('4번 핵심 줄 두 줄로')

# ---- (추가) 5번에 상위 문서 구조(HLS) 표기 (이사님 10/14)
prs3 = Presentation(dst); s5 = prs3.slides[4]
def _set(sid, text):
    sh = shape(s5, sid); p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
_set(5, '▶ ISO 14001은 상위 문서 구조(HLS) 위의 환경 모듈, 전과정평가(LCA)는 구조 밖의 별개 표준 — 그 사이가 연계다.')
_set(8, '왼쪽 공통 구조(4~10절) = 상위 문서 구조(HLS: High Level Structure) — 품질·안전보건·환경 등 모든 경영시스템 표준이 같이 쓴다 → ISO 14001은 그 위의 환경 모듈 — 전과정 관점의 고려(6.1.2)·일관성(8.1)을 요구하되 방법은 정하지 않음 → 방법인 전과정평가(ISO 14040, ISO 14044)는 구조 밖 별개 표준 → 그 사이의 연계 프로세스가 본 연구의 대상')
tf = s5.notes_slide.notes_text_frame; t = tf.text
old = '왼쪽은 품질, 안전보건, 환경 등 모든 경영시스템 표준이 같이 쓰는 공통 구조로, 조직 상황부터 개선까지 4절에서 10절이고 PDCA 순환입니다.'
assert old in t
tf.text = t.replace(old, '왼쪽은 품질, 안전보건, 환경 등 모든 경영시스템 표준이 같이 쓰는 공통 구조로, 상위 문서 구조, 줄여서 HLS라고 합니다. 조직 상황부터 개선까지 4절에서 10절이고 PDCA 순환입니다.')
prs3.save(dst); print('5번 HLS 표기')
