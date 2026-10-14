# v32b -> v32c: ⑦ 6번 ISO 14044 정식 명칭 ⑧ 5번 ISO 14001 정식 명칭 ⑨⑩ 매핑 장 복귀(10 ↔ B-1), 노트 세 문장, 백업 B-1b 조항별 경로 표
import sys, json, re, copy
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptlib import LOG, shape
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); IN = 914400
def set_one(sn, sid, text, why):
    sh = shape(prs.slides[sn-1], sid); old = sh.text_frame.text; p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for q in list(sh.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
    LOG.append((sn, f'[{sid}]', old, text, why))
def renumber_ids(s):
    i = 2
    for el in s.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        nv = el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        if nv is None: nv = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if nv is not None: nv.set('id', str(i)); i += 1
def add_line(sn, text, top, h, size=11.5, color='595959'):
    """7번 근거 상자(회색 작은 글)를 복사해 한 줄 추가"""
    s = prs.slides[sn-1]; s.shapes._spTree.append(copy.deepcopy(shape(prs.slides[6], 7)._element)); renumber_ids(s)
    sh = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.startswith('근거')][-1]
    p = sh.text_frame.paragraphs[0]; p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].font.size = Pt(size); p.runs[0].font.color.rgb = RGBColor.from_string(color); p.alignment = PP_ALIGN.LEFT
    sh.left = Emu(int(0.54*IN)); sh.top = Emu(int(top*IN)); sh.width = Emu(int(10.6*IN)); sh.height = Emu(int(h*IN))
    LOG.append((sn, '정식 명칭 줄', '(없음)', text, '이사님 10/14')); return sh
def shrink_pic(sn, sid, new_top, bottom):
    pic = shape(prs.slides[sn-1], sid); a = pic.width / pic.height; H = bottom - new_top; w = H * a
    if w > 10.6*IN: w = 10.6*IN; H = w / a
    pic.top = Emu(int(new_top)); pic.height = Emu(int(H)); pic.width = Emu(int(w)); pic.left = Emu(int((0.54*IN + 11.14*IN - w) / 2))
# ⑧ 5번
add_line(5, 'KS I ISO 14001:2015  환경경영시스템 — 요구사항 및 사용지침  (ISO 14001:2015)', 1.30, 0.30)
shrink_pic(5, 9, 1.62*IN, 6.55*IN)
# ⑦ 6번
add_line(6, 'KS I ISO 14044:2006  환경경영 — 전과정 평가 — 요구사항 및 지침  (ISO 14044:2006/Amd 1:2017/Amd 2:2020)  ·  원칙과 기본 구조는 KS I ISO 14040:2006', 1.30, 0.30, size=10.5)
shrink_pic(6, 8, 1.62*IN, 6.4*IN)
tf6 = prs.slides[5].notes_slide.notes_text_frame; t6 = tf6.text
tf6.text = '전과정평가(LCA)의 요구사항 표준은 KS I ISO 14044, 「환경경영 — 전과정 평가 — 요구사항 및 지침」입니다. 원칙과 틀은 ISO 14040에 있고 수행 요건은 14044에 있어 이 연구는 14044를 기준으로 삼았습니다.\n' + t6
tf6.text = tf6.text.replace('〔50초.', '〔55초.'); LOG.append((6, '발표 노트', '(첫머리)', '정식 명칭 한 문장', '이사님 10/14'))
tf5 = prs.slides[4].notes_slide.notes_text_frame; tf5.text = tf5.text.replace('두 표준이 각각 무엇인지부터 보겠습니다.', '두 표준이 각각 무엇인지부터 보겠습니다. 환경경영시스템(ISO 14001)의 정식 이름은 KS I ISO 14001, 「환경경영시스템 — 요구사항 및 사용지침」입니다.'); LOG.append((5, '발표 노트', '(첫머리)', '정식 명칭 한 문장', '이사님 10/14'))
# ⑨⑩ 매핑 장(35) 복귀 → 본 10번
m = prs.slides[34]
set_one(35, 2, '두 표준의 조항 매핑', '복귀 — 제목 단순화')
# 장 표시: 7번 것 복사(상자 모양)
old3 = shape(m, 3); old3._element.getparent().remove(old3._element)
m.shapes._spTree.append(copy.deepcopy(shape(prs.slides[6], 3)._element)); renumber_ids(m)
tag = [x for x in m.shapes if x.has_text_frame and x.text_frame.text.startswith('Ⅰ')][0]
tag.text_frame.paragraphs[0].runs[0].text = 'Ⅲ. 모델링 — 무엇을 만들었나'
for r in tag.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
tag.width = Emu(int(3.08*IN)); tag.left = Emu(int((11.15-3.08)*IN))
ids = {x.text_frame.text[:12]: x.shape_id for x in m.shapes if x.has_text_frame}
def by_prefix(pref):
    return [x for x in m.shapes if x.has_text_frame and x.text_frame.text.startswith(pref)][0]
g = by_prefix('▶'); g.text_frame.paragraphs[0].runs[-1].text = ('' if g.text_frame.paragraphs[0].runs[0].text.startswith('▶') and len(g.text_frame.paragraphs[0].runs)>1 else '▶ ') + '두 표준의 조항을 하나하나 대응시켜, 조항마다 네 가지 물음으로 네 등급을 매겼다 — 매우 높음 네 조항이 핵심 문항 4종.'
for r in g.text_frame.paragraphs[0].runs[1:-1]: r._r.getparent().remove(r._r)
LOG.append((35, '요지', '(shall·should 전부 대조…)', '두 표준의 조항을 하나하나 대응시켜 …', '쉬운 말'))
hd = by_prefix('네 가지 판정 질문'); hd.text_frame.paragraphs[0].runs[0].text = '네 가지 판정 질문 — ISO 14001 요구사항 조항 30개에 같은 순서로(조항마다 두세 번 물으면 등급이 정해짐)'
for r in hd.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
bt = by_prefix('ISO 14001 — 요구사항'); tfb = bt.text_frame
for q in list(tfb.paragraphs)[1:]: q._p.getparent().remove(q._p)
p0 = tfb.paragraphs[0]; p0.runs[0].text = 'ISO 14001 요구사항 조항 30개 전부 판정 · 매우 높음 4개 조항 = 핵심 문항 4종(P-9·P-14·D-1·C-1) — 관리 주기의 계획 둘·실행 하나·점검 하나 · 판정 흐름도는 [그림 3-3], 조항별 결과는 부록 A.1'
for r in p0.runs[1:]: r._r.getparent().remove(r._r)
bt.height = Emu(int(0.6*IN)); bt.top = Emu(int(6.3*IN)); LOG.append((35, '아래 줄', '30·25·28·3 셈', '30 전부 · 핵심 4 · 흐름도·A.1 위치', '셈은 노트북으로'))
rg = by_prefix('근거'); rg.text_frame.paragraphs[0].runs[0].text = '근거  3.2.2, [표 3-1], [표 3-2], [그림 3-3], 부록 A.1·B.1'
m.notes_slide.notes_text_frame.text = '3장의 출발은 매핑입니다. 두 표준의 조항을 하나하나 대응시켰습니다.\n조항마다 LCA가 그 조항을 지키는 데 얼마나 필요한지 네 가지 물음으로 네 등급을 매겼습니다. 전과정 관점을 직접 요구하는가, 없으면 의무를 못 지키는가, 주요 입력인가, 보조 정보라도 되는가이고, 예가 앞에서 나올수록 등급이 높습니다. 조항마다 두 번 또는 세 번 물으면 등급이 정해집니다.\n매우 높음이 네 조항, 환경측면, 환경목표, 운용기획 및 관리, 모니터링입니다. 이 네 곳이 뒤의 핵심 문항 4종이 되고, 관리 주기의 계획에 둘, 실행과 점검에 하나씩 놓입니다. ISO 14001의 요구사항 조항 서른 개는 모두 문항이 되었습니다.\n〔1분 10초. 30 = 부록 A.1 원표 34행에서 7.4.1~7.4.3을 하나로, 상위 조항 6.1·6.2 제외(3.2.2). ISO 14044는 전조항 28개 가운데 25개가 문항의 근거, 나머지 3개(4.2.1·4.2.3.1 일반 조항, 4.4.3)도 세부 조항 문항·체크포인트로 반영 → 빠진 조항 없음(부록 A.4). 대조 대상은 shall·should — may·can 제외(3.1.1). 55문항 근거 성격은 노트북. 기준은 B.1(등급 정의·징후·질문 흐름), 조항별 결과는 A.1(연계 가능성 열·연계 방안 열) — 「기준을 먼저 정하고 전 조항에 적용」(B551·B555). 흐름도 [그림 3-3]은 백업 B-1, 조항별 질문 경로(재구성)는 B-1b, 등급 요건은 B-1a.〕'
LOG.append((35, '발표 노트(전체)', '(B-1 지시문)', '매핑 세 문장 + 지시문', '이사님 10/14'))
# [그림 3-3] 장(10) → 백업 B-1
f = prs.slides[9]; t3 = shape(f, 3); t3._element.getparent().remove(t3._element)
f.shapes._spTree.append(copy.deepcopy(shape(prs.slides[34], 3)._element)) if False else None
# 백업 표시는 B-1a(36번) 것 복사
f.shapes._spTree.append(copy.deepcopy(shape(prs.slides[35], 3)._element)); renumber_ids(f)
bt3 = [x for x in f.shapes if x.has_text_frame and x.text_frame.text.startswith('백업')][0]
bt3.text_frame.paragraphs[0].runs[0].text = '백업  B-1'
for r in bt3.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
f.notes_slide.notes_text_frame.text = '「어떤 순서로 물었나」를 묻는 경우 — 논문 [그림 3-3]. 조항 하나를 넣으면 두세 번의 물음을 거쳐 오른쪽 네 상자 가운데 하나로 떨어진다. Q1 예→Q2(예 매우 높음 / 아니오 높음), Q1 아니오→Q3(예 높음 / 아니오 Q4: 예 중간 / 아니오 낮음). 3.2.2, 부록 B.1. 조항별 경로는 B-1b.'
LOG.append((10, '[그림 3-3] 장', '본 10번', '백업 B-1', '이사님 10/14'))
# 백업 B-1b: 조항별 판정 질문 경로(재구성) 표
tpl = prs.slides[6]; nb = prs.slides.add_slide(tpl.slide_layout)
for sh in tpl.shapes:
    if sh.shape_id in (2, 4, 5, 6): nb.shapes._spTree.append(copy.deepcopy(sh._element))
nb.shapes._spTree.append(copy.deepcopy(shape(prs.slides[35], 3)._element)); renumber_ids(nb)
tt = shape(nb, 2); tt.text_frame.paragraphs[0].runs[0].text = '조항별 판정 질문 경로 — 부록 A.1 등급에서 재구성'
for r in tt.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
tt.text_frame.paragraphs[0].runs[0].font.size = Pt(23)
bb = [x for x in nb.shapes if x.has_text_frame and x.text_frame.text.startswith('백업')][0]; bb.text_frame.paragraphs[0].runs[0].text = '백업  B-1b'
for r in bb.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
gg = [x for x in nb.shapes if x.has_text_frame and x.text_frame.text.startswith('▶')][0]
gg.text_frame.paragraphs[0].runs[-1].text = '등급은 질문 흐름에서 기계적으로 나오므로 조항별 경로는 등급에서 역으로 확인된다 — 조항마다 두 번 또는 세 번 묻는다.'
for r in gg.text_frame.paragraphs[0].runs[1:-1]: r._r.getparent().remove(r._r)
rows = [('등급', '질문 경로', '물음', '조항(부록 A.1 등급)'),
        ('매우 높음 4', 'Q1 예 → Q2 예', '2', '6.1.2 환경측면 · 6.2.1 환경목표 · 8.1 운용기획 및 관리 · 9.1.1 모니터링·측정·분석·평가'),
        ('높음 12 (a)', 'Q1 예 → Q2 아니오 — 전과정 관점이나 그 결과(중대한 환경측면·환경목표)를 직접 받지만 LCA 없이도 이행은 됨', '2', '6.1.1 리스크·기회 일반 · 6.1.4 조치 계획 · 6.2.2 목표 달성 조치 · 9.3 경영검토 · 10.1 개선 일반 · 10.3 지속적 개선'),
        ('높음 12 (b)', 'Q1 아니오 → Q3 예 — 전과정 관점 요구는 없지만 LCA 결과가 주요 입력', '2', '4.2 이해관계자 · 4.3 적용범위 · 5.2 환경방침 · 7.2 역량 · 7.4 의사소통 · 7.5.1 문서화된 정보 일반'),
        ('중간 10', 'Q1 아니오 → Q3 아니오 → Q4 예', '3', '4.1 · 4.4 · 5.1 · 5.3 · 7.1 · 7.3 · 7.5.2 · 7.5.3 · 9.2.1 · 9.2.2'),
        ('낮음 4', 'Q1 아니오 → Q3 아니오 → Q4 아니오', '3', '6.1.3 준수의무사항 · 8.2 비상사태 · 9.1.2 준수평가 · 10.2 부적합·시정조치')]
tb = nb.shapes.add_table(len(rows), 4, Emu(int(0.54*IN)), Emu(int(1.4*IN)), Emu(int(10.6*IN)), Emu(int(4.6*IN))).table
for j, w in enumerate((1.3, 4.0, 0.6, 4.7)): tb.columns[j].width = Emu(int(w*IN))
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tb.cell(i, j); c.text = v
        for pp in c.text_frame.paragraphs:
            for r in pp.runs: r.font.size = Pt(12 if i else 12.5); r.font.name = '맑은 고딕'; r.font.bold = (i == 0 or j == 0)
        c.margin_left = c.margin_right = Emu(int(0.06*IN))
nb.shapes._spTree.append(copy.deepcopy(shape(prs.slides[6], 9)._element)); renumber_ids(nb)
nt = [x for x in nb.shapes if x.has_text_frame and x.text_frame.text.startswith('상태 A')][0]
nt.top = Emu(int(6.2*IN)); nt.height = Emu(int(1.0*IN)); p = nt.text_frame.paragraphs[0]; p.runs[0].text = '기준(등급 정의·징후·질문 흐름)은 부록 B.1, 조항별 결과(연계 가능성 등급·연계 방안)는 부록 A.1 — 기준을 먼저 정하고 전 조항에 적용(3.2.2). 부록 A.1에는 질문별 예·아니오 열이 없어 이 표는 등급에서 재구성한 것이며, 높음 12개의 (a)·(b) 나눔은 A.1의 근거 문구(환경측면·환경목표를 참조하면 Q1 예, LCA 결과를 입력으로 하면 Q3 예)에 따름. 본심사본 부록 A.1에 「판정 경로」 열로 추가 예정.'
for r in p.runs[1:]: r._r.getparent().remove(r._r)
for q in list(nt.text_frame.paragraphs)[1:]: q._p.getparent().remove(q._p)
for r in p.runs: r.font.size = Pt(11.5)
nb.notes_slide.notes_text_frame.text = '「조항별로 Q1~Q4에 어떻게 답했나, 기록이 있나」 — 기준은 B.1, 결과는 A.1에 있고, 경로는 등급에서 도출된다. 질문별 예·아니오 표는 따로 두지 않았고 본심사본 A.1에 열로 추가하겠다. 높음 12개의 두 경로 나눔은 이사님 확인 사항.'
n_b1b = len(prs.slides); LOG.append((n_b1b, '새 백업 B-1b', '(없음)', '조항별 판정 질문 경로 — 재구성 표', '이사님 10/14'))
# ---- 순서: 1..9, 35(매핑), 11..34, 10([그림 3-3]→B-1), B-1b, 36..N-1
N = len(prs.slides)
order = list(range(1, 10)) + [35] + list(range(11, 35)) + [10, N] + list(range(36, N))
assert sorted(order) == list(range(1, N+1)), (len(order), N)
old2new = {o: i+1 for i, o in enumerate(order)}
L = prs.slides._sldIdLst; ids = list(L)
for e in ids: L.remove(e)
for o in order: L.append(ids[o-1])
def tr_nums(text):
    def one(mm):
        nums = re.findall(r'\d+', mm.group(0)); seps = re.findall(r'[·~,]\s?', mm.group(0))
        out = mm.group(0)[:mm.group(0).index(nums[0])]
        for k, v in enumerate(nums): out += str(old2new.get(int(v), int(v))) + (seps[k] if k < len(seps) else '')
        return out
    return re.sub(r'백업\s*\d+(?:[·~,]\s?\d+)*', one, text)
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(7.3*IN)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tf = sl.notes_slide.notes_text_frame; tx = tf.text; t2 = tr_nums(tx)
    if t2 != tx: tf.text = t2
prs.save(dst); json.dump({'log': LOG, 'old2new': old2new}, open('log32c.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ok', len(LOG), len(prs.slides))
